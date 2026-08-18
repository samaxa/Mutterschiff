#!/usr/bin/env python3
"""
Importiert Bank-Umsätze aus einer CSV-Datei in den Finanzplaner.

Aufruf:
    python import_umsaetze.py Umsaetze.csv
    python import_umsaetze.py Umsaetze.csv --excel Finanzplaner_Sarah.xlsx
    python import_umsaetze.py Umsaetze.csv --dry-run     # nur anzeigen, nichts schreiben

Getestet mit dem Sparkassen-Export "CSV-CAMT". Andere Banken (ING, DKB,
Comdirect, N26 ...) werden über die Spaltenerkennung unten meist automatisch
mitgenommen.

Benötigt nur openpyxl:  pip install openpyxl
"""

import argparse
import csv
import datetime as dt
import hashlib
import io
import re
import warnings
from pathlib import Path

warnings.filterwarnings("ignore", module="openpyxl")

from openpyxl import load_workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Font

# =====================================================================
#  KATEGORIE-REGELN  –  hier trägst du mit der Zeit deine eigenen ein
# =====================================================================
# Regel = (Suchmuster, Kategorie). Das Suchmuster ist ein regulärer Ausdruck
# und wird gegen Empfänger + Verwendungszweck + Buchungstext geprüft
# (Groß-/Kleinschreibung egal). Die ERSTE passende Regel gewinnt –
# spezielle Regeln also nach oben, allgemeine nach unten.

REGELN = [
    # --- Einnahmen ---
    (r"uniper|gehalt|lohn|bez[üu]ge",                      "Gehalt Werkstudent"),
    (r"b(oe|ö|o)ckler|stipendi|f[öo]rderung",              "Stipendium"),

    # --- Wohnen ---
    (r"miete|hausverwaltung|nebenkosten|kaltmiete",        "Miete & Nebenkosten"),
    (r"stadtwerke|vattenfall|eon|e\.on|rwe|innogy|"
     r"telekom|vodafone|1und1|1&1|o2|unitymedia|"
     r"pyur|gasag|strom|internet",                         "Strom & Internet"),

    # --- Versicherung & Finanzen ---
    (r"versicherung|allianz|huk|axa|ergo|debeka|barmenia|"
     r"techniker krankenkasse|aok|barmer|dak|kkh",         "Versicherungen"),

    # --- Essen ---
    (r"rewe|edeka|aldi|lidl|penny|netto|kaufland|"
     r"real markt|nahkauf|tegut|denns|alnatura|bio company|"
     r"supermarkt|backerei|b[äa]ckerei|metzger",           "Lebensmittel"),
    (r"lieferando|wolt|uber ?eats|too good to go|"
     r"restaurant|pizzeria|mcdonald|burger king|kfc|"
     r"subway|starbucks|cafe|caf[ée]|imbiss|doner|d[öo]ner",
                                                           "Restaurant & Lieferdienst"),

    # --- Mobilität ---
    (r"tankstelle|aral|shell|esso|jet |total |sprit|"
     r"deutsche bahn|db vertrieb|bahn\.de|rheinbahn|vrr|"
     r"deutschlandticket|9-euro|flixbus|adac|"
     r"werkstatt|tuv|t[üu]v|dekra|kfz|autohaus|parkhaus",  "Mobilität & Auto"),

    # --- Drogerie & Gesundheit ---
    (r"\bdm\b|dm-drogerie|rossmann|m[üu]ller drogerie|"
     r"apotheke|zahnarzt|arztpraxis|praxis dr|"
     r"doctolib|shop apotheke",                            "Gesundheit & Drogerie"),

    # --- Freizeit & Sport ---
    (r"fitness|mcfit|urban sports|gym|schwimmbad|"
     r"kino|cinemaxx|uci |theater|museum|tanzschule",       "Sport & Freizeit"),

    # --- Kleidung ---
    (r"zalando|h&m|h und m|zara|c&a|about you|"
     r"uniqlo|snipes|deichmann|tk maxx|primark|decathlon",  "Kleidung"),

    # --- Abos & Software ---
    (r"spotify|netflix|disney|amazon prime|apple\.com|"
     r"itunes|google \*|microsoft|adobe|dropbox|"
     r"audible|youtube ?premium|patreon|openai|anthropic",  "Abos & Software"),

    # --- Studium ---
    (r"th k[öo]ln|hochschule|semesterbeitrag|studierenden|"
     r"asta|thalia|hugendubel|buchhandlung|springer|elsevier",
                                                           "Studium & Bücher"),

    # --- Reisen ---
    (r"booking\.com|airbnb|hotel|hostel|ryanair|eurowings|"
     r"lufthansa|easyjet|flug|reise|urlaub",                "Urlaub & Reisen"),

    # --- Geschenke ---
    (r"geschenk|blumen|floristik",                          "Geschenke"),
]

# Fallback, wenn keine Regel greift
FALLBACK_AUSGABE = "Sonstiges"
FALLBACK_EINNAHME = "Sonstige Einnahmen"

# Standard-Konto, das in Spalte H eingetragen wird
STANDARD_KONTO = "Girokonto"

# =====================================================================
#  Ab hier musst du nichts mehr anfassen
# =====================================================================

BUCHUNGEN_ERSTE_ZEILE = 4
BUCHUNGEN_LETZTE_ZEILE = 1003
HASH_SPALTE = 10  # Spalte J – technische Import-ID zur Dublettenerkennung

# Spaltennamen, die in Bank-Exporten für dasselbe stehen
SPALTEN_DATUM = ["buchungstag", "buchungsdatum", "datum", "buchung", "wertstellung", "valutadatum"]
SPALTEN_BETRAG = ["betrag", "betrag (eur)", "umsatz", "soll/haben", "amount"]
SPALTEN_ZWECK = ["verwendungszweck", "buchungstext", "beschreibung", "vorgang/verwendungszweck", "referenz"]
SPALTEN_PARTNER = ["beguenstigter/zahlungspflichtiger", "begünstigter/zahlungspflichtiger",
                   "auftraggeber/empfänger", "auftraggeber/empfaenger", "empfänger",
                   "name zahlungsbeteiligter", "zahlungsbeteiligter", "payee"]
SPALTEN_INFO = ["info", "status"]


def lese_csv(pfad: Path):
    """Liest die CSV robust ein: Kodierung und Trennzeichen werden erkannt."""
    rohdaten = pfad.read_bytes()
    text = None
    for kodierung in ("utf-8-sig", "utf-8", "cp1252", "iso-8859-1"):
        try:
            text = rohdaten.decode(kodierung)
            break
        except UnicodeDecodeError:
            continue
    if text is None:
        raise SystemExit("Kodierung der Datei konnte nicht erkannt werden.")

    zeilen = [z for z in text.splitlines() if z.strip()]
    if not zeilen:
        raise SystemExit("Die Datei ist leer.")

    # Manche Banken schreiben Vorspann-Zeilen vor die eigentliche Kopfzeile.
    kopf_index = 0
    for i, zeile in enumerate(zeilen[:15]):
        klein = zeile.lower()
        if any(n in klein for n in SPALTEN_DATUM) and any(n in klein for n in SPALTEN_BETRAG):
            kopf_index = i
            break

    nutzdaten = "\n".join(zeilen[kopf_index:])
    trenner = ";" if nutzdaten.count(";") >= nutzdaten.count(",") else ","
    if nutzdaten.count("\t") > nutzdaten.count(trenner):
        trenner = "\t"

    leser = csv.DictReader(io.StringIO(nutzdaten), delimiter=trenner)
    return list(leser), leser.fieldnames or []


def finde_spalte(kopfzeilen, kandidaten):
    """Sucht die passende Spalte, erst exakt, dann als Teilstring."""
    normal = {(k or "").strip().lower().strip('"'): k for k in kopfzeilen}
    for kandidat in kandidaten:
        if kandidat in normal:
            return normal[kandidat]
    for name, original in normal.items():
        for kandidat in kandidaten:
            if kandidat in name:
                return original
    return None


def parse_betrag(wert: str):
    """'-1.234,56' oder '832,9' oder '-190' -> float"""
    if wert is None:
        return None
    w = str(wert).strip().replace('"', "").replace("\u00a0", " ").replace(" ", "")
    w = re.sub(r"(EUR|€)", "", w, flags=re.I)
    if not w:
        return None
    negativ = w.startswith("-") or w.endswith("-") or w.upper().endswith("S")
    w = w.strip("-").rstrip("SHsh")
    if "," in w:                      # deutsches Format
        w = w.replace(".", "").replace(",", ".")
    elif w.count(".") == 1 and len(w.split(".")[-1]) in (1, 2):
        pass                          # englisches Format, passt schon
    else:
        w = w.replace(".", "")
    try:
        betrag = float(w)
    except ValueError:
        return None
    return -betrag if negativ and betrag > 0 else betrag


def parse_datum(wert: str):
    """Erkennt 03.08.26, 03.08.2026, 2026-08-03 ..."""
    if not wert:
        return None
    w = str(wert).strip().replace('"', "")
    for muster in ("%d.%m.%y", "%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
        try:
            datum = dt.datetime.strptime(w, muster).date()
            if datum.year < 1990:     # zweistellige Jahre wie "26"
                datum = datum.replace(year=datum.year + 100)
            return datum
        except ValueError:
            continue
    return None


def kategorisiere(text: str, betrag: float, typ_je_kategorie: dict):
    """Ordnet eine Buchung anhand der Regeln einer Kategorie zu."""
    for muster, kategorie in REGELN:
        if re.search(muster, text, flags=re.I):
            erwartet = "Einnahme" if betrag > 0 else "Ausgabe"
            # Regel passt inhaltlich, aber Vorzeichen widerspricht dem Typ
            # (z. B. Rückerstattung von REWE) -> Fallback, damit nichts falsch
            # als Einnahme/Ausgabe verbucht wird.
            if typ_je_kategorie.get(kategorie) == erwartet:
                return kategorie, True
            break
    return (FALLBACK_EINNAHME if betrag > 0 else FALLBACK_AUSGABE), False


def import_id(datum, betrag, partner, zweck):
    roh = f"{datum.isoformat()}|{betrag:.2f}|{partner.strip().lower()}|{zweck.strip().lower()}"
    return hashlib.sha1(roh.encode("utf-8")).hexdigest()[:16]


def diagramme_neu_aufbauen(wb):
    """
    openpyxl verliert beim Öffnen/Speichern alle Diagramme einer Datei.
    Deshalb werden die fünf Dashboard-Diagramme hier neu erzeugt.
    """
    da, mo, ve, pr, ka = (wb["Dashboard"], wb["Monatsübersicht"], wb["Vermögen"],
                          wb["Prognose"], wb["Kategorien"])
    da._charts = []

    monate = Reference(mo, min_col=1, min_row=4, max_row=27)

    c1 = BarChart(); c1.type = "col"; c1.title = "Einnahmen vs. Ausgaben pro Monat"
    c1.y_axis.title = "€"; c1.height, c1.width = 8.5, 20
    c1.add_data(Reference(mo, min_col=4, max_col=5, min_row=3, max_row=27), titles_from_data=True)
    c1.set_categories(monate); da.add_chart(c1, "D4")

    c2 = LineChart(); c2.title = "Kumulierter Saldo"; c2.y_axis.title = "€"
    c2.height, c2.width = 8.5, 20
    c2.add_data(Reference(mo, min_col=8, min_row=3, max_row=27), titles_from_data=True)
    c2.set_categories(monate); da.add_chart(c2, "D22")

    c3 = LineChart(); c3.title = "Nettovermögen"; c3.y_axis.title = "€"
    c3.height, c3.width = 8.5, 20
    c3.add_data(Reference(ve, min_col=9, min_row=3, max_row=27), titles_from_data=True)
    c3.set_categories(Reference(ve, min_col=1, min_row=4, max_row=27)); da.add_chart(c3, "P4")

    c4 = LineChart(); c4.title = "Prognose Nettovermögen (12 Monate)"; c4.y_axis.title = "€"
    c4.height, c4.width = 8.5, 20
    c4.add_data(Reference(pr, min_col=3, min_row=16, max_row=29), titles_from_data=True)
    c4.set_categories(Reference(pr, min_col=2, min_row=17, max_row=29)); da.add_chart(c4, "P22")

    c5 = BarChart(); c5.type = "bar"; c5.title = "Ausgaben je Kategorie (gewähltes Jahr)"
    c5.height, c5.width = 12, 20
    c5.add_data(Reference(ka, min_col=19, min_row=5, max_row=42), titles_from_data=True)
    c5.set_categories(Reference(ka, min_col=1, min_row=6, max_row=42)); da.add_chart(c5, "D40")


def main():
    parser = argparse.ArgumentParser(description="Bank-CSV in den Finanzplaner importieren")
    parser.add_argument("csv_datei", help="Export aus dem Online-Banking")
    parser.add_argument("--excel", default="Finanzplaner_Sarah.xlsx", help="Pfad zur Excel-Datei")
    parser.add_argument("--konto", default=STANDARD_KONTO, help="Konto-Bezeichnung für Spalte H")
    parser.add_argument("--dry-run", action="store_true", help="nur anzeigen, nichts schreiben")
    args = parser.parse_args()

    csv_pfad, excel_pfad = Path(args.csv_datei), Path(args.excel)
    for pfad in (csv_pfad, excel_pfad):
        if not pfad.exists():
            raise SystemExit(f"Datei nicht gefunden: {pfad}")

    zeilen, kopf = lese_csv(csv_pfad)
    if not zeilen:
        raise SystemExit("Keine Datenzeilen in der CSV gefunden.")

    sp_datum = finde_spalte(kopf, SPALTEN_DATUM)
    sp_betrag = finde_spalte(kopf, SPALTEN_BETRAG)
    sp_zweck = finde_spalte(kopf, SPALTEN_ZWECK)
    sp_partner = finde_spalte(kopf, SPALTEN_PARTNER)
    sp_info = finde_spalte(kopf, SPALTEN_INFO)

    if not sp_datum or not sp_betrag:
        raise SystemExit(f"Datums- oder Betragsspalte nicht erkannt.\nGefundene Spalten: {kopf}")

    print(f"Spalten erkannt  →  Datum: {sp_datum} | Betrag: {sp_betrag} | "
          f"Zweck: {sp_zweck} | Empfänger: {sp_partner}\n")

    wb = load_workbook(excel_pfad)
    bu, ein = wb["Buchungen"], wb["Einstellungen"]

    # Kategorien und ihre Typen aus dem Blatt Einstellungen einlesen
    typ_je_kategorie = {}
    for r in range(4, 41):
        name, typ = ein.cell(row=r, column=1).value, ein.cell(row=r, column=2).value
        if name and typ:
            typ_je_kategorie[str(name).strip()] = str(typ).strip()

    # Bereits importierte Buchungen einsammeln
    bekannt = set()
    for r in range(BUCHUNGEN_ERSTE_ZEILE, BUCHUNGEN_LETZTE_ZEILE + 1):
        h = bu.cell(row=r, column=HASH_SPALTE).value
        if h:
            bekannt.add(str(h))

    # Erste freie Zeile suchen
    naechste_zeile = None
    for r in range(BUCHUNGEN_ERSTE_ZEILE, BUCHUNGEN_LETZTE_ZEILE + 1):
        if bu.cell(row=r, column=1).value is None:
            naechste_zeile = r
            break
    if naechste_zeile is None:
        raise SystemExit("Das Blatt Buchungen ist voll (1000 Zeilen).")

    neu, doppelt, uebersprungen, ohne_regel = [], 0, 0, []

    for zeile in zeilen:
        info = (zeile.get(sp_info) or "") if sp_info else ""
        if "vorgemerkt" in info.lower():
            uebersprungen += 1          # noch nicht endgültig gebucht
            continue

        datum = parse_datum(zeile.get(sp_datum))
        betrag = parse_betrag(zeile.get(sp_betrag))
        if datum is None or betrag is None or betrag == 0:
            uebersprungen += 1
            continue

        partner = (zeile.get(sp_partner) or "").strip() if sp_partner else ""
        zweck = (zeile.get(sp_zweck) or "").strip() if sp_zweck else ""
        suchtext = " ".join(str(v) for v in zeile.values() if v)

        kennung = import_id(datum, betrag, partner, zweck)
        if kennung in bekannt:
            doppelt += 1
            continue
        bekannt.add(kennung)

        kategorie, getroffen = kategorisiere(suchtext, betrag, typ_je_kategorie)
        if not getroffen:
            ohne_regel.append((datum, betrag, partner or zweck[:40]))

        beschreibung = " – ".join(t for t in (partner, zweck) if t)[:200]
        neu.append((datum, kategorie, beschreibung, abs(betrag), kennung))

    neu.sort(key=lambda x: x[0])

    print(f"{len(neu)} neue Buchungen | {doppelt} bereits vorhanden | "
          f"{uebersprungen} übersprungen (vorgemerkt/unlesbar)")

    if ohne_regel:
        print(f"\n{len(ohne_regel)} Buchungen ohne passende Regel "
              f"(gelandet in '{FALLBACK_AUSGABE}' / '{FALLBACK_EINNAHME}'):")
        for datum, betrag, wer in ohne_regel[:25]:
            print(f"   {datum:%d.%m.%Y}  {betrag:>10.2f} €   {wer}")
        print("   → passende Regel oben in REGELN ergänzen, dann erneut importieren.")

    if args.dry_run:
        print("\n--dry-run: es wurde nichts geschrieben.")
        return

    if not neu:
        print("\nNichts zu tun.")
        return

    grau = Font(name="Arial", size=8, color="A6A6A6")
    for datum, kategorie, beschreibung, betrag, kennung in neu:
        if naechste_zeile > BUCHUNGEN_LETZTE_ZEILE:
            print(f"Achtung: Platz ist voll, ab {datum:%d.%m.%Y} wurde nichts mehr importiert.")
            break
        r = naechste_zeile
        bu.cell(row=r, column=1).value = datum
        bu.cell(row=r, column=1).number_format = "DD.MM.YYYY"
        bu.cell(row=r, column=4).value = kategorie
        bu.cell(row=r, column=5).value = beschreibung
        bu.cell(row=r, column=6).value = betrag
        bu.cell(row=r, column=8).value = args.konto
        h = bu.cell(row=r, column=HASH_SPALTE)
        h.value = kennung
        h.font = grau
        naechste_zeile += 1

    kopfzelle = bu.cell(row=3, column=HASH_SPALTE)
    if not kopfzelle.value:
        kopfzelle.value = "Import-ID"
        kopfzelle.font = grau

    diagramme_neu_aufbauen(wb)
    wb.save(excel_pfad)
    print(f"\nFertig – {len(neu)} Buchungen geschrieben nach {excel_pfad}")
    print("Excel berechnet beim Öffnen automatisch neu.")


if __name__ == "__main__":
    main()
