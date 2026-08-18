#!/usr/bin/env python3
"""
Einmal-Skript: macht den Finanzplaner übersichtlicher.

    python finanzplaner_verschoenern.py
    python finanzplaner_verschoenern.py --excel Finanzplaner_Sarah.xlsx

Was es ändert:
  * Blatt "Anleitung"  – neue Übersicht deiner Konten (mit Balken), Kurzanleitung
                         zum Import, kompakte Blatt- und Farblegende.
  * Blatt "Vermögen"   – Spaltenköpfe auf deine echten Konten umbenannt, der
                         Stand vom 18.08.2026 eingetragen (nur falls noch leer).
  * Blatt "Dashboard"  – Block mit den aktuellen Kontoständen, den das neue
                         Diagramm "Kontostände je Konto" auswertet.
  * Diagramme          – neu aufgebaut mit sprechenden Titeln und Erklärsatz.

Deine Buchungen werden nicht angefasst. Das Skript ist mehrfach ausführbar.

Benötigt nur openpyxl:  pip install openpyxl
"""

import argparse
import warnings
from pathlib import Path

warnings.filterwarnings("ignore", module="openpyxl")

from openpyxl import load_workbook
from openpyxl.formatting.rule import DataBarRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from import_umsaetze import diagramme_neu_aufbauen

# =====================================================================
#  Deine Konten – Quelle: Sparkassen-Finanzübersicht vom 18.08.2026
#  (Spalte = Spalte im Blatt Vermögen, in der der Stand landet)
# =====================================================================
KONTEN = [
    # Bezeichnung,            Inhaber,                        Nummer,                        Vermögen-Spalte, Stand
    ("S-POOL light+",         "Sarah Maxara",                 "DE25 3055 0000 1005 4402 90", 2,  64.23,  "Girokonto"),
    ("Giro Express Plus",     "Marvin Liskatin, Sarah Maxara", "DE77 3055 0000 1005 4403 24", 6,  62.36,  "Gemeinschaftskonto"),
    ("Tagesgeldkonto+",       "Sarah Maxara",                 "DE59 3055 0000 1005 4403 57", 3,   0.00,  "Tagesgeld / Sparkonto"),
    ("Gold Privat (Kreditkarte)", "Sarah Maxara",             "5232 **** **** 6059",         8, 522.74,  "Kreditkarte"),
]
STAND_ZEILE = 11          # Blatt Vermögen: Zeile 4 = Januar 2026 -> Zeile 11 = August 2026
STAND_TEXT = "18.08.2026"

# Farben und Schriften der Datei (nicht ändern – sonst passt es nicht zum Rest)
DUNKELBLAU = "1F3864"
GELB = "FFF2CC"
GRAU = "595959"
EURO = '#,##0.00" €";[RED]\\-#,##0.00" €"'

F_TITEL = Font(name="Arial", size=14, bold=True, color=DUNKELBLAU)
F_UNTERTITEL = Font(name="Arial", size=10, color=GRAU)
F_ABSCHNITT = Font(name="Arial", size=11, bold=True, color=DUNKELBLAU)
F_KOPF = Font(name="Arial", size=10, bold=True, color="FFFFFF")
F_NORMAL = Font(name="Arial", size=10)
F_FETT = Font(name="Arial", size=10, bold=True)
F_KLEIN = Font(name="Arial", size=9, italic=True, color=GRAU)
F_EINGABE = Font(name="Arial", size=10, color="0000FF")

FILL_KOPF = PatternFill("solid", fgColor=DUNKELBLAU)
FILL_GELB = PatternFill("solid", fgColor=GELB)
RAHMEN = Border(*(Side(style="thin", color="BFBFBF"),) * 4)
OBEN = Alignment(vertical="top", wrap_text=True)


def _zelle(ws, koord, wert, schrift=F_NORMAL, fill=None, format=None, align=None):
    z = ws[koord]
    z.value = wert
    z.font = schrift
    if fill:
        z.fill = fill
    if format:
        z.number_format = format
    if align:
        z.alignment = align
    return z


def _absatz(ws, zeile, label, text, hoehe=30):
    """Zeile mit Stichwort in A und beschreibendem Text über B:E."""
    _zelle(ws, f"A{zeile}", label, F_FETT, align=Alignment(vertical="top"))
    ws.merge_cells(f"B{zeile}:E{zeile}")
    _zelle(ws, f"B{zeile}", text, F_NORMAL, align=OBEN)
    ws.row_dimensions[zeile].height = hoehe


def anleitung_neu(wb):
    """Blatt Anleitung komplett neu aufbauen."""
    position = wb.sheetnames.index("Anleitung")
    wb.remove(wb["Anleitung"])
    ws = wb.create_sheet("Anleitung", position)
    ws.sheet_view.showGridLines = False
    for spalte, breite in (("A", 26), ("B", 24), ("C", 28), ("D", 15), ("E", 26)):
        ws.column_dimensions[spalte].width = breite

    _zelle(ws, "A1", "Finanzplaner – Anleitung", F_TITEL)
    _zelle(ws, "A2", "Erstellt für Sarah · Alle Beträge in Euro (€)", F_UNTERTITEL)

    # ---------------- Konten ----------------
    _zelle(ws, "A4", "Deine Konten und Karten", F_ABSCHNITT)
    ws.merge_cells("A5:E5")
    _zelle(ws, "A5", f"Stand {STAND_TEXT}. Die Salden holt sich die Tabelle aus dem Blatt "
                     "Vermögen – trag dort einmal im Monat die Kontostände ein, dann bleibt "
                     "diese Übersicht aktuell.", F_KLEIN, align=OBEN)
    ws.row_dimensions[5].height = 28

    for spalte, titel in zip("ABCDE", ("Konto", "Inhaber", "IBAN / Kartennummer",
                                       "Saldo (€)", 'Import: --konto "…"')):
        _zelle(ws, f"{spalte}6", titel, F_KOPF, FILL_KOPF)

    for i, (name, inhaber, nummer, spalte, _stand, kontoname) in enumerate(KONTEN):
        z = 7 + i
        buchstabe = chr(ord("A") + spalte - 1)
        vorzeichen = "-" if spalte == 8 else ""        # Schulden nach unten
        _zelle(ws, f"A{z}", name, F_NORMAL)
        _zelle(ws, f"B{z}", inhaber, F_NORMAL)
        _zelle(ws, f"C{z}", nummer, F_NORMAL)
        _zelle(ws, f"D{z}",
               f"={vorzeichen}IFERROR(LOOKUP(9.99999999999999E+307,"
               f"Vermögen!${buchstabe}$4:${buchstabe}$27),0)", F_NORMAL, format=EURO)
        _zelle(ws, f"E{z}", kontoname, F_NORMAL)
        for spaltenbuchstabe in "ABCDE":
            ws[f"{spaltenbuchstabe}{z}"].border = RAHMEN

    summenzeile = 7 + len(KONTEN)
    _zelle(ws, f"A{summenzeile}", "Nettovermögen (Summe)", F_FETT)
    _zelle(ws, f"D{summenzeile}", f"=SUM(D7:D{summenzeile - 1})", F_FETT, format=EURO)
    for spaltenbuchstabe in "ABCDE":
        ws[f"{spaltenbuchstabe}{summenzeile}"].border = RAHMEN

    # Balken direkt in der Saldo-Spalte – zeigt die Größenverhältnisse auf einen Blick
    ws.conditional_formatting.add(
        f"D7:D{summenzeile - 1}",
        DataBarRule(start_type="num", start_value=-600, end_type="num", end_value=600,
                    color="4472C4", showValue=True))

    hinweis = summenzeile + 1
    ws.merge_cells(f"A{hinweis}:E{hinweis}")
    _zelle(ws, f"A{hinweis}",
           "Hinweis: Das Giro Express Plus führst du gemeinsam mit Marvin – hier zählt es "
           "voll als dein Guthaben. Willst du nur die Hälfte rechnen, trag im Blatt Vermögen "
           "entsprechend den halben Betrag ein.", F_KLEIN, align=OBEN)
    ws.row_dimensions[hinweis].height = 28

    # ---------------- Import ----------------
    z = hinweis + 2
    _zelle(ws, f"A{z}", "In drei Schritten aktuell halten", F_ABSCHNITT)
    _absatz(ws, z + 1, "1. Exportieren",
            "Für jedes der vier Konten im Online-Banking die Umsätze als CSV herunterladen "
            "und in denselben Ordner legen wie diese Datei.")
    _absatz(ws, z + 2, "2. Importieren",
            'Diese Datei schließen. Dann im Terminal je CSV einmal aufrufen, z. B.:  '
            'python import_umsaetze.py umsaetze.CSV --konto "Girokonto"   '
            "– den passenden Namen findest du oben in der letzten Spalte.", hoehe=42)
    _absatz(ws, z + 3, "3. Öffnen",
            "Datei wieder öffnen. Monatsübersicht, Kategorien, Dashboard und alle Diagramme "
            "haben sich automatisch neu berechnet – nichts abtippen.")
    _absatz(ws, z + 4, "Gut zu wissen",
            "Ein Import kann nichts doppelt eintragen: bereits vorhandene Buchungen werden "
            "an der Import-ID in Spalte J erkannt und übersprungen. Mit --dry-run siehst du "
            "vorab, was passieren würde, ohne dass etwas geschrieben wird.", hoehe=42)

    # ---------------- Blätter ----------------
    z += 6
    _zelle(ws, f"A{z}", "Die Blätter der Datei", F_ABSCHNITT)
    blaetter = [
        ("Dashboard", "Startseite: Kennzahlen des laufenden Monats, deine Kontostände und "
                      "die fünf Diagramme. Über jedem Diagramm steht, was es zeigt."),
        ("Buchungen", "Das Herzstück – jede Einnahme und Ausgabe als eigene Zeile. Füllt "
                      "normalerweise das Import-Skript; du kannst aber jederzeit von Hand "
                      "korrigieren, vor allem die Kategorie in Spalte D."),
        ("Einstellungen", "Deine Kategorien, das Monatsbudget je Kategorie und die Liste der "
                          "Konten. Neue Kategorie? Einfach in eine freie Zeile schreiben."),
        ("Fixkosten", "Wiederkehrende Kosten mit Rhythmus. Jährliche und vierteljährliche "
                      "Beträge werden auf den Monat umgerechnet."),
        ("Monatsübersicht", "Pro Monat: Einnahmen, Ausgaben, Saldo, Sparquote und der "
                            "kumulierte Verlauf. Rechnet sich vollständig selbst."),
        ("Kategorien", "Kreuztabelle Kategorie × Monat für ein Jahr, mit Vergleich gegen dein "
                       "Budget. Das Jahr stellst du oben in B3 ein."),
        ("Sparziele", "Ziele mit Zielbetrag und Monatsrate, dazu Fortschritt und "
                      "voraussichtliches Zieldatum."),
        ("Vermögen", "Einmal im Monat die Kontostände eintragen. Daraus entstehen "
                     "Nettovermögen, Veränderung und die Konten-Übersicht ganz oben."),
        ("Prognose", "Hochrechnung deines Vermögens für 12 Monate auf Basis deiner "
                     "tatsächlichen Sparrate."),
    ]
    for i, (name, text) in enumerate(blaetter):
        _absatz(ws, z + 1 + i, name, text)

    # ---------------- Legenden ----------------
    z += len(blaetter) + 2
    _zelle(ws, f"A{z}", "Farblegende", F_ABSCHNITT)
    legende = [
        ("Blaue Schrift", "Fest eingetragener Wert – hier darfst du tippen.", F_EINGABE),
        ("Gelb hinterlegt", "Zellen, die du ausfüllen bzw. anpassen solltest.", F_NORMAL),
        ("Schwarze Schrift", "Berechnete Formel – bitte nicht überschreiben.", F_NORMAL),
        ("Grüne Schrift", "Verweis auf ein anderes Blatt – ebenfalls Formel.", F_NORMAL),
    ]
    for i, (label, text, schrift) in enumerate(legende):
        zeile = z + 1 + i
        _zelle(ws, f"A{zeile}", label, schrift)
        if label == "Gelb hinterlegt":
            ws[f"A{zeile}"].fill = FILL_GELB
        ws.merge_cells(f"B{zeile}:E{zeile}")
        _zelle(ws, f"B{zeile}", text, F_NORMAL, align=OBEN)

    z += len(legende) + 2
    _zelle(ws, f"A{z}", "Wichtige Hinweise", F_ABSCHNITT)
    hinweise = [
        ("Beispielzeilen", "In Buchungen, Fixkosten und Sparziele steht je eine graue "
                           "BEISPIEL-Zeile als Formatvorlage. Einfach überschreiben oder leeren."),
        ("Zeitraum", "Monatsübersicht, Vermögen und Prognose laufen von Januar 2026 bis "
                     "Dezember 2027. Das Blatt Buchungen fasst 1.000 Zeilen."),
        ("Sparen ist keine Ausgabe", "Überweisungen auf dein Tagesgeld sind keine Ausgabe – "
                                     "sonst zählt die Sparquote doppelt. Solche Umbuchungen "
                                     "am besten aus den Buchungen löschen und nur im Blatt "
                                     "Vermögen abbilden."),
        ("Eigene Diagramme", "Das Import-Skript baut die Diagramme auf dem Dashboard bei "
                             "jedem Lauf neu. Eigene Grafiken deshalb lieber auf einem "
                             "separaten Blatt anlegen, sonst sind sie nach dem nächsten "
                             "Import weg."),
    ]
    for i, (label, text) in enumerate(hinweise):
        _absatz(ws, z + 1 + i, label, text, hoehe=42)

    ws.sheet_properties.tabColor = DUNKELBLAU
    return ws


def vermoegen_anpassen(wb):
    """Spaltenköpfe auf die echten Konten umbenennen und den Stand eintragen."""
    ve = wb["Vermögen"]
    koepfe = {2: "Girokonto S-POOL (€)", 3: "Tagesgeld (€)", 4: "Depot (€)",
              5: "Bargeld (€)", 6: "Gemeinschaftskonto (€)", 8: "Kreditkarte / Schulden (€)"}
    for spalte, titel in koepfe.items():
        ve.cell(row=3, column=spalte).value = titel

    schon_belegt = any(ve.cell(row=STAND_ZEILE, column=c).value not in (None, "")
                       for c in (2, 3, 4, 5, 6, 8))
    if schon_belegt:
        print(f"Blatt Vermögen: Zeile {STAND_ZEILE} ist bereits gefüllt – nicht überschrieben.")
        return

    for *_, spalte, stand, _name in KONTEN:
        z = ve.cell(row=STAND_ZEILE, column=spalte)
        z.value = stand
        z.font = F_EINGABE
        z.fill = FILL_GELB
        z.number_format = EURO
    ve.cell(row=STAND_ZEILE, column=12).value = (
        f"Kontostände laut Sparkassen-Finanzübersicht vom {STAND_TEXT} (von Sarah bereitgestellt)")
    ve.cell(row=STAND_ZEILE, column=12).font = F_KLEIN
    print(f"Blatt Vermögen: Stand {STAND_TEXT} in Zeile {STAND_ZEILE} eingetragen.")


def dashboard_konten(wb):
    """Datenblock mit den aktuellen Kontoständen – Grundlage des Kontodiagramms."""
    da = wb["Dashboard"]
    _zelle(da, "A26", "Kontostände (letzter Eintrag im Blatt Vermögen)", F_ABSCHNITT)
    _zelle(da, "A27", "Konto", F_KOPF, FILL_KOPF)
    _zelle(da, "B27", "Saldo (€)", F_KOPF, FILL_KOPF)
    for i, (name, _inhaber, _nummer, spalte, _stand, _kontoname) in enumerate(KONTEN):
        zeile = 28 + i
        buchstabe = chr(ord("A") + spalte - 1)
        vorzeichen = "-" if spalte == 8 else ""
        _zelle(da, f"A{zeile}", name, F_NORMAL)
        _zelle(da, f"B{zeile}",
               f"={vorzeichen}IFERROR(LOOKUP(9.99999999999999E+307,"
               f"Vermögen!${buchstabe}$4:${buchstabe}$27),0)", F_NORMAL, format=EURO)
    zeile = 28 + len(KONTEN)
    _zelle(da, f"A{zeile}", "Summe", F_FETT)
    _zelle(da, f"B{zeile}", f"=SUM(B28:B{zeile - 1})", F_FETT, format=EURO)


def main():
    parser = argparse.ArgumentParser(description="Finanzplaner übersichtlicher machen")
    parser.add_argument("--excel", default="Finanzplaner_Sarah.xlsx", help="Pfad zur Excel-Datei")
    args = parser.parse_args()

    pfad = Path(args.excel)
    if not pfad.exists():
        raise SystemExit(f"Datei nicht gefunden: {pfad}")

    wb = load_workbook(pfad)
    anleitung_neu(wb)
    vermoegen_anpassen(wb)
    dashboard_konten(wb)
    diagramme_neu_aufbauen(wb)
    wb.save(pfad)

    print(f"\nFertig – {pfad} überarbeitet.")
    print("Excel berechnet beim Öffnen automatisch neu.")


if __name__ == "__main__":
    main()
