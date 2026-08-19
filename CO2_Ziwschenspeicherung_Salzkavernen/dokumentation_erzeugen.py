# -*- coding: utf-8 -*-
"""
Erzeugt je Projektordner eine Word-Dokumentation der enthaltenen Skripte
========================================================================
Struktur (Konstanten, Funktionen, Ausgabedateien) wird aus dem Quelltext
gelesen, die Erklaerung stammt aus der kuratierten Tabelle unten. So bleibt
die Dokumentation automatisch aktuell, wo sie es sein kann, und lesbar, wo
sie es sein muss.

Aufruf:  python dokumentation_erzeugen.py
Ergebnis: in jedem Ordner eine Datei  Dokumentation_<Ordner>.docx
"""
import ast
import io
import os
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor, Inches

BASIS = os.path.dirname(os.path.abspath(__file__))

NAVY = RGBColor(0x00, 0x30, 0x4F)
BLAU = RGBColor(0x04, 0x76, 0xD9)
GRAU = RGBColor(0x4A, 0x4A, 0x4A)

# =====================================================================
# Kuratierte Erklaerungen. Schluessel = Dateiname.
#   zweck    kurze Antwort auf "wofuer ist das da"
#   methode  Liste von Zeilen: die tragenden Gleichungen und Schritte
#   ergebnis wichtigste Zahlen, die herauskommen
#   annahmen Liste der Setzungen, die das Ergebnis tragen
#   hinweis  optionale Warnung
# =====================================================================
DOKU = {
    "co2_kaverne3.py": dict(
        zweck="Kernmodul der Kavernenauslegung. Rechnet Gassäule, Teufenprofile, "
              "Druckfenster, Arbeitsgas und die Zustandspunkte für Ein- und "
              "Ausspeicherpfad. Alle anderen Module beziehen sich auf die hier "
              "erzeugten Zahlen.",
        methode=[
            "Gassäule abwärts:  dp/dz = ρ(p,T)·g, schrittweise integriert (600 Schritte),",
            "                   T(z) folgt dem Gebirgsprofil T = 10 °C + 0,03 K/m · z.",
            "Kopfdruck:         nicht vorgegeben, sondern per Sekantenverfahren so gelöst,",
            "                   dass unten an der LCCS genau p_max erreicht wird.",
            "Ausspeichern:      adiabater Grenzfall h + g·z = konstant; Dichte und",
            "                   Temperatur aus (p, h), weil dieses Paar auch im",
            "                   Zweiphasengebiet eindeutig ist.",
            "Pumpe/Verdichter:  w = [h₂ₛ(p₂,s₁) − h₁] / η_s",
            "Drossel:           isenthalp, h₂ = h₁",
            "Arbeitsgas:        m = ρ(p,T) · V, Differenz zwischen p_max und p_min",
        ],
        ergebnis=[
            "Kopfdruck Einspeichern 107,6 bar, Säulengewinn +102,4 bar auf 210 bar",
            "Pumpenarbeit 3,21 kJ/kg, Leistung 88–176 kW",
            "Bohrlochkopf Ausspeichern: aus der Säule erreichbar 116,3 bar / 33,8 °C",
            "Kavernentemperatur 46 °C, kritische Teufe 699 m",
        ],
        annahmen=[
            "Oberflächentemperatur 10 °C und Gradient 0,03 K/m (Annahme der Quelle Q-028)",
            "η_s = 0,80 für Pumpe und Verdichter",
            "statische Säule, keine Reibungsdruckverluste",
            "p_kopf_aus = 120 bar ist gesetzt — die Säule liefert nur 116,3 bar",
        ],
        hinweis="Das Modul gibt selbst eine Warnung aus, wenn der gesetzte "
                "Ausspeicher-Sollwert über dem liegt, was die Gassäule hergibt. "
                "Diese Warnung ist derzeit aktiv.",
    ),
    "zyklus.py": dict(
        zweck="Vereinfachte Kavernensimulation über einen kompletten Speicherzyklus. "
              "Beantwortet, wie stark die Ausspeicherrate das nutzbare Arbeitsgas "
              "bestimmt.",
        methode=[
            "Massenbilanz:   dm/dt = ṁ",
            "Energiebilanz:  dU/dt = ṁ·h + UA·(T_Gebirge − T_Kaverne)",
            "Zustand aus (ρ, u) — dieses Paar ist auch zweiphasig eindeutig.",
            "Ratenbegrenzung auf dpdt_max = 6 bar/Tag.",
            "Phasenfolge: Ausspeichern → Stillstand → Einspeichern → Stillstand → Ausspeichern",
        ],
        ergebnis=[
            "55 kg/s: Kaverne kühlt auf 33,8 °C, bleibt einphasig → 71 % Arbeitsgas",
            "234 kg/s: Kaverne kühlt auf 28,7 °C, unter T_krit → nur 33 %",
            "Zyklusdauer 278 Tage, Inventar 139,6 bis 485,7 kt, Kissengas 29 %",
        ],
        annahmen=[
            "Kaverne ideal durchmischt",
            "UA = 546 kW/K — geht direkt in jedes Ergebnis ein",
            "Bohrloch nicht modelliert",
            "Einspeichertemperatur 30 °C an der LCCS",
        ],
    ),
    "zyklus_folie.py": dict(
        zweck="Folientaugliche Abbildung zur Zyklussimulation. Gleiche Rechnung wie "
              "zyklus.py, aber im Querformat und mit der Aussage im Vordergrund.",
        methode=["Ruft zyklus.simuliere() für zwei Raten auf und stellt "
                 "Kaverneninnentemperatur und Inventar gegenüber."],
        ergebnis=["abb_zyklus_folie.png, 16,6 × 6,9 Zoll, für die volle Folienbreite"],
        annahmen=["identisch mit zyklus.py"],
    ),
    "teufenoptimierung.py": dict(
        zweck="Leitet das zulässige Teufenfenster ab und untersucht, welche "
              "Ausspeicherrate je Teufe zulässig ist.",
        methode=[
            "Schranke 1:  T_Gebirge > T_krit + Marge  →  z > 799 m",
            "Schranke 2:  p_min < p_krit              →  z < 1139 m",
            "Schranke 3:  Geomechanik                 →  z > 700 m",
            "Grenzrate je Teufe: größte Rate, bei der die Kaverne einphasig bleibt.",
        ],
        ergebnis=[
            "zulässiges Teufenfenster 799 – 1139 m",
            "zwei Optima: 799 m mit 84 % Arbeitsgas und 144 kt/a,",
            "             1139 m mit 67 % Arbeitsgas und 644 kt/a",
        ],
        annahmen=[
            "p_max = 80 %, p_min = 30 % des lithostatischen Drucks (Faustwerte)",
            "Sicherheitsmarge 3 K zur kritischen Temperatur (eigene Wahl)",
            "UA = 546 kW/K geht direkt in die Grenzrate ein",
        ],
        hinweis="Der Auslegungsdurchsatz von 54,9 kg/s liegt über der höchsten "
                "zulässigen Rate im Fenster (53 kg/s bei 1139 m).",
    ),
    "rechenweg_teufenfenster.py": dict(
        zweck="Folienabbildung: macht die Herleitung des Teufenfensters nachvollziehbar "
              "und trennt sichtbar, welche Eingabe aus der Quelle stammt und welche "
              "eigene Setzung ist.",
        methode=[
            "Untergrenze:  z > (T_krit + 3 K − 10 °C) / 0,03 = 799 m",
            "Obergrenze:   Gradient ρ_Salz·g = 0,2158 bar/m, davon 30 % = 0,0647 bar/m,",
            "              z < p_krit / 0,0647 = 1139 m",
        ],
        ergebnis=["abb_rechenweg_teufenfenster.png"],
        annahmen=["siehe teufenoptimierung.py"],
    ),
    "teufenfenster_herleitung.py": dict(
        zweck="Folienabbildung: Vorgehen, Validierung gegen Q-028 und der Zielkonflikt "
              "zwischen Arbeitsgasanteil und Jahresdurchsatz.",
        methode=["Fenstergrenzen live gerechnet; die ratenabhängigen Werte stammen "
                 "aus teufenoptimierung.py (report_teufe.txt)."],
        ergebnis=["abb_teufenfenster_herleitung.png"],
        annahmen=["siehe teufenoptimierung.py"],
    ),
    "vorwaermer.py": dict(
        zweck="Vorwärmleistung vor der Druckreduzierung, im Vergleich CO₂ gegen Erdgas.",
        methode=[
            "Die Drossel ist isenthalp, der gesamte Enthalpiehub kommt vom Vorwärmer:",
            "    Q = ṁ · [ h(p₂, T_ziel) − h(p₁, T_ein) ]",
            "Die Temperatur zwischen Vorwärmer und Drossel kommt darin nicht vor;",
            "sie ergibt sich hinterher durch Umkehrung von h bei p₁.",
        ],
        ergebnis=[
            "Drosselung 210 → 30 bar, Ziel 17,5 °C: 167,0 kJ/kg = 9 168 kW",
            "ohne Vorwärmung käme der Strom bei −5,6 °C zweiphasig aus der Drossel",
            "CO₂ braucht das 7,5- bis 11-fache von Erdgas",
        ],
        annahmen=[
            "Eintritt 210 bar / 50 °C — bewusst nicht an die 1200-m-Kaverne gekoppelt;",
            "dort läge die Kopftemperatur bei rund 34 °C",
            "einstufige Drosselung, also obere Schranke",
        ],
    ),
    "rechenweg_pumpe.py": dict(
        zweck="Folienabbildung: wie aus dem Anlieferzustand die 3,21 kJ/kg und daraus "
              "die Motorgröße werden.",
        methode=[
            "1. verlustfreie Maschine erzeugt keine Entropie → s₂ = s₁",
            "2. damit ist h₂ₛ(p₂, s₁) bestimmt",
            "3. w_s = h₂ₛ − h₁ = 2,572 kJ/kg",
            "4. w = w_s / η_s = 3,214 kJ/kg",
            "Gegenprobe: w ≈ Δp/(ρ·η) = 3,23 kJ/kg, 0,5 % Abweichung",
        ],
        ergebnis=[
            "Leistungskette 142 → 177 → 186 kW, Normmotor 200 kW",
            "Auslegungspunkt für die Anfrage: 113–226 m³/h bei H = 263 m",
        ],
        annahmen=["η_s = 0,80, η_Motor = 0,955, Antriebsmarge nach API 610"],
    ),
    "pumpenwahl_beleg.py": dict(
        zweck="Folienabbildung: warum am Auslegungspunkt kein Verdichter in Frage kommt, "
              "belegt mit eigener Rechnung und Herstellerangaben.",
        methode=[
            "Realgasfaktor Z = p/(ρ·R·T) an allen Prozesspunkten gegen das",
            "Herstellerkriterium Z > 0,65 (Siemens Energy).",
            "Gegenrechnung: was es kostet, den Strom turbotauglich zu machen.",
        ],
        ergebnis=[
            "Z am Auslegungspunkt 0,179 — deutlich unter der Grenze",
            "Z = 0,65 wird bei 85 bar erst bei 62 °C erreicht",
            "Aufheizen dorthin kostet 12,3 MW, die Verdichtung wäre danach 3,9-mal teurer",
        ],
        annahmen=["Z-Grenze 0,65 ist ein Software-Gültigkeitsbereich, keine "
                  "physikalische Grenze"],
    ),
    "rechenweg_verdichter.py": dict(
        zweck="Folienabbildung: wie die Stufenzahl zustande kommt und warum der "
              "Verdichter bei 80 bar an die Pumpe übergibt.",
        methode=[
            "Stufenzahl: n = ⌈ ln(p₂/p₁) / ln(r_max) ⌉ mit r_max = 2,2 → n = 2",
            "je Stufe: w = [h₂ₛ(p₂,s₁) − h₁] / η_s, dazwischen Zwischenkühlung",
            "Vergleich der letzten 28 bar: 3. Verdichterstufe gegen Verflüssigen + Pumpe",
        ],
        ergebnis=[
            "2 Stufen: 57,22 kJ/kg, höchste Austrittstemperatur 84 °C",
            "3. Verdichterstufe 11,2 kJ/kg gegen Pumpe 4,4 kJ/kg — Faktor 2,6",
        ],
        annahmen=[
            "r_max = 2,2 ist eine Auslegungskonvention, keine physikalische Grenze",
            "Zwischenkühlung 40 °C, Verflüssigung 25 °C",
        ],
        hinweis="Die frühere Begründung „einstufig über 180 °C\" ist falsch. "
                "30 → 80 bar einstufig ergibt 102 °C. Die 183 °C gelten für 30 → 210 bar.",
    ),
    "verdichterwahl_beleg.py": dict(
        zweck="Folienabbildung: Betriebspunkt des Verdichters, Bauartvergleich und "
              "Kandidaten.",
        methode=["Betriebspunkt aus Massenstrom und Dichte am Stufeneintritt; "
                 "Bauartvergleich aus Q-047, Figure 5."],
        ergebnis=[
            "Eintrittsvolumenstrom Stufe 1: 1444–2888 m³/h — 13-mal so groß wie bei der Pumpe",
            "Z = 0,805, also klar im Turbogebiet",
            "Wellenleistung 1,57–3,14 MW, Kühllast 7,0–14,1 MW",
        ],
        annahmen=["Der Herstellervergleich gilt für 1 Mt/a von 2 auf 150 bar — "
                  "relative Aussagen übertragbar, absolute nicht"],
    ),
    "gaspfad_vergleich.py": dict(
        zweck="Vergleicht in Szenario 2 die beiden Wege auf 107,6 bar: verflüssigen "
              "und pumpen gegen gasförmig durchverdichten.",
        methode=["Wellenarbeit und Kühllast beider Pfade als gestapelte Bilanz."],
        ergebnis=[
            "Pfad 2: 61,6 kJ/kg Welle, 256,5 kJ/kg Kühllast",
            "Pfad 3: 69,7 kJ/kg Welle, 263,1 kJ/kg Kühllast",
            "Pfad 2 spart 11,6 % Wellenarbeit bei sogar geringerer Kühllast",
        ],
        annahmen=["η_s = 0,80, Zwischenkühlung 40 °C, Verflüssigung 80 bar / 25 °C"],
    ),
    "pumpe_oder_verdichter.py": dict(
        zweck="Zeigt, ab welchem Anlieferdruck eine Pumpe überhaupt anwendbar ist.",
        methode=["Siedelinie, Dichtesprung am Siededruck und spezifische Wellenarbeit "
                 "über dem Anlieferdruck."],
        ergebnis=["Siededruck bei 15 °C: 50,9 bar; Dichtesprung 161 → 821 kg/m³"],
        annahmen=["η_s = 0,80, maximales Druckverhältnis je Stufe 2,2"],
    ),
    "pfadentscheidung.py": dict(
        zweck="Stellt die drei gerechneten Prozesspfade nebeneinander und ordnet ihnen "
              "laufende CCS-Projekte zu.",
        methode=["Darstellung, keine eigene Rechnung — Zahlen aus gaspfad_vergleich.py "
                 "und pumpe_oder_verdichter.py."],
        ergebnis=["abb_pfadentscheidung.png"],
        annahmen=["Herstellerangaben Stand 08/2026"],
    ),
    "CO2_Prozesspfade.py": dict(
        zweck="Zeichnet die Prozesspfade beider Anlieferszenarien in das p-T-Phasen"
              "diagramm. Punktbezeichnungen identisch mit den Blockfließbildern.",
        methode=[
            "Sättigungslinie und kritischer Punkt aus CoolProp.",
            "Bohrlochsäule live integriert (dp/dz = ρ·g), nicht mehr aus der alten Excel.",
            "Zwei Grenzfälle als Band: Fluid folgt dem Gebirge (langsam) gegen",
            "Fluid behält die Pumpenaustrittstemperatur (schnell).",
        ],
        ergebnis=["abbC_prozesspfad_szenarioD.png, abbD_prozesspfad_szenarioG.png"],
        annahmen=["Kopfdruck 107,6 bar, Gebirge 10 °C + 0,03 K/m, Teufe 1200 m"],
    ),
    "blockfliessbild.py": dict(
        zweck="Blockfließbilder der Übertageanlage für beide Anlieferszenarien.",
        methode=["Darstellung; die Zustandspunkte stammen aus dem Auslegungsmodul."],
        ergebnis=["bfb_sz1.png / bfb_sz2.png bzw. bfb_szD.png / bfb_szG.png"],
        annahmen=["Zustandspunkte wie im Auslegungsmodul"],
    ),
}

ALT = {
    "co2_kaverne.py": "Frühere Fassung des Auslegungsmoduls. Durch co2_kaverne3.py ersetzt.",
    "co2_kaverne2.py": "Frühere Fassung des Auslegungsmoduls. Durch co2_kaverne3.py ersetzt.",
}


# =====================================================================
# Auslesen aus dem Quelltext
# =====================================================================
def analysiere(pfad):
    """Liest Docstring, Modulkonstanten, Funktionen und Ausgabedateien."""
    quelle = io.open(pfad, encoding="utf-8", errors="replace").read()
    try:
        baum = ast.parse(quelle)
    except SyntaxError:
        return dict(docstring="", konstanten=[], funktionen=[], ausgaben=[])

    doc = ast.get_docstring(baum) or ""

    # Modulkonstanten mit dem Kommentar dahinter
    zeilen = quelle.splitlines()
    konstanten = []
    for k in baum.body:
        if not isinstance(k, ast.Assign):
            continue
        ziele = [t.id for t in k.targets if isinstance(t, ast.Name)]
        gross = [n for n in ziele if n.isupper() or re.fullmatch(r"[a-z_]+", n)]
        if not gross:
            continue
        try:
            wert = ast.literal_eval(k.value)
        except Exception:
            continue
        if isinstance(wert, (list, dict, tuple)) and len(str(wert)) > 60:
            continue
        zeile = zeilen[k.lineno - 1] if k.lineno <= len(zeilen) else ""
        kommentar = zeile.split("#", 1)[1].strip() if "#" in zeile else ""
        konstanten.append((", ".join(ziele), wert, kommentar))

    funktionen = []
    for k in baum.body:
        if isinstance(k, ast.FunctionDef) and not k.name.startswith("_"):
            d = (ast.get_docstring(k) or "").strip().splitlines()
            args = ", ".join(a.arg for a in k.args.args)
            funktionen.append((f"{k.name}({args})", d[0] if d else ""))

    ausgaben = sorted(set(
        re.findall(r'savefig\(\s*[fr]?["\']([^"\']+)', quelle)
        + re.findall(r'savefig\(\s*f?["\']?\{[^}]+\}/([^"\']+)', quelle)
        + re.findall(r'open\(\s*["\']([^"\']+\.(?:txt|csv|json|xlsx))', quelle)
    ))
    return dict(docstring=doc, konstanten=konstanten,
                funktionen=funktionen, ausgaben=ausgaben)


# =====================================================================
# Word-Ausgabe
# =====================================================================
def absatz(doc, text, groesse=10, fett=False, farbe=None, kursiv=False,
           abstand_vor=0, abstand_nach=4, mono=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(abstand_vor)
    p.paragraph_format.space_after = Pt(abstand_nach)
    r = p.add_run(text)
    r.font.size = Pt(groesse)
    r.bold = fett
    r.italic = kursiv
    r.font.name = "Consolas" if mono else "Calibri"
    if farbe is not None:
        r.font.color.rgb = farbe
    return p


def schreibe_ordner(ordner, dateien):
    name = os.path.basename(ordner) or "Projekt"
    doc = Document()
    for s in doc.sections:
        s.left_margin = s.right_margin = Inches(0.9)

    t = doc.add_paragraph()
    r = t.add_run(f"Dokumentation — {name}")
    r.font.size = Pt(20); r.bold = True; r.font.color.rgb = NAVY

    absatz(doc, "CO₂-Zwischenspeicherung in Salzkavernen · Masterarbeit Sarah Maxara",
           10, farbe=GRAU)
    absatz(doc, "Automatisch erzeugt aus dem Quelltext (Konstanten, Funktionen, "
                "Ausgabedateien) und ergänzt um die inhaltliche Erklärung. "
                "Bei Änderungen am Code neu erzeugen mit dokumentation_erzeugen.py.",
           9, farbe=GRAU, kursiv=True, abstand_nach=14)

    for datei in dateien:
        pfad = os.path.join(ordner, datei)
        a = analysiere(pfad)
        d = DOKU.get(datei)

        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(16)
        r = h.add_run(datei)
        r.font.size = Pt(14); r.bold = True; r.font.color.rgb = BLAU

        if datei in ALT:
            absatz(doc, "⚠  " + ALT[datei], 10, fett=True, farbe=RGBColor(0xCA, 0x22, 0x0E))

        # --- Zweck
        absatz(doc, "Zweck", 11, fett=True, farbe=NAVY, abstand_vor=8, abstand_nach=2)
        if d and d.get("zweck"):
            absatz(doc, d["zweck"], 10)
        elif a["docstring"]:
            erste = [z for z in a["docstring"].splitlines()
                     if z.strip() and not set(z.strip()) <= set("=-")]
            absatz(doc, " ".join(erste[:3]), 10)
        else:
            absatz(doc, "— keine Beschreibung hinterlegt —", 10, kursiv=True, farbe=GRAU)

        # --- Methode
        if d and d.get("methode"):
            absatz(doc, "Methode und tragende Gleichungen", 11, fett=True,
                   farbe=NAVY, abstand_vor=8, abstand_nach=2)
            for z in d["methode"]:
                absatz(doc, z, 9.5, mono=True, abstand_nach=1)

        # --- Eingaben
        if a["konstanten"]:
            absatz(doc, "Eingaben und Randbedingungen", 11, fett=True,
                   farbe=NAVY, abstand_vor=10, abstand_nach=3)
            tab = doc.add_table(rows=1, cols=3)
            tab.style = "Light Grid Accent 1"
            for i, kopf in enumerate(("Größe", "Wert", "Bedeutung")):
                z = tab.rows[0].cells[i].paragraphs[0].add_run(kopf)
                z.bold = True; z.font.size = Pt(9)
            for nm, wert, kom in a["konstanten"][:26]:
                zeile = tab.add_row().cells
                for i, txt in enumerate((nm, str(wert), kom)):
                    p = zeile[i].paragraphs[0]
                    rr = p.add_run(txt)
                    rr.font.size = Pt(8.5)
                    if i < 2:
                        rr.font.name = "Consolas"

        # --- Funktionen
        if a["funktionen"]:
            absatz(doc, "Funktionen", 11, fett=True, farbe=NAVY,
                   abstand_vor=10, abstand_nach=3)
            for sig, beschr in a["funktionen"]:
                p = doc.add_paragraph()
                p.paragraph_format.space_after = Pt(2)
                r1 = p.add_run(sig)
                r1.font.name = "Consolas"; r1.font.size = Pt(9); r1.bold = True
                if beschr:
                    r2 = p.add_run("   " + beschr)
                    r2.font.size = Pt(9); r2.font.color.rgb = GRAU

        # --- Ergebnis
        if d and d.get("ergebnis"):
            absatz(doc, "Was herauskommt", 11, fett=True, farbe=NAVY,
                   abstand_vor=10, abstand_nach=2)
            for z in d["ergebnis"]:
                absatz(doc, "•  " + z, 10, abstand_nach=1)

        if a["ausgaben"]:
            absatz(doc, "Erzeugte Dateien: " + " · ".join(a["ausgaben"]),
                   9, farbe=GRAU, kursiv=True, abstand_vor=4)

        # --- Annahmen
        if d and d.get("annahmen"):
            absatz(doc, "Annahmen, die das Ergebnis tragen", 11, fett=True,
                   farbe=NAVY, abstand_vor=10, abstand_nach=2)
            for z in d["annahmen"]:
                absatz(doc, "•  " + z, 10, abstand_nach=1)

        if d and d.get("hinweis"):
            absatz(doc, "⚠  " + d["hinweis"], 10, fett=True,
                   farbe=RGBColor(0xEE, 0x72, 0x03), abstand_vor=8)

        absatz(doc, f"Ausführen:  python {datei}", 9, mono=True,
               farbe=GRAU, abstand_vor=8)

    ziel = os.path.join(ordner, f"Dokumentation_{name}.docx")
    doc.save(ziel)
    return ziel


def main():
    gesamt = 0
    for wurzel, _, dateien in os.walk(BASIS):
        if ".git" in wurzel or "__pycache__" in wurzel:
            continue
        py = sorted(f for f in dateien
                    if f.endswith(".py") and f != "dokumentation_erzeugen.py")
        if not py:
            continue
        ziel = schreibe_ordner(wurzel, py)
        print(f"  {len(py):>2} Skripte  ->  {os.path.relpath(ziel, BASIS)}")
        gesamt += len(py)
    print(f"\n{gesamt} Skripte dokumentiert.")


if __name__ == "__main__":
    main()
