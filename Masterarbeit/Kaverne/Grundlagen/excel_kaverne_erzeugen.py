# -*- coding: utf-8 -*-
"""
Excel-Rechenübersicht der Kaverne erzeugen (reines CO₂ und Worst-Case-Gemisch)
============================================================================
Baut aus den Ergebnisdateien (JSON) der Kaverne-Skripte je Stoff eine Excel-Mappe im
Stil der Einspeicherpfad-Mappen (Farblegende, Annahmen mit Quelle und Status, Blätter je
Rechenschritt, Diagramme, Quellen):

  Reines_CO2/CO2_Kaverne_Rechenuebersicht.xlsx                   (aus Reines_CO2/*.json)
  Begleitstoffe_CO2/CO2_Kaverne_Rechenuebersicht_Gemisch.xlsx    (aus Begleitstoffe_CO2/*.json)

Graue Werte kommen aus dem Python-Code und gelten nur für die Annahmen in der Spalte
"gerechnet mit" der Übersicht. Wird dort ein blauer Wert geändert, färben sich alle grauen
Werte rot -> Skripte 01-04 (und 06) neu rechnen und diese Datei neu ausführen. Weiße
Werte sind Excel-Formeln und rechnen sofort mit.

Reihenfolge: 01-04 im Stoffordner (Gemisch zusätzlich 05), Ausspeichergas/06, dann
    python excel_kaverne_erzeugen.py            (beide Mappen)
    python excel_kaverne_erzeugen.py rein|gem   (nur eine)
Fehlt eine JSON-Datei, bleibt das zugehörige Blatt mit einem Hinweis leer.
"""
import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference, ScatterChart, Series
from openpyxl.chart.axis import ChartLines
from openpyxl.chart.text import RichText, Text
from openpyxl.chart.title import Title
from openpyxl.drawing.text import CharacterProperties, Paragraph, ParagraphProperties, RegularTextRun
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.styles.differential import DifferentialStyle
from openpyxl.formatting.rule import Rule
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.hyperlink import Hyperlink
from openpyxl.worksheet.page import PageMargins
from openpyxl.worksheet.properties import PageSetupProperties

GRUNDLAGEN = Path(__file__).resolve().parent
KAVERNE = GRUNDLAGEN.parent
ORDNER = {"rein": KAVERNE / "Reines_CO2", "gem": KAVERNE / "Begleitstoffe_CO2"}
DATEI = {"rein": "CO2_Kaverne_Rechenuebersicht.xlsx", "gem": "CO2_Kaverne_Rechenuebersicht_Gemisch.xlsx"}
NAME = {"rein": "reines CO₂", "gem": "Worst-Case-Gemisch"}
RATE_A, RATE_B = "100.000 Nm³/h", "50.000 Nm³/h"
RATEN = (RATE_B, RATE_A)
ST_A, ST_B = "Stillstand 21,67 d (Q-028)", "kontinuierlich (0 d)"
BLAU, ROT, GRUEN, ORANGE, GRAU = "00304F", "CA220E", "007335", "EE7203", "7F7F7F"
STAND = "15.10.2026"

# ------------------------------------------------------------------ Hausstil
DUENN = Side(style="thin", color="BFBFBF")
RAND = Border(left=DUENN, right=DUENN, top=DUENN, bottom=DUENN)
KRAND = Side(style="thin", color="EE7203")
STIL = {
    "titel": dict(font=Font(name="Calibri", size=20, bold=True, color=BLAU), al=Alignment(vertical="center")),
    "untertitel": dict(font=Font(name="Calibri", size=10, italic=True, color="595959"), al=Alignment(vertical="center")),
    "link": dict(font=Font(name="Calibri", size=10, bold=True, color="0476D9"), al=Alignment(vertical="center")),
    "h2": dict(font=Font(name="Calibri", size=11, bold=True, color=BLAU), al=Alignment(vertical="center")),
    "abschnitt": dict(font=Font(name="Calibri", size=12, bold=True, color="FFFFFF"), fill="00304F",
                      al=Alignment(vertical="center")),
    "kasten": dict(font=Font(name="Calibri", size=10), fill="FFF7EE", al=Alignment(vertical="top", wrap_text=True),
                   rand=Border(left=KRAND, right=KRAND, top=KRAND, bottom=KRAND)),
    "kopf": dict(font=Font(name="Calibri", size=10, bold=True, color="FFFFFF"), fill="1F4E79",
                 al=Alignment(horizontal="center", vertical="center", wrap_text=True), rand=RAND),
    "band": dict(font=Font(name="Calibri", size=10, bold=True, color=BLAU), fill="DDE7F0",
                 al=Alignment(vertical="center"), rand=RAND),
    "text": dict(font=Font(name="Calibri", size=10), al=Alignment(vertical="center", wrap_text=True), rand=RAND),
    "textfett": dict(font=Font(name="Calibri", size=10, bold=True), al=Alignment(vertical="center", wrap_text=True),
                     rand=RAND),
    "mitte": dict(font=Font(name="Calibri", size=10), al=Alignment(horizontal="center", vertical="center",
                                                                    wrap_text=True), rand=RAND),
    "eingabe": dict(font=Font(name="Calibri", size=10, bold=True, color="0000FF"), fill="FFF2CC",
                    al=Alignment(horizontal="center", vertical="center"), rand=RAND),
    "grau": dict(font=Font(name="Calibri", size=10, color="404040"), fill="E7E6E6",
                 al=Alignment(horizontal="center", vertical="center"), rand=RAND),
    "formel": dict(font=Font(name="Calibri", size=10, color="000000"), fill="FFFFFF",
                   al=Alignment(horizontal="center", vertical="center"), rand=RAND),
    "ergebnis": dict(font=Font(name="Calibri", size=10, bold=True, color="000000"), fill="E2EFDA",
                     al=Alignment(horizontal="center", vertical="center"), rand=RAND),
    "notiz": dict(font=Font(name="Calibri", size=9, italic=True, color="595959"), al=Alignment(vertical="center")),
    "warn": dict(font=Font(name="Calibri", size=10, bold=True, color="9C0006"), fill="FFE5E5",
                 al=Alignment(vertical="center", wrap_text=True),
                 rand=Border(left=Side(style="thin", color="C00000"), right=Side(style="thin", color="C00000"),
                             top=Side(style="thin", color="C00000"), bottom=Side(style="thin", color="C00000"))),
}
DXF_VERALTET = DifferentialStyle(font=Font(color="9C0006"), fill=PatternFill(bgColor="F8CBAD"))
DXF_WARN = DifferentialStyle(font=Font(color="9C0006", bold=True), fill=PatternFill(bgColor="FFE5E5"))
DXF_OK = DifferentialStyle(font=Font(color="007335", bold=True))


def lade(stoff, datei):
    for pfad in (ORDNER[stoff] / datei, KAVERNE / "Ausspeichergas" / datei):
        if pfad.exists():
            with open(pfad, encoding="utf-8") as f:
                return json.load(f)
    return None


def de(x, nd=1):
    if x is None:
        return "–"
    s = f"{x:,.{nd}f}"
    return s.replace(",", "X").replace(".", ",").replace("X", ".")


def z(ws, ref, wert=None, art="text", nf=None, fett=None, ausr=None):
    c = ws[ref]
    s = STIL[art]
    c.font = s["font"]
    if "fill" in s:
        c.fill = PatternFill("solid", fgColor=s["fill"])
    if "al" in s:
        c.alignment = s["al"]
    if "rand" in s:
        c.border = s["rand"]
    if wert is not None:
        c.value = wert
        if isinstance(wert, str) and wert.startswith("= "):
            c.data_type = "s"
    if nf:
        c.number_format = nf
    if fett is not None:
        f = c.font.copy(bold=fett)
        c.font = f
    if ausr:
        c.alignment = Alignment(horizontal=ausr, vertical="center", wrap_text=True)
    return c


def reich(text):
    """Kastentext mit fetten Stichwörtern ("STICHWORT:  Text") als Rich Text."""
    import re
    from openpyxl.cell.rich_text import CellRichText, TextBlock
    from openpyxl.cell.text import InlineFont
    fett, normal = InlineFont(rFont="Calibri", sz=10, b=True), InlineFont(rFont="Calibri", sz=10)
    teile = []
    zeilen = text.split("\n")
    for i, zl in enumerate(zeilen):
        ende = "\n" if i < len(zeilen) - 1 else ""
        m = re.match(r"^(.{2,95}?[:?])  (.*)$", zl, re.S)
        if m:
            lab = m.group(1)
            bu = [ch for ch in lab if ch.isalpha()]
            if bu and sum(ch.isupper() for ch in bu) / len(bu) >= 0.6:
                teile += [TextBlock(fett, lab), TextBlock(normal, "  " + m.group(2) + ende)]
                continue
        teile.append(TextBlock(normal, zl + ende))
    return CellRichText(teile)


def kasten(ws, rng, text, hoehe, art="kasten"):
    from openpyxl.utils.cell import range_boundaries
    c1, r1, c2, r2 = range_boundaries(rng)
    for r in range(r1, r2 + 1):
        for c in range(c1, c2 + 1):
            z(ws, f"{get_column_letter(c)}{r}", None, art)
    ws.merge_cells(rng)
    ws.cell(r1, c1).value = reich(text) if art == "kasten" else text
    ws.row_dimensions[r1].height = hoehe


def link(ws, ref, text, blatt):
    c = z(ws, ref, text, "link")
    c.hyperlink = Hyperlink(ref=ref, location=f"'{blatt}'!A1", display=text)


def kopfzeile(ws, zeile, spalten, texte, hoehe=30):
    for s, t in zip(spalten, texte):
        z(ws, f"{s}{zeile}", t, "kopf")
    ws.row_dimensions[zeile].height = hoehe


def band(ws, zeile, von, bis, text):
    for c in range(von, bis + 1):
        z(ws, f"{get_column_letter(c)}{zeile}", None, "band")
    z(ws, f"{get_column_letter(von)}{zeile}", text, "band")
    ws.merge_cells(start_row=zeile, start_column=von, end_row=zeile, end_column=bis)


def veraltet(ws, rng):
    ws.conditional_formatting.add(rng, Rule(type="expression", formula=["NOT(Stand_ok)"], dxf=DXF_VERALTET))


def name(wb, n, ziel):
    wb.defined_names[n] = DefinedName(n, attr_text=ziel)


def ref(blatt, zelle):
    s, r = "".join(ch for ch in zelle if ch.isalpha()), "".join(ch for ch in zelle if ch.isdigit())
    return f"'{blatt}'!${s}${r}"


def blatt(wb, titel, untertitel, tab, breiten):
    ws = wb.create_sheet(titel[:31] if len(titel) <= 31 else titel[:31])
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = tab
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = 9
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.page_margins = PageMargins(left=0.4, right=0.4, top=0.5, bottom=0.5, header=0.3, footer=0.3)
    ws.column_dimensions["A"].width = 2.7
    for s, b in breiten.items():
        ws.column_dimensions[s].width = b
    return ws


def kopf_blatt(ws, titel, untertitel):
    z(ws, "B1", titel, "titel")
    ws.row_dimensions[1].height = 30
    z(ws, "B2", untertitel, "untertitel")
    link(ws, "B3", "← zur Übersicht", "Übersicht")


def diagrammtitel(ch, text, sz=1100):
    cp = CharacterProperties(sz=sz, b=True, solidFill=BLAU)
    para = Paragraph(pPr=ParagraphProperties(defRPr=cp), r=[RegularTextRun(rPr=cp, t=text)])
    ch.title = Title(tx=Text(rich=RichText(p=[para])), overlay=False)


def streu(titel, xt, yt, b=16, h=9, xmin=None, xmax=None, ymin=None, ymax=None, y_umkehren=False):
    ch = ScatterChart()
    diagrammtitel(ch, titel)
    ch.style = 13
    ch.x_axis.title, ch.y_axis.title = xt, yt
    ch.width, ch.height = b, h
    ch.x_axis.delete = ch.y_axis.delete = False
    for ax, lo, hi in ((ch.x_axis, xmin, xmax), (ch.y_axis, ymin, ymax)):
        if lo is not None:
            ax.scaling.min = lo
        if hi is not None:
            ax.scaling.max = hi
    if y_umkehren:
        ch.y_axis.scaling.orientation = "maxMin"
    ch.legend.position = "b"
    return ch


def linie(ws, ch, xs, ys, r1, r2, titel, farbe, breite=2.0, strich=None, marker=None):
    s = Series(Reference(ws, min_col=ys, min_row=r1, max_row=r2), Reference(ws, min_col=xs, min_row=r1, max_row=r2),
               title=titel)
    s.graphicalProperties.line.solidFill = farbe
    s.graphicalProperties.line.width = int(breite * 12700)
    if strich:
        s.graphicalProperties.line.dashStyle = strich
    if marker:
        s.marker.symbol, s.marker.size = marker, 6
        s.marker.graphicalProperties.solidFill = farbe
        s.marker.graphicalProperties.line.solidFill = farbe
    else:
        s.marker.symbol = "none"
    s.smooth = False
    ch.series.append(s)
    return s


def fehlt(ws, zeile, datei, skript):
    kasten(ws, f"B{zeile}:L{zeile}", f"NOCH NICHT GERECHNET:  {datei} fehlt – zuerst {skript} ausführen, dann diese Datei "
                                     f"neu erzeugen (excel_kaverne_erzeugen.py).", 32, art="warn")


# ======================================================================================
# Übersicht
# ======================================================================================
ANNAHMEN = [
    "Referenzkaverne (nur tief – die flache Kaverne entfällt, Q-030 S. 17)",
    ("Teufe LCCS", "z_lccs", 1200, "m", "Q-030 Tab. 1, Option 2 (Uniper-Bestand); abgestimmt mit M. Wall", "0"),
    ("Bezugsteufe des Kaverneninhalts unter der LCCS", "dz_bezug", 205, "m",
     "Q-028 Tab. 1: LCCS 1150 m, Kavernentemperatur in 1355 m angegeben → gleiche Form", "0"),
    ("Geometrisches Volumen", "V_kav", 650000, "m³",
     "Q-028 Tab. 1 („accessible for storage“); bestätigt in der Besprechung K. Vosbeck 15.10.2026", "#,##0"),
    ("Maximaldruck an der LCCS", "p_max", 210, "bar", "Vorgabe M. Wall; zwischen Q-028 Tab. 1 (200) und Q-030 Tab. 1 (240)", "0"),
    ("Minimaldruck an der LCCS", "p_min", 70, "bar", "Q-028 Tab. 1: geomechanisches Minimum", "0"),
    "Gebirge",
    ("Gebirgstemperatur an der Oberfläche", "T0", 10, "°C", "Q-028 Kap. 3 (S. 5), Fig. 3", "0.0"),
    ("Geothermischer Gradient", "gradT", 0.03, "K/m", "Q-028 Kap. 3 (S. 5), Fig. 3", "0.000"),
    ("Wärmeübergang Kaverne–Gebirge (konstant)", "UA", 546, "kW/K",
     "an der Simulation aus Q-028 (S. 10) kalibriert – Nachrechnung Blatt „03 Zyklus“", "0"),
    "Betrieb",
    ("Normvolumenstrom, unterer Fall", "Vn_min", 50000, "Nm³/h", "wie Einspeicherung (abgestimmt mit M. Wall)", "#,##0"),
    ("Normvolumenstrom, oberer Fall (Auslegung)", "Vn_max", 100000, "Nm³/h", "wie Einspeicherung", "#,##0"),
    ("Zulässige Druckänderung an der LCCS", "dpdt_max", 10, "bar/d", "Q-030 Kap. 1.2.1; Q-089 Tab. 2", "0"),
    ("Zulässige Druckänderung, vorsichtig (Vergleich)", "dpdt_vors", 6, "bar/d", "Q-028 Tab. 1 (Simulation)", "0"),
    ("Stillstand nach Aus- und nach Einspeichern", "t_still", 21.67, "d",
     "Q-028 S. 10 (Kap. 4): Betriebsszenario „Standstill phase (21.67 d)“ – gleiche Begründung wie dort; "
     "Vergleich kontinuierlich (0 d)", "0.00"),
    "Förderstrang und Übergabe",
    ("Innendurchmesser Förderstrang (8 5/8\" 36# J55)", "D_strang", 0.198755, "m", "Q-030 Tab. 1 und Kap. 1.2.1 (S. 12)",
     "0.0000"),
    ("Rohrreibungszahl (Darcy)", "f_darcy", 0.016, "–", "Q-003 Kap. 6.2: 0,015–0,017 für neues Stahlrohr", "0.000"),
    ("Strömungsgeschwindigkeit höchstens", "v_max", 20, "m/s", "Q-030 Kap. 1.2.1", "0"),
    ("Strömung im Strang", "–", "adiabat", "–",
     "Besprechung K. Vosbeck 15.10.2026: schnelle Förderung, kein Wärmeaustausch mit dem Gebirge (h + g·z = konst.)", None),
    ("Rückgabe ins Netz (dicht)", "p_rueck", 91, "bar", "OGE-Austausch 16.09.2026 (wie Übergabe S1)", "0"),
    ("Gasförmige Übergabe / Drosselung auf", "p_gas", 30, "bar", "wie Übergabe S2", "0"),
    ("Kopftemperatur mindestens (Vergleich Erdgas)", "T_kopf_min", 5, "°C", "Q-035 Folie 36 (Speicher Epe, Erdgas)", "0"),
    "Einspeisung (gekoppelt an die Einspeicherung)",
    ("Kopftemperatur beim Einspeisen", "–", "Pumpenaustritt", "–",
     "S1: 91 bar / 15 °C + Pumpe; S2: Kühler 91 bar / 26 °C + Pumpe (Einspeicherung/Begleitstoffe_CO2/Grundlagen); "
     "bei kleinerem Kopfdruck isenthalp gedrosselt (Q-030 S. 15)", None),
]


def uebersicht(wb, stoff, K):
    ws = wb.active
    ws.title = "Übersicht"
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = BLAU
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    for s, b in dict(A=2.7, B=12, C=12, D=14, E=11, F=12, G=9, H=12, I=12, J=12, K=12, L=12, M=12, N=12, O=12).items():
        ws.column_dimensions[s].width = b
    z(ws, "B1", f"CO₂-Zwischenspeicherung in Salzkavernen – Kaverne · {NAME[stoff]}", "titel")
    ws.row_dimensions[1].height = 30
    z(ws, "B2", "Rechenübersicht Kaverne · Masterarbeit Sarah Maxara · TH Köln / Uniper Energy Storage", "untertitel")
    z(ws, "B3", f"Stand: {STAND} (nach der Besprechung mit K. Vosbeck) · Stoffdaten: CoolProp 8.0.0 ("
                + ("Span & Wagner 1996" if stoff == "rein" else "GERG-2008-Mischungsregel, 95 % CO₂ + 5 % Begleitstoffe")
                + f") · Abbild der Skripte Kaverne/{ORDNER[stoff].name}/01–04"
                + (", 05" if stoff == "gem" else "") + " und Kaverne/Ausspeichergas/06", "untertitel")
    z(ws, "B5", "Worum geht es?", "h2")
    kasten(ws, "B6:H11",
           "Diese Mappe zeigt, was in der Kaverne zwischen Einspeicherung und Ausspeicherung passiert – so weit es für die "
           "Obertageanlage wichtig ist:\n"
           "01  Welcher Kopfdruck im Stillstand und – neu – im Betrieb (= Zieldruck der Pumpe)?\n"
           "02  Wie viel CO₂ lässt sich umschlagen (Arbeitsgas), welche Teufe ist günstig?\n"
           "03  Welche Rate verträgt die Kaverne, wie verläuft ein Zyklus – mit Stillstand 21,67 d (Q-028) oder kontinuierlich?\n"
           "04  Welcher Kopfdruck und welche Temperatur beim Ein- und Ausspeichern (Kopfdruckfenster für die Regelung)?\n"
           "05  Wie viele Kavernen braucht der Bedarf (Q-002), wie oft wird umgeschlagen?\n"
           "06  Mit welcher Zusammensetzung startet der Ausspeicherpfad (Min- und Max-Fall nach Q-089)?", 22)
    for r in range(7, 12):
        ws.row_dimensions[r].height = 22
    z(ws, "J5", "Farblegende der Zellen", "h2")
    for i, (art, wert, txt) in enumerate((("eingabe", 210, "Annahme – darf geändert werden"),
                                          ("grau", 806.2, "Wert aus dem Python-Code – fest"),
                                          ("formel", 100, "Excel-Formel – rechnet sofort mit"),
                                          ("ergebnis", 114.3, "wichtiges Ergebnis"))):
        z(ws, f"J{6 + i}", wert, art, nf="0.0")
        z(ws, f"K{6 + i}", txt, "notiz")
    c = z(ws, "J10", "rot", "text")
    c.fill = PatternFill("solid", fgColor="F8CBAD")
    c.font = Font(name="Calibri", size=10, color="9C0006")
    z(ws, "K10", "fester Wert passt nicht mehr zu den Annahmen → Skripte neu rechnen", "notiz")
    z(ws, "B13", "Inhalt", "h2")
    inhalt = [("01 Stillstand & Zieldruck", "Gassäule für voll, Grenzfall und leer; strömende Säule beim Einspeisen – Zieldruck S1 und S2."),
              ("02 Arbeitsgas & Teufe", "Inhalt, Kissengas und Arbeitsgas (isotherm / isentrop); Kontrolle gegen DBI; Teufenfenster."),
              ("03 Zyklus", "Rate aus 10 bar/d; Kavernenbilanz über einen Zyklus (S1/S2, 2 Raten, mit und ohne Stillstand); Nachrechnung Q-028."),
              ("04 Bohrlochkopf", "Kopfdruckfenster beim Ein- und Ausspeichern, Pumpe oder Drossel, Phase im Strang, Reibung."),
              ("05 Speicherbedarf", "Arbeitsgas je Kaverne gegen den Bedarf aus Q-002; Zyklen pro Jahr (Q-029). Rechnet live mit."),
              ("06 Ausspeichergas", "Zusammensetzung nach der Kaverne (Min-/Max-Fall), Wasser, Rückgabegrenzen OGE.")]
    if stoff == "gem":
        inhalt.append(("Rein vs. Gemisch", "Alle Kernzahlen mit dem Worst-Case-Gemisch neben reinem CO₂ (Skript 05)."))
    inhalt += [("Quellen", "Herkunft aller Annahmen mit Fundstelle."), ("Diagrammdaten", "Hilfstabellen für die Diagramme (nicht ändern).")]
    for i, (b, t) in enumerate(inhalt):
        link(ws, f"B{14 + i}", "→ " + b, b)
        z(ws, f"E{14 + i}", t, "text").border = Border()
        ws.merge_cells(f"E{14 + i}:O{14 + i}")
    r0 = 14 + len(inhalt) + 1
    for c in range(2, 16):
        z(ws, f"{get_column_letter(c)}{r0}", None, "abschnitt")
    z(ws, f"B{r0}", "Annahmen", "abschnitt")
    ws.merge_cells(f"B{r0}:O{r0}")
    z(ws, f"B{r0 + 1}", "Blaue Werte dürfen geändert werden. Graue Werte kommen aus dem Python-Code (Kavernenbilanz, "
                        "CoolProp) und passen nur zu den Annahmen in „gerechnet mit“ – wird eine Annahme geändert, färben "
                        "sich alle grauen Werte rot: Skripte neu rechnen und diese Mappe neu erzeugen.", "notiz")
    zs = r0 + 2
    z(ws, f"B{zs}", "Status", "textfett")
    ws.merge_cells(f"B{zs}:D{zs}")
    z(ws, f"E{zs}", None, "textfett")
    ws.merge_cells(f"E{zs}:O{zs}")
    kopf = zs + 2
    kopfzeile(ws, kopf, "BEFGHMN", ["Größe", "Symbol", "Wert", "Einheit", "Quelle / Begründung", "gerechnet mit", "Status"],
              20)
    ws.merge_cells(f"B{kopf}:D{kopf}")
    ws.merge_cells(f"H{kopf}:L{kopf}")
    ws.merge_cells(f"N{kopf}:O{kopf}")
    r = kopf + 1
    erste = r
    for e in ANNAHMEN:
        if isinstance(e, str):
            band(ws, r, 2, 15, e)
            r += 1
            continue
        groesse, sym, wert, einheit, quelle, nf = e
        z(ws, f"B{r}", groesse, "text")
        ws.merge_cells(f"B{r}:D{r}")
        z(ws, f"E{r}", sym, "text")
        z(ws, f"G{r}", einheit, "text")
        z(ws, f"H{r}", quelle, "text")
        ws.merge_cells(f"H{r}:L{r}")
        if isinstance(wert, str):
            z(ws, f"F{r}", wert, "grau")
            z(ws, f"M{r}", "–", "grau")
            z(ws, f"N{r}", "Modellannahme", "text")
        else:
            z(ws, f"F{r}", wert, "eingabe", nf=nf)
            z(ws, f"M{r}", wert, "grau", nf=nf)
            z(ws, f"N{r}", f'=IF(ABS(F{r}-M{r})<0.000000001,"✔ passt","⚠ geändert – neu rechnen")', "text")
            name(wb, sym, f"Übersicht!$F${r}")
        ws.merge_cells(f"N{r}:O{r}")
        ws.row_dimensions[r].height = 27 if len(quelle) > 70 else 18
        r += 1
    letzte = r - 1
    ws[f"Q{zs}"] = f'=COUNTIF($N${erste}:$N${letzte},"⚠*")=0'
    ws[f"Q{zs}"].font = Font(color="FFFFFF", size=8)
    name(wb, "Stand_ok", f"Übersicht!$Q${zs}")
    ws[f"E{zs}"] = ('=IF(Stand_ok,"✔ Alle grauen Werte passen zu den Annahmen.","⚠ Mindestens eine Annahme wurde geändert – '
                    'rot markierte Werte sind veraltet: Skripte neu rechnen.")')
    ws.conditional_formatting.add(f"E{zs}:O{zs}", Rule(type="expression", formula=["NOT(Stand_ok)"], dxf=DXF_WARN))
    ws.conditional_formatting.add(f"E{zs}:O{zs}", Rule(type="expression", formula=["Stand_ok"], dxf=DXF_OK))
    ws.conditional_formatting.add(f"N{erste}:O{letzte}", Rule(type="expression", formula=[f'LEFT(N{erste},1)="⚠"'],
                                                                dxf=DXF_WARN))
    ws.conditional_formatting.add(f"N{erste}:O{letzte}", Rule(type="expression", formula=[f'LEFT(N{erste},1)="✔"'],
                                                                dxf=DXF_OK))
    # Kernergebnisse (Formeln auf die Ergebnisblätter, siehe K)
    r = letzte + 2
    for c in range(2, 16):
        z(ws, f"{get_column_letter(c)}{r}", None, "abschnitt")
    z(ws, f"B{r}", "Kernergebnisse auf einen Blick", "abschnitt")
    ws.merge_cells(f"B{r}:O{r}")
    r += 1
    for txt, formel, einheit, nf, bem in K:
        z(ws, f"B{r}", txt, "text")
        ws.merge_cells(f"B{r}:E{r}")
        z(ws, f"F{r}", formel, "ergebnis", nf=nf)
        z(ws, f"G{r}", einheit, "text")
        z(ws, f"H{r}", bem, "text")
        ws.merge_cells(f"H{r}:O{r}")
        ws.row_dimensions[r].height = 27 if len(bem) > 80 or len(txt) > 45 else 18
        r += 1
    return ws


# ======================================================================================
# 01 Stillstand & Zieldruck
# ======================================================================================
def blatt_01(wb, stoff, D, DD):
    ws = blatt(wb, "01 Stillstand & Zieldruck", "", "1F4E79", dict(B=30, C=12, D=12, E=12, F=12, G=12, H=12, I=13, J=26,
                                                                   K=12, L=12))
    kopf_blatt(ws, "01 Stillstand & Zieldruck – welcher Druck am Bohrlochkopf?",
               f"Referenzkaverne (LCCS 1200 m) · {NAME[stoff]} · Code: Kaverne/{ORDNER[stoff].name}/01_stillstand_und_zieldruck.py "
               "(Bausteine kaverne_bausteine.py)")
    if D is None:
        fehlt(ws, 5, "kaverne_stillstand.json", "01_stillstand_und_zieldruck.py")
        return {}
    S, B = D["saeulen"], D["betrieb"]
    grenze = "Cricondenbar 82,2 bar" if stoff == "gem" else "p_krit 73,8 bar"
    kasten(ws, "B4:L4",
           "STILLSTAND:  Ruhendes CO₂ hat Gebirgstemperatur T(z) = T₀ + grad T·z. Die Säule wird in 2-m-Schritten von unten "
           "(bekannter Druck an der LCCS) nach oben gerechnet: dp/dz = ρ(p,T)·g. Drei Füllstände wie Q-028 Fig. 3: voll, "
           f"Grenzfall (Kopf = {grenze}: darüber ganzes Bohrloch einphasig) und leer.\n"
           "BETRIEB (NEU):  Beim Einspeisen strömt das CO₂ mit der Temperatur nach der Pumpe nach unten – adiabat "
           "(h + g·z = konst., kein Wärmeaustausch mit dem Gebirge; Besprechung K. Vosbeck 15.10.2026), mit Reibung im "
           "8 5/8\"-Strang (Darcy f = 0,016, Q-003 Kap. 6.2). Pumpenaustritt und Kopfdruck hängen voneinander ab → Iteration. "
           "S1 (15 °C) gibt eine kältere, schwerere Säule als im Stillstand, S2 (26 °C + Pumpe) eine wärmere, leichtere.\n"
           f"ERGEBNIS:  Zieldruck der Pumpe bei {RATE_A}: S1 {de(B['S1'][RATE_A]['p_kopf'])} bar, S2 "
           f"{de(B['S2'][RATE_A]['p_kopf'])} bar (Stillstand voll {de(S['voll']['p_kopf'])} bar).\n"
           "GRENZE:  homogen und einphasig gerechnet; ein zweiphasiger Abschnitt wird nur angezeigt – genauer nur mit "
           "Mehrphasen-Bohrlochmodell (Q-028 S. 11).", 120)
    z(ws, "B6", "Stillstand: drei Füllstände der Kaverne (Gebirgstemperatur)", "h2")
    kopfzeile(ws, 7, "BCDEFGHIJ", ["Fall", "Druck LCCS [bar]", "Kopfdruck [bar]", "T Kopf [°C]", "ρ Kopf [kg/m³]",
                                   "ρ LCCS [kg/m³]", "Säulengewinn [bar]", "zweiphasig im Bohrloch [m]", "Bedeutung"], 42)
    bed = {"voll": "Kopfdruck im Stillstand bei voller Kaverne", "Grenzfall": f"Kopf = {grenze}",
           "leer": "niedrigster Kopfdruck im Stillstand"}
    out = {}
    for i, f in enumerate(("voll", "Grenzfall", "leer")):
        r = 8 + i
        s = S[f]
        z(ws, f"B{r}", f, "textfett")
        z(ws, f"C{r}", s["p_lccs"], "grau", nf="0.0")
        z(ws, f"D{r}", s["p_kopf"], "ergebnis" if f == "voll" else "grau", nf="0.0")
        z(ws, f"E{r}", 10.0, "grau", nf="0.0")
        z(ws, f"F{r}", s["rho_kopf"], "grau", nf="0")
        z(ws, f"G{r}", s["rho_lccs"], "grau", nf="0")
        z(ws, f"H{r}", f"=C{r}-D{r}", "formel", nf="0.0")
        zb = s.get("zweiphasig_bereich")
        z(ws, f"I{r}", f"{de(zb[0], 0)}–{de(zb[1], 0)}" if zb else "–", "mitte")
        z(ws, f"J{r}", bed[f], "text")
        out[f] = ref(ws.title, f"D{r}")
    veraltet(ws, "C8:G10")
    name(wb, "Kopf_voll", out["voll"])
    name(wb, "Kopf_leer", out["leer"])
    name(wb, "p_grenzfall", ref(ws.title, "C9"))

    z(ws, "B12", f"Betrieb: Einspeisen mit strömender Säule (Pumpenaustritt am Kopf, 210 bar an der LCCS)", "h2")
    kopfzeile(ws, 13, "BCDEFGHIJ", ["Fall", "ṁ [kg/s]", "Kopfdruck [bar]", "T Kopf [°C]", "T LCCS [°C]", "v Kopf [m/s]",
                                    "davon Reibung [bar]", "Säulengewinn [bar]", "Bedeutung"], 42)
    r = 14
    for szen in ("S1", "S2"):
        for rate in RATEN:
            b = B[szen][rate]
            z(ws, f"B{r}", f"Einspeisen {szen}, {rate}", "textfett")
            z(ws, f"C{r}", b["m_dot"], "grau", nf="0.00")
            z(ws, f"D{r}", b["p_kopf"], "ergebnis" if rate == RATE_A else "grau", nf="0.0")
            z(ws, f"E{r}", b["T_kopf"], "grau", nf="0.0")
            z(ws, f"F{r}", b["T_lccs"], "grau", nf="0.0")
            z(ws, f"G{r}", b["v_kopf"], "grau", nf="0.00")
            rb = b["p_kopf"] - b["p_kopf_ohne_reibung"] if "p_kopf_ohne_reibung" in b else None
            z(ws, f"H{r}", rb if rb is not None else "–", "grau", nf="0.00")
            z(ws, f"I{r}", f"=210-D{r}", "formel", nf="0.0")
            z(ws, f"J{r}", f"Zieldruck Pumpe {szen} (Auslegung)" if rate == RATE_A else "", "text")
            if rate == RATE_A:
                name(wb, f"Kopf_betr_{szen}", ref(ws.title, f"D{r}"))
                out[f"betr_{szen}"] = ref(ws.title, f"D{r}")
            r += 1
    veraltet(ws, "C14:H17")
    z(ws, "B18", "Reibung = Kopfdruck mit minus ohne Rohrreibung (nur bei 100.000 Nm³/h gerechnet; bei 50.000 Nm³/h etwa ein "
                 "Viertel, da ~ṁ²). Säulengewinn = 210 bar − Kopfdruck (enthält die Reibung).", "notiz")

    # Profile -> Diagrammdaten
    st = DD.cell
    c0 = DD.max_column + 2 if DD.max_column > 1 else 2
    titel = ["z [m]", "T Gebirge [°C]", "p voll [bar]", "p Grenzfall [bar]", "p leer [bar]", "p S1 Betrieb [bar]",
             "T S1 Betrieb [°C]", "p S2 Betrieb [bar]", "T S2 Betrieb [°C]"]
    DD.cell(5, c0, "01 Profile (alle 50 m)").font = Font(bold=True, color=BLAU)
    for j, t in enumerate(titel):
        DD.cell(6, c0 + j, t).font = Font(bold=True)
    pv, pg, pl = S["voll"]["profil_50m"], S["Grenzfall"]["profil_50m"], S["leer"]["profil_50m"]
    b1, b2 = B["S1"][RATE_A]["profil_50m"], B["S2"][RATE_A]["profil_50m"]
    n = len(pv)
    for i in range(n):
        werte = [pv[i][0], pv[i][2], pv[i][1], pg[i][1], pl[i][1], b1[i][1], b1[i][2], b2[i][1], b2[i][2]]
        for j, v in enumerate(werte):
            DD.cell(7 + i, c0 + j, v)
    r1, r2 = 7, 6 + n
    ch = streu("Druck im Bohrloch: Stillstand und Einspeisen (100.000 Nm³/h)", "Druck [bar]", "Teufe [m]", b=15, h=10,
               xmin=0, xmax=220, ymin=0, ymax=1200, y_umkehren=True)
    for col, t, farbe, strich in ((c0 + 2, "Stillstand voll", BLAU, None), (c0 + 3, "Stillstand Grenzfall", GRAU, "dash"),
                                  (c0 + 4, "Stillstand leer", "8EA9C1", "sysDot"), (c0 + 5, "Einspeisen S1", GRUEN, None),
                                  (c0 + 7, "Einspeisen S2", ORANGE, None)):
        linie(DD, ch, col, c0, r1, r2, t, farbe, 2, strich)
    ws.add_chart(ch, "B20")
    ch2 = streu("Temperatur im Bohrloch: Gebirge (Stillstand) und strömende Säule", "Temperatur [°C]", "Teufe [m]", b=15,
                h=10, xmin=0, xmax=60, ymin=0, ymax=1200, y_umkehren=True)
    for col, t, farbe in ((c0 + 1, "Gebirge = Stillstand", BLAU), (c0 + 6, "Einspeisen S1", GRUEN),
                          (c0 + 8, "Einspeisen S2", ORANGE)):
        linie(DD, ch2, col, c0, r1, r2, t, farbe, 2)
    ws.add_chart(ch2, "G20")
    return out


# ======================================================================================
# 02 Arbeitsgas & Teufe
# ======================================================================================
def blatt_02(wb, stoff, A, DD):
    ws = blatt(wb, "02 Arbeitsgas & Teufe", "", "1F4E79", dict(B=30, C=12, D=13, E=13, F=12, G=13, H=13, I=12, J=12, K=12))
    kopf_blatt(ws, "02 Arbeitsgas & Teufe – wie viel CO₂ lässt sich umschlagen?",
               f"{NAME[stoff]} · Code: Kaverne/{ORDNER[stoff].name}/02_arbeitsgas_teufenfenster.py")
    if A is None:
        fehlt(ws, 5, "arbeitsgas.json", "02_arbeitsgas_teufenfenster.py")
        return {}
    kasten(ws, "B4:K4",
           "INHALT (TGV):  Masse in der vollen Kaverne bei 210 bar an der LCCS und Gebirgstemperatur: m = ρ(p_Mitte, T)·V.\n"
           "KISSENGAS (CGV):  was bei p_min in der Kaverne bleibt – isotherm (langsam, das Gebirge hält die Temperatur) oder "
           "isentrop (schnell, das CO₂ kühlt beim Entspannen ab und ist dichter → mehr bleibt drin).\n"
           "ARBEITSGAS (WGV):  TGV − CGV. Realistisch liegt der Zyklus dazwischen – die Kavernenbilanz (Blatt 03) rechnet "
           "den Wärmeübergang mit.\n"
           "TEUFENFENSTER:  ab welcher Teufe bleibt der Kaverneninhalt (Gebirgstemperatur) mit Abstand über der "
           + ("Cricondentherm" if stoff == "gem" else "kritischen Temperatur") + "; bis welche Teufe passt p_max zum Gebirge.",
           92)
    z(ws, "B6", "Inhalt, Kissengas und Arbeitsgas", "h2")
    z(ws, "B7", "Inhalt bei p_max (TGV) [t]", "textfett")
    z(ws, "C7", A["TGV"], "ergebnis", nf="#,##0")
    name(wb, "TGV", ref(ws.title, "C7"))
    kopfzeile(ws, 9, "BCDEFGHIJ", ["p_min-Variante", "p_min LCCS [bar]", "CGV isotherm [t]", "CGV isentrop [t]",
                                   "T isentrop [°C]", "WGV isotherm [t]", "WGV isentrop [t]", "Anteil isotherm",
                                   "Anteil isentrop"], 42)
    for i, v in enumerate(A["varianten"]):
        r = 10 + i
        z(ws, f"B{r}", v["variante"], "textfett")
        z(ws, f"C{r}", v["p_min"], "grau", nf="0.0")
        z(ws, f"D{r}", v["CGV_iso"], "grau", nf="#,##0")
        z(ws, f"E{r}", v["CGV_ise"] if v["CGV_ise"] is not None else "–", "grau", nf="#,##0")
        z(ws, f"F{r}", v["T_ise"] if v["T_ise"] is not None else "–", "grau", nf="0.0")
        z(ws, f"G{r}", f"=TGV-D{r}", "formel", nf="#,##0")
        z(ws, f"H{r}", f"=IF(ISNUMBER(E{r}),TGV-E{r},\"–\")", "formel", nf="#,##0")
        z(ws, f"I{r}", f"=G{r}/TGV", "ergebnis", nf="0%")
        z(ws, f"J{r}", f"=IF(ISNUMBER(H{r}),H{r}/TGV,\"–\")", "formel", nf="0%")
    rv = 9 + len(A["varianten"])
    veraltet(ws, f"C10:F{rv}")
    veraltet(ws, "C7")
    r = rv + 2
    z(ws, f"B{r}", "Kontrolle der Stoffdaten gegen DBI (Q-030 Tab. 1 / Q-029)", "h2")
    kopfzeile(ws, r + 1, "BCDEFG", ["Option", "p [bar]", "T [°C]", "ρ DBI [kg/m³]", "ρ CoolProp [kg/m³]", "Abweichung"], 30)
    for i, k in enumerate(A["dbi_kontrolle"]):
        rr = r + 2 + i
        z(ws, f"B{rr}", k["option"], "text")
        z(ws, f"C{rr}", k["p"], "grau", nf="0")
        z(ws, f"D{rr}", k["T"], "grau", nf="0")
        z(ws, f"E{rr}", k["rho_dbi"], "grau", nf="0.0")
        z(ws, f"F{rr}", k["rho_coolprop"], "grau", nf="0.0")
        z(ws, f"G{rr}", f"=F{rr}/E{rr}-1", "formel", nf="0.0%")
    r = r + 2 + len(A["dbi_kontrolle"]) + 1
    if stoff == "gem":
        z(ws, f"B{r}", "Beim Gemisch weicht ρ von DBI ab – DBI rechnet mit reinem CO₂; die Abweichung ist der Effekt der "
                       "Begleitstoffe.", "notiz")
        r += 1
    T = A["teufenfenster"]
    r += 1
    z(ws, f"B{r}", "Teufenfenster", "h2")
    kopfzeile(ws, r + 1, "BCD", ["Größe", "Wert", "Einheit"], 20)
    for i, (t, v, e) in enumerate((("Grenztemperatur (" + ("Cricondentherm" if stoff == "gem" else "T_krit") + ")",
                                    T["T_grenze"], "°C"), ("Grenzdruck", T["p_grenze"], "bar"),
                                   ("Sicherheitsabstand", T["T_marge"], "K"),
                                   ("Bezugsteufe mindestens", T["z_bezug_min"], "m"),
                                   ("LCCS-Teufe mindestens", T["z_lccs_min"], "m"),
                                   ("LCCS-Teufe höchstens (p_max-Gradient)", T["z_lccs_max"], "m"))):
        rr = r + 2 + i
        z(ws, f"B{rr}", t, "text")
        z(ws, f"C{rr}", v, "grau", nf="0.0")
        z(ws, f"D{rr}", e, "mitte")
    z(ws, f"B{r + 8}", "Die Referenzkaverne (LCCS 1200 m) liegt im Fenster. Flache Kavernen fallen heraus – DBI verwirft die "
                       "flache Option selbst (Q-030 S. 17).", "notiz")
    # Dichtekurve
    c0 = DD.max_column + 2
    DD.cell(5, c0, "02 Dichte über Druck (Gebirgstemperatur Kavernenmitte)").font = Font(bold=True, color=BLAU)
    DD.cell(6, c0, "p [bar]").font = Font(bold=True)
    DD.cell(6, c0 + 1, "ρ [kg/m³]").font = Font(bold=True)
    for i, (p, rho) in enumerate(A["dichtekurve"]):
        DD.cell(7 + i, c0, p)
        DD.cell(7 + i, c0 + 1, rho)
    ch = streu("Dichte des Kaverneninhalts über dem Druck (Gebirgstemperatur)", "Druck Kavernenmitte [bar]",
               "Dichte [kg/m³]", b=16, h=9, xmin=0, xmax=240, ymin=0)
    linie(DD, ch, c0, c0 + 1, 7, 6 + len(A["dichtekurve"]), NAME[stoff], BLAU if stoff == "rein" else ROT, 2.25)
    ws.add_chart(ch, "F" + str(r + 1))
    return {"TGV": ref(ws.title, "C7"), "AG_anteil_iso": ref(ws.title, "I10")}


# ======================================================================================
# 03 Zyklus
# ======================================================================================
def blatt_03(wb, stoff, Z, DD):
    ws = blatt(wb, "03 Zyklus", "", "1F4E79", dict(B=34, C=13, D=13, E=13, F=13, G=13, H=13, I=13, J=13, K=13, L=13, M=13,
                                                   N=13, O=13))
    kopf_blatt(ws, "03 Raten und Kavernenzyklus – mit Stillstand oder kontinuierlich",
               f"{NAME[stoff]} · Code: Kaverne/{ORDNER[stoff].name}/03_kavernenzyklus.py · Kavernenbilanz als ideal "
               "durchmischter Knoten mit Wärmeübergang UA")
    if Z is None:
        fehlt(ws, 5, "kavernenzyklus.json", "03_kavernenzyklus.py")
        return {}
    kasten(ws, "B4:O4",
           "RATE:  Die Kaverne darf an der LCCS höchstens 10 bar/d Druck ändern (Q-030 Kap. 1.2.1). Daraus folgt eine "
           "zulässige Massenrate ṁ = V·(∂ρ/∂p)·dp/dt – sie hängt vom Zustand ab (isotherm: langsam; isentrop: schnell).\n"
           "ZYKLUS:  voll (210 bar, Gebirgstemperatur) → Ausspeichern bis 70 bar → Stillstand → Einspeichern bis 210 bar → "
           "Stillstand. Die Einspeiseenthalpie kommt aus der Einspeicherung (Pumpenaustritt S1 bzw. S2, adiabater Strang). "
           "Die Rate wird gedrosselt, wo 10 bar/d sonst überschritten würden.\n"
           "STILLSTAND 21,67 d:  aus Q-028 S. 10 (Kap. 4) – Betriebsszenario der dortigen Simulation („Standstill phase "
           "(21.67 d)“); gleiche Begründung wie in der Quelle, und an dieser Simulation ist UA kalibriert. Vergleich: "
           "kontinuierlich (0 d).\n"
           "GRENZE:  konstantes UA statt instationärer Wärmeleitung – im Stillstand gleicht sich die Temperatur in Tagen an; "
           "Q-035 beschreibt für Erdgas Monate. Das Ende eines Zyklus hängt deshalb kaum vom Stillstand ab, der Start der "
           "nächsten Phase schon.", 120)
    # Validierung
    v = Z.get("validierung_q028")
    r = 6
    z(ws, f"B{r}", "Nachrechnung Q-028 (reines CO₂, Einspeisung 30 °C an der LCCS) – Kalibrierung von UA", "h2")
    kopfzeile(ws, r + 1, "BCDE", ["Größe", "Q-028", "Kavernenbilanz", "Abweichung"], 20)
    if v:
        for i, (t, k, nf) in enumerate((("Druck LCCS am Ende der Ausspeicherung [bar]", "p_ende", "0.0"),
                                        ("Temperatur am Ende [°C]", "T_ende", "0.0"),
                                        ("Rate min. [t/h]", "rate_min_th", "0"), ("Rate max. [t/h]", "rate_max_th", "0"))):
            rr = r + 2 + i
            z(ws, f"B{rr}", t, "text")
            z(ws, f"C{rr}", v["quelle"][k], "grau", nf=nf)
            z(ws, f"D{rr}", v[k], "grau", nf=nf)
            z(ws, f"E{rr}", f"=D{rr}-C{rr}", "formel", nf=nf)
        z(ws, f"B{r + 6}", "Q-028 S. 10 (Text: 77,6 bar / 33,1 °C am Ende der ersten Ausspeicherung) und Fig. 7/8 (S. 12: Druck, Temperatur, Raten ca. 90–600 t/h); Nachrechnung nur für reines CO₂ – sie kalibriert UA für beide Stoffe.",
          "notiz")
    else:
        z(ws, f"B{r + 2}", "nur in der Mappe für reines CO₂ (UA ist dort kalibriert und gilt für beide Stoffe)", "notiz")
    # Ratengrenze
    rg = Z["ratengrenze"]
    r = 14
    z(ws, f"B{r}", "Zulässige Rate aus der Druckänderung an der LCCS", "h2")
    kopfzeile(ws, r + 1, "BCDEF", ["p LCCS [bar]", "isotherm 6 bar/d [kg/s]", "isotherm 10 bar/d [kg/s]",
                                   "isentrop 6 bar/d [kg/s]", "isentrop 10 bar/d [kg/s]"], 42)
    for i, p in enumerate(rg["p"]):
        rr = r + 2 + i
        z(ws, f"B{rr}", p, "grau", nf="0")
        for j, k in enumerate(("isotherm_6", "isotherm_10", "isentrop_6", "isentrop_10")):
            z(ws, f"{'CDEF'[j]}{rr}", rg[k][i], "grau", nf="0.0")
    rre = r + 1 + len(rg["p"])
    veraltet(ws, f"B{r + 2}:F{rre}")
    gr = Z["grenzrate"]
    z(ws, f"H{r + 1}", "Auslegungsraten [kg/s]", "textfett")
    for i, (k, mm) in enumerate(Z["raten"].items()):
        z(ws, f"H{r + 2 + i}", k, "text")
        z(ws, f"I{r + 2 + i}", mm, "grau", nf="0.0")
    z(ws, f"H{r + 5}", "Grenzrate: Kaverneninhalt bleibt beim Ausspeichern einphasig bis [kg/s]", "text")
    ws.merge_cells(f"H{r + 5}:K{r + 5}")
    z(ws, f"L{r + 5}", gr["rate"], "ergebnis", nf="0.0")
    z(ws, f"H{r + 6}", "bindend (unterhalb der Suchgrenze 150 kg/s)?", "text")
    ws.merge_cells(f"H{r + 6}:K{r + 6}")
    z(ws, f"L{r + 6}", "ja" if gr["bindend"] else "nein", "mitte")
    name(wb, "Grenzrate", ref(ws.title, f"L{r + 5}"))
    ch = streu("Zulässige Rate aus der Druckänderung (10 bar/d bzw. 6 bar/d)", "Druck LCCS [bar]", "Rate [kg/s]", b=16, h=9,
               xmin=60, xmax=220, ymin=0)
    linie(ws, ch, 2, 4, r + 2, rre, "isotherm, 10 bar/d", BLAU, 2)
    linie(ws, ch, 2, 3, r + 2, rre, "isotherm, 6 bar/d", BLAU, 1.5, "dash")
    linie(ws, ch, 2, 6, r + 2, rre, "isentrop, 10 bar/d", ORANGE, 2)
    linie(ws, ch, 2, 5, r + 2, rre, "isentrop, 6 bar/d", ORANGE, 1.5, "dash")
    ws.add_chart(ch, f"H{r + 8}")
    # Zyklen
    r = max(rre + 2, r + 28)
    z(ws, f"B{r}", "Kavernenzyklus: Ergebnisse je Fall", "h2")
    sp = ["Fall", "Arbeitsgas [t]", "Anteil am Inhalt", "Zyklusdauer [d]", "Zyklen pro Jahr", "Aus: Dauer [d]",
          "Aus: T min [°C]", "Aus: Rate gedrosselt [%]", "Stillstand 1: p LCCS danach [bar]", "Ein: Dauer [d]",
          "Ein: T max [°C]", "Stillstand 2: p LCCS danach [bar]", "Kaverne zweiphasig?", "T an der LCCS beim Einspeisen [°C]"]
    kopfzeile(ws, r + 1, "BCDEFGHIJKLMNO", sp, 56)
    out = {}
    rr = r + 2
    for key, zk in Z["zyklen"].items():
        kw = {k["art"] + str(i): k for i, k in enumerate(zk["kennwerte"])}
        aus = zk["kennwerte"][0]
        ein = [k for k in zk["kennwerte"] if k["art"] == "ein"][0]
        still = [k for k in zk["kennwerte"] if k["art"] == "still"]
        z(ws, f"B{rr}", f"{zk['szenario']}, {zk['rate']}, {zk['betriebsweise']}", "textfett")
        z(ws, f"C{rr}", zk["arbeitsgas_t"], "grau", nf="#,##0")
        z(ws, f"D{rr}", zk["anteil_arbeitsgas"], "grau", nf="0%")
        z(ws, f"E{rr}", zk["zyklusdauer_d"], "grau", nf="0")
        z(ws, f"F{rr}", f"=365/E{rr}", "formel", nf="0.0")
        z(ws, f"G{rr}", aus["dauer"], "grau", nf="0")
        z(ws, f"H{rr}", aus["T_min"], "grau", nf="0.0")
        z(ws, f"I{rr}", 100 * aus["anteil_begrenzt"], "grau", nf="0")
        z(ws, f"J{rr}", still[0]["p_lccs_ende"] if still else "–", "grau", nf="0.0")
        z(ws, f"K{rr}", ein["dauer"], "grau", nf="0")
        z(ws, f"L{rr}", ein["T_max"], "grau", nf="0.0")
        z(ws, f"M{rr}", still[1]["p_lccs_ende"] if len(still) > 1 else "–", "grau", nf="0.0")
        zp = any(k["zweiphasig"] for k in zk["kennwerte"])
        z(ws, f"N{rr}", "ja" if zp else "nein", "mitte")
        z(ws, f"O{rr}", zk["T_ein_lccs"], "grau", nf="0.0")
        if key == f"S2|{RATE_A}|{ST_A}":
            for k_, sp_ in (("AG_zyklus", "C"), ("AG_anteil", "D"), ("Zyklusdauer", "E"), ("Zyklen_a", "F")):
                name(wb, k_, ref(ws.title, f"{sp_}{rr}"))
            out["zeile_auslegung"] = rr
        ws.row_dimensions[rr].height = 27
        rr += 1
    veraltet(ws, f"C{r + 2}:E{rr - 1}")
    veraltet(ws, f"G{r + 2}:M{rr - 1}")
    veraltet(ws, f"O{r + 2}:O{rr - 1}")
    z(ws, f"B{rr}", "Arbeitsgas = in einem Zyklus umgeschlagene Masse. Mit Stillstand 21,67 d nach jeder Phase; kontinuierlich = "
                    "ohne Pause. Auslegungsfall (Kernergebnisse, Blatt 05): S2, 100.000 Nm³/h, Stillstand 21,67 d.", "notiz")
    # Verläufe -> Diagrammdaten (jeder n-te Punkt)
    c0 = DD.max_column + 2
    DD.cell(5, c0, f"03 Zyklusverlauf {RATE_A}, S2").font = Font(bold=True, color=BLAU)
    spalten = []
    for bw in (ST_A, ST_B):
        vl = Z["zyklen"][f"S2|{RATE_A}|{bw}"]["verlauf"]
        schritt = max(1, len(vl) // 300)
        pts = vl[::schritt] + [vl[-1]]
        j = len(spalten) * 3
        for k, t in enumerate(("t [d]", f"p LCCS [bar] ({bw})", f"T Kaverne [°C] ({bw})")):
            DD.cell(6, c0 + j + k, t).font = Font(bold=True)
        for i, x in enumerate(pts):
            DD.cell(7 + i, c0 + j, x["t"])
            DD.cell(7 + i, c0 + j + 1, x["p_lccs"])
            DD.cell(7 + i, c0 + j + 2, x["T"])
        spalten.append((c0 + j, 7, 6 + len(pts), bw))
    rz = rr + 2
    ch = streu("Druck an der LCCS über einen Zyklus (S2, 100.000 Nm³/h)", "Zeit [d]", "p LCCS [bar]", b=16, h=9, ymin=50,
               ymax=220)
    ch2 = streu("Temperatur des Kaverneninhalts über einen Zyklus (S2, 100.000 Nm³/h)", "Zeit [d]", "T [°C]", b=16, h=9)
    for (cx, a, b, bw), farbe in zip(spalten, (BLAU, ORANGE)):
        linie(DD, ch, cx, cx + 1, a, b, bw, farbe, 2)
        linie(DD, ch2, cx, cx + 2, a, b, bw, farbe, 2)
    ws.add_chart(ch, f"B{rz}")
    ws.add_chart(ch2, f"G{rz}")
    return out


# ======================================================================================
# 04 Bohrlochkopf
# ======================================================================================
def blatt_04(wb, stoff, BK, DD):
    ws = blatt(wb, "04 Bohrlochkopf", "", "1F4E79", dict(B=36, C=13, D=13, E=13, F=13, G=13, H=13, I=13, J=13, K=13))
    kopf_blatt(ws, "04 Bohrlochkopf – Kopfdruckfenster beim Ein- und Ausspeichern",
               f"{NAME[stoff]} · Code: Kaverne/{ORDNER[stoff].name}/04_bohrlochkopf.py · Strang adiabat, homogen, Reibung "
               "f = 0,016")
    if BK is None:
        fehlt(ws, 5, "bohrlochkopf.json", "04_bohrlochkopf.py")
        return {}
    F = BK["kopfdruckfenster"]
    kasten(ws, "B4:K4",
           "REGELUNG ÜBER DEN KOPFDRUCK (Besprechung K. Vosbeck 15.10.2026, Pkt. 6):  Die Anlage muss über den ganzen Zyklus "
           "den passenden Kopfdruck liefern bzw. aufnehmen. Höchster Wert = Auslegung der Pumpe (Einspeisen in die volle "
           "Kaverne bei 100.000 Nm³/h), kleinster Wert = Rückverdichtung bzw. Drossel beim Ausspeichern.\n"
           "EINSPEISEN:  Reicht der Lieferdruck (91 bar) für den nötigen Kopfdruck, läuft die Pumpe nicht – dann wird am Kopf "
           "isenthalp gedrosselt (kalter Kopf, unten zweiphasiger Strang möglich). Erst ab einem höheren Kavernendruck läuft "
           "die Pumpe.\n"
           "AUSSPEICHERN:  Strömung nach oben, adiabat: das CO₂ entspannt und kühlt im Strang; bei niedrigem Kavernendruck "
           "wird der Strang zweiphasig (Band unten). Danach Drosselung auf 30 bar (gasförmige Übergabe) oder Rückgabe dicht "
           "mit 91 bar (Rückverdichtung nötig, wenn der Kopf darunter liegt).\n"
           "GRENZE:  homogenes Modell – zweiphasige Abschnitte nur angezeigt (Q-028 S. 11).", 120)
    z(ws, "B6", "Kopfdruckfenster (Übersicht)", "h2")
    kopfzeile(ws, 7, "BCDEF", ["Betriebsfall", RATE_B, RATE_A, "Einheit", "Bedeutung"], 30)
    ws.merge_cells("F7:H7")
    zeilen = [("Stillstand, Kaverne voll", F.get("stillstand_voll"), F.get("stillstand_voll"), "bar", "Kopfdruck ruhend"),
              ("Stillstand, Kaverne leer", F.get("stillstand_leer"), F.get("stillstand_leer"), "bar", "Kopfdruck ruhend")]
    for s in ("S1", "S2"):
        zeilen.append((f"Einspeisen {s}: Kopfdruck höchstens", F.get(f"einspeisen_{s}_{RATE_B}_max"),
                       F.get(f"einspeisen_{s}_{RATE_A}_max"), "bar", "Ende der Einspeicherung (Pumpe)"))
        zeilen.append((f"Einspeisen {s}: Kopfdruck mindestens", F.get(f"einspeisen_{s}_{RATE_B}_min"),
                       F.get(f"einspeisen_{s}_{RATE_A}_min"), "bar", "Start in die leere Kaverne (gedrosselt)"))
    zeilen += [("Ausspeisen: Kopfdruck höchstens", F.get(f"ausspeisen_{RATE_B}_max"), F.get(f"ausspeisen_{RATE_A}_max"), "bar",
                "Start (Kaverne voll)"),
               ("Ausspeisen: Kopfdruck mindestens", F.get(f"ausspeisen_{RATE_B}_min"), F.get(f"ausspeisen_{RATE_A}_min"), "bar",
                "Ende (Kaverne leer) → Rückverdichtung auf 91 bar"),
               ("Ausspeisen: Kopftemperatur mindestens", F.get(f"ausspeisen_{RATE_B}_T_min"),
                F.get(f"ausspeisen_{RATE_A}_T_min"), "°C", "Vergleich Epe: 5 °C (Q-035 Folie 36)"),
               ("Ausspeisen: T nach Drossel auf 30 bar mindestens", F.get(f"ausspeisen_{RATE_B}_T_drossel_min"),
                F.get(f"ausspeisen_{RATE_A}_T_drossel_min"), "°C", "→ Vorwärmung vor der Drossel nötig"),
               ("Ausspeisen: Geschwindigkeit am Kopf höchstens", F.get(f"ausspeisen_{RATE_B}_v_max"),
                F.get(f"ausspeisen_{RATE_A}_v_max"), "m/s", "Grenze 20 m/s (Q-030 Kap. 1.2.1)")]
    out = {}
    for i, (t, a, b, e, bed) in enumerate(zeilen):
        r = 8 + i
        z(ws, f"B{r}", t, "textfett")
        z(ws, f"C{r}", a if a is not None else "–", "grau", nf="0.0")
        z(ws, f"D{r}", b if b is not None else "–", "ergebnis" if "höchstens" in t or "mindestens" in t else "grau", nf="0.0")
        z(ws, f"E{r}", e, "mitte")
        z(ws, f"F{r}", bed, "text")
        ws.merge_cells(f"F{r}:H{r}")
        if t.startswith("Ausspeisen: Kopfdruck mindestens"):
            name(wb, "Kopf_aus_min", ref(ws.title, f"D{r}"))
        if t.startswith("Ausspeisen: T nach Drossel"):
            name(wb, "T_drossel_min", ref(ws.title, f"D{r}"))
        if t.startswith("Ausspeisen: Kopftemperatur"):
            name(wb, "T_kopf_aus_min", ref(ws.title, f"D{r}"))
    re_ = 8 + len(zeilen) - 1
    veraltet(ws, f"C8:D{re_}")
    # Schwellen Einspeisen
    r = re_ + 2
    z(ws, f"B{r}", "Einspeisen: ab welchem Kavernendruck einphasig, ab wann Pumpe?", "h2")
    kopfzeile(ws, r + 1, "BCDEF", ["Fall", "Zieldruck Pumpe (Auslegung) [bar]", "Strang einphasig ab p LCCS [bar]",
                                   "Pumpe nötig ab p LCCS [bar]", "darunter"], 42)
    ws.merge_cells(f"F{r + 1}:H{r + 1}")
    rr = r + 2
    for key, e in BK["einspeisen"].items():
        z(ws, f"B{rr}", f"Einspeisen {e['szenario']}, {e['rate']}", "textfett")
        z(ws, f"C{rr}", e["p_kopf_auslegung"], "grau", nf="0.0")
        z(ws, f"D{rr}", e["p_lccs_einphasig_ab"] if e["p_lccs_einphasig_ab"] is not None else "–", "grau", nf="0.0")
        z(ws, f"E{rr}", e["p_lccs_pumpe_ab"] if e["p_lccs_pumpe_ab"] is not None else "–", "grau", nf="0.0")
        z(ws, f"F{rr}", "gedrosselt, Strang unten zweiphasig möglich", "text")
        ws.merge_cells(f"F{rr}:H{rr}")
        rr += 1
    veraltet(ws, f"C{r + 2}:E{rr - 1}")
    # Ausspeisen: zweiphasiges Band
    r = rr + 1
    z(ws, f"B{r}", "Ausspeisen: Kavernendruck, bei dem der Strang zweiphasig ist", "h2")
    kopfzeile(ws, r + 1, "BCDE", ["Rate", "zweiphasig unter p LCCS [bar]", "wieder einphasig unter p LCCS [bar]",
                                  "Vergleich DBI"], 42)
    rr = r + 2
    for rate in RATEN:
        b = BK["ausspeisen"][rate]["strang_zweiphasig"]
        z(ws, f"B{rr}", rate, "textfett")
        z(ws, f"C{rr}", b["p_lccs_oben"] if b else "–", "grau", nf="0.0")
        z(ws, f"D{rr}", (b["p_lccs_unten"] if b and b["p_lccs_unten"] is not None else "bis p_min"), "grau", nf="0.0")
        z(ws, f"E{rr}", "Q-030 S. 17: 160 bar Mindestdruck (Option 2)", "text")
        rr += 1
    veraltet(ws, f"C{r + 2}:D{rr - 1}")
    # Reibung
    r = rr + 1
    z(ws, f"B{r}", "Reibung im Strang (100.000 Nm³/h)", "h2")
    kopfzeile(ws, r + 1, "BCDEFG", ["Fall", "p LCCS [bar]", "Kopfdruck [bar]", "ohne Reibung [bar]", "Δp Reibung [bar]",
                                    "ΔT Kopf [K]"], 30)
    rr = r + 2
    for x in BK["reibung_ausspeisen"]:
        z(ws, f"B{rr}", "Ausspeisen", "text")
        z(ws, f"C{rr}", x["p_lccs"], "grau", nf="0.0")
        z(ws, f"D{rr}", x["p_kopf"], "grau", nf="0.0")
        z(ws, f"E{rr}", x["p_kopf_f0"], "grau", nf="0.0")
        z(ws, f"F{rr}", f"=D{rr}-E{rr}", "formel", nf="0.0")
        z(ws, f"G{rr}", x["dT"], "grau", nf="0.0")
        rr += 1
    for s, x in BK["reibung_einspeisen"].items():
        z(ws, f"B{rr}", f"Einspeisen {s} (Kaverne voll)", "text")
        z(ws, f"C{rr}", 210, "grau", nf="0.0")
        z(ws, f"D{rr}", x["p_kopf"], "grau", nf="0.0")
        z(ws, f"E{rr}", x["p_kopf_f0"], "grau", nf="0.0")
        z(ws, f"F{rr}", f"=D{rr}-E{rr}", "formel", nf="0.0")
        z(ws, f"G{rr}", "–", "mitte")
        rr += 1
    veraltet(ws, f"C{r + 2}:E{rr - 1}")
    z(ws, f"B{rr}", "Beim Ausspeichern kostet die Reibung Kopfdruck und – über den Joule-Thomson-Effekt – Temperatur, am "
                    "stärksten am Ende (niedriger Druck, hohe Geschwindigkeit).", "notiz")
    # Diagramm: Kopfdruck über Kavernendruck
    c0 = DD.max_column + 2
    DD.cell(5, c0, "04 Kopfdruck über Kavernendruck (100.000 Nm³/h)").font = Font(bold=True, color=BLAU)
    reihen = [("Ausspeisen", BK["ausspeisen"][RATE_A]["punkte"])]
    for s in ("S1", "S2"):
        k = f"{s}|{RATE_A}"
        if k in BK["einspeisen"]:
            reihen.append((f"Einspeisen {s}", BK["einspeisen"][k]["punkte"]))
    ch = streu("Kopfdruck über dem Kavernendruck (100.000 Nm³/h) – Kopfdruckfenster", "Druck LCCS [bar]", "Kopfdruck [bar]",
               b=17, h=10, xmin=60, xmax=220, ymin=0)
    ch2 = streu("Kopftemperatur über dem Kavernendruck (100.000 Nm³/h)", "Druck LCCS [bar]", "Kopftemperatur [°C]", b=17,
                h=10, xmin=60, xmax=220)
    for j, ((t, pts), farbe) in enumerate(zip(reihen, (BLAU, GRUEN, ORANGE))):
        cx = c0 + 3 * j
        DD.cell(6, cx, "p LCCS").font = Font(bold=True)
        DD.cell(6, cx + 1, f"p Kopf {t}").font = Font(bold=True)
        DD.cell(6, cx + 2, f"T Kopf {t}").font = Font(bold=True)
        pts = sorted([p for p in pts if p["kopf"]], key=lambda p: p["p_lccs"])
        for i, p in enumerate(pts):
            DD.cell(7 + i, cx, p["p_lccs"])
            DD.cell(7 + i, cx + 1, p["kopf"]["p"])
            DD.cell(7 + i, cx + 2, p["kopf"]["T"])
        linie(DD, ch, cx, cx + 1, 7, 6 + len(pts), t, farbe, 2, marker="circle")
        linie(DD, ch2, cx, cx + 2, 7, 6 + len(pts), t, farbe, 2, marker="circle")
    ws.add_chart(ch, f"B{rr + 2}")
    ws.add_chart(ch2, f"F{rr + 2}")
    return out


# ======================================================================================
# 05 Speicherbedarf (live)
# ======================================================================================
def blatt_05(wb, stoff, Z):
    ws = blatt(wb, "05 Speicherbedarf", "", "007335", dict(B=46, C=14, D=14, E=12, F=60))
    kopf_blatt(ws, "05 Speicherbedarf – wie viele Kavernen, wie oft umgeschlagen?",
               f"{NAME[stoff]} · rechnet live mit den Ergebnissen aus Blatt 03 · Bedarf nach Q-002, Zyklen nach Q-029")
    kasten(ws, "B4:F4",
           "FRAGE:  Reicht eine Referenzkaverne für den Zwischenspeicherbedarf? Q-002 (CO₂START, S. 94–95) nennt für die "
           "Szenarien 1–8 einen Bedarf von 0,21 bis 1,21 Mt. Eine Kaverne schlägt je Zyklus das Arbeitsgas aus Blatt 03 um.\n"
           "ZYKLEN:  Q-029 (DBI für Uniper, Management Summary und Kap. 5) rechnet zwei Fälle: multizyklisch mit 4 Zyklen pro "
           "Jahr (Approach 2) und saisonal mit 1 Zyklus pro Jahr (Approach 8). Die Kavernenbilanz hier ergibt bei 100.000 "
           "Nm³/h und Stillstand 21,67 d weniger als 2 Zyklen pro Jahr – 4 Zyklen bräuchten eine höhere Rate oder kürzere "
           "Stillstände.\n"
           "BLAU = darf geändert werden (Bedarf, Zyklen); alles andere rechnet mit.", 92)
    if Z is None:
        fehlt(ws, 6, "kavernenzyklus.json", "03_kavernenzyklus.py")
        return
    kopfzeile(ws, 6, "BCDEF", ["Größe", "unterer Bedarf", "oberer Bedarf", "Einheit", "Quelle / Formel"], 20)
    rows = [("Zwischenspeicherbedarf", 0.21, 1.21, "Mt", "Q-002 S. 94–95 (Szenarien 1–8)", "eingabe", "0.00"),
            ("Arbeitsgas je Kaverne und Zyklus", "=AG_zyklus/1000000", "=AG_zyklus/1000000", "Mt",
             "Blatt 03: S2, 100.000 Nm³/h, Stillstand 21,67 d", "formel", "0.000"),
            ("Kavernen nötig (Bedarf in einem Zyklus)", "=ROUNDUP(C7/C8,0)", "=ROUNDUP(D7/D8,0)", "–",
             "Bedarf / Arbeitsgas, aufgerundet", "ergebnis", "0"),
            ("Zyklen pro Jahr (Kavernenbilanz)", "=Zyklen_a", "=Zyklen_a", "1/a", "365 d / Zyklusdauer (Blatt 03)", "formel",
             "0.0"),
            ("Umschlag je Kaverne und Jahr (Kavernenbilanz)", "=C8*C10", "=D8*D10", "Mt/a", "Arbeitsgas × Zyklen", "formel",
             "0.00"),
            ("Zyklen pro Jahr laut Q-029 (multizyklisch)", 4, 4, "1/a", "Q-029 Management Summary (PDF-S. 3), Kap. 5 "
                                                                        "(PDF-S. 41–43): Approach 2", "eingabe", "0"),
            ("Umschlag je Kaverne und Jahr bei 4 Zyklen", "=C8*C12", "=D8*D12", "Mt/a", "Arbeitsgas × 4", "formel", "0.00"),
            ("Zyklusdauer, die 4 Zyklen erlauben würde", "=365/C12", "=365/D12", "d", "zum Vergleich mit Blatt 03",
             "formel", "0")]
    for i, (t, a, b, e, q, art, nf) in enumerate(rows):
        r = 7 + i
        z(ws, f"B{r}", t, "textfett" if art == "ergebnis" else "text")
        z(ws, f"C{r}", a, art, nf=nf)
        z(ws, f"D{r}", b, art, nf=nf)
        z(ws, f"E{r}", e, "mitte")
        z(ws, f"F{r}", q, "text")
        ws.row_dimensions[r].height = 20
    name(wb, "Kavernen_min", ref(ws.title, "C9"))
    name(wb, "Kavernen_max", ref(ws.title, "D9"))
    z(ws, "B16", "Grobe Abschätzung: Umschlag in einem Zyklus. Q-029 kommt mit ca. zwei Kavernen bis 2035 aus "
                 "(Management Summary) – dort mit kleinerem Volumen (250.000 m³) und anderer Betriebsweise.", "notiz")


# ======================================================================================
# 06 Ausspeichergas
# ======================================================================================
def blatt_06(wb, stoff, AG):
    ws = blatt(wb, "06 Ausspeichergas", "", "CA220E", dict(B=40, C=14, D=14, E=14, F=14, G=40))
    kopf_blatt(ws, "06 Ausspeichergas – mit welcher Zusammensetzung startet der Ausspeicherpfad?",
               "Code: Kaverne/Ausspeichergas/06_zusammensetzung_nach_kaverne.py (Bausteine ausspeichergas.py) · Bandbreite "
               "nach Q-089 (H₂-Studie UGS für DGMK 2024, übertragen auf CO₂)")
    if AG is None:
        fehlt(ws, 5, "zusammensetzung_nach_kaverne.json", "Ausspeichergas/06_zusammensetzung_nach_kaverne.py")
        return
    kasten(ws, "B4:G4",
           "BANDBREITE STATT EINES WORST CASE (Besprechung K. Vosbeck 15.10.2026):  Min-Fall mit den unteren, Max-Fall mit "
           "den oberen Werten aus Q-089 – Wasser (Abschn. 6.5, Tab. 19), Resterdgas im Finger (Abschn. 6.3, Tab. 16/17), "
           "Diesel-Blanket (Tab. 10), Verdichteröl (Tab. 20), Mikrobiologie (Abschn. 6.4).\n"
           "RÜCKGABE (OGE 16.09.2026):  H₂O ≤ 30 ppm, H₂S < 0,5 ppm, mindestens 91 bar – die Spalte „entfernen“ zeigt, was die "
           "Aufbereitung leisten muss.", 64)
    kl = "gem" if stoff == "gem" else "rein"
    r = 6
    z(ws, f"B{r}", "Einträge in ppm (mol) – erster Zyklus", "h2")
    kopfzeile(ws, r + 1, "BCDE", ["Eintrag", "Min-Fall", "Max-Fall", "Einheit"], 20)
    det = AG.get("details", {})
    dmin, dmax = det.get(f"{kl}_min"), det.get(f"{kl}_max")
    rr = r + 2
    if dmin and dmax:
        for k in dmin["ppm"]:
            z(ws, f"B{rr}", k.replace("_", " "), "text")
            z(ws, f"C{rr}", dmin["ppm"][k], "grau", nf="#,##0.0#")
            z(ws, f"D{rr}", dmax["ppm"][k], "grau", nf="#,##0.0#")
            z(ws, f"E{rr}", "ppm", "mitte")
            rr += 1
        z(ws, f"B{rr}", "Wasser (bezogen auf Normvolumen)", "text")
        z(ws, f"C{rr}", dmin["g_Nm3"], "grau", nf="0.00")
        z(ws, f"D{rr}", dmax["g_Nm3"], "grau", nf="0.00")
        z(ws, f"E{rr}", "g/Nm³", "mitte")
        rr += 1
        veraltet(ws, f"C{r + 2}:D{rr - 1}")
        rr += 1
        z(ws, f"B{rr}", "Rückgabe ins Netz: was die Aufbereitung entfernen muss", "h2")
        kopfzeile(ws, rr + 1, "BCDE", ["Größe", "Min-Fall", "Max-Fall", "Einheit"], 20)
        rr += 2
        for k in dmin.get("rueckgabe", {}):
            z(ws, f"B{rr}", k.replace("_", " "), "text")
            z(ws, f"C{rr}", dmin["rueckgabe"][k], "grau", nf="#,##0.0")
            z(ws, f"D{rr}", dmax["rueckgabe"][k], "grau", nf="#,##0.0")
            z(ws, f"E{rr}", "ppm" if k.endswith("ppm") else ("kg/h" if k.endswith("kg_h") else ""), "mitte")
            rr += 1
    rr += 1
    z(ws, f"B{rr}", "Zusammensetzung feucht (Molanteil) – Übergabe an den Ausspeicherpfad", "h2")
    faelle = AG.get("faelle", {})
    komp = []
    for k in (f"{kl}_min", f"{kl}_max"):
        if faelle.get(k) and faelle[k].get("feucht"):
            komp += [c for c in faelle[k]["feucht"] if c not in komp]
    kopfzeile(ws, rr + 1, "BCDE", ["Komponente", "Min-Fall", "Max-Fall", "ab 2. Zyklus (Max, trocken)"], 30)
    r0 = rr + 2
    for i, c in enumerate(komp):
        z(ws, f"B{r0 + i}", c, "text")
        for sp, k, feld in (("C", f"{kl}_min", "feucht"), ("D", f"{kl}_max", "feucht"),
                            ("E", f"{kl}_max_ab_2_zyklus", "trocken_coolprop")):
            w = (faelle.get(k) or {}).get(feld) or {}
            z(ws, f"{sp}{r0 + i}", w.get(c, 0.0), "grau", nf="0.000000")
    if komp:
        veraltet(ws, f"C{r0}:E{r0 + len(komp) - 1}")
    z(ws, f"B{r0 + len(komp) + 1}", "Molanteile wie in zusammensetzung_nach_kaverne.json (Übergabe an HYSYS über "
                                    "ausspeichergas.py). Max-Fall ist der Auslegungsfall.", "notiz")


# ======================================================================================
# Rein vs. Gemisch (nur Gemisch-Mappe)
# ======================================================================================
def blatt_vergleich(wb, V):
    ws = blatt(wb, "Rein vs. Gemisch", "", "CA220E", dict(B=46, C=14, D=14, E=14, F=9, G=12, H=44))
    kopf_blatt(ws, "Rein vs. Gemisch – was die Begleitstoffe an der Kaverne ändern",
               "Code: Kaverne/Begleitstoffe_CO2/05_vergleich_rein_gemisch.py · alle Werte aus den beiden Stoffordnern")
    if V is None:
        fehlt(ws, 5, "vergleich_rein_gemisch.json", "Begleitstoffe_CO2/05_vergleich_rein_gemisch.py")
        return
    kasten(ws, "B4:H4",
           "FRAGE (Besprechung K. Vosbeck 15.10.2026, Pkt. 7):  Ist reines CO₂ konservativ? Antwort: nicht pauschal – je "
           "Auslegungsgröße ist mal das Gemisch, mal reines CO₂ ungünstiger. Deshalb werden beide gerechnet und je Größe der "
           "ungünstigere Wert für die Auslegung der Obertageanlage genommen (Spalte „ungünstiger“).", 48)
    kopfzeile(ws, 6, "BCDEFGH", ["Auslegungsgröße", "reines CO₂", "Gemisch", "Gemisch − rein", "Einheit", "ungünstiger",
                                 "betrifft"], 30)
    r = 7
    for g in V.get("groessen", []):
        nf = "#,##0" if g["einheit"] == "t" else ("0%" if g["einheit"] == "–" else "0.0")
        z(ws, f"B{r}", g["groesse"], "text")
        for sp, k in (("C", "rein"), ("D", "gemisch")):
            w = g[k]
            z(ws, f"{sp}{r}", w if isinstance(w, (int, float)) and w == w else "–", "grau", nf=nf)
        z(ws, f"E{r}", f'=IF(AND(ISNUMBER(C{r}),ISNUMBER(D{r})),D{r}-C{r},"–")', "formel", nf=nf)
        z(ws, f"F{r}", g["einheit"], "mitte")
        c = z(ws, f"G{r}", g["ungünstiger_fall"], "mitte")
        c.font = Font(name="Calibri", size=10, bold=True, color=ROT if g["ungünstiger_fall"] == "Gemisch" else BLAU)
        z(ws, f"H{r}", g["betrifft"], "text")
        ws.row_dimensions[r].height = 20
        r += 1
    veraltet(ws, f"C7:D{r - 1}")


# ======================================================================================
# Quellen
# ======================================================================================
QUELLEN = [
    ("Q-002", "DBI; DVGW-EBI am KIT; GWI (2026): CO₂-Management – Phase I: Grundlagenermittlung / CO₂START. DVGW-Forschungsbericht G 202434.",
     "S. 94–95: Zwischenspeicherbedarf 0,21–1,21 Mt (Szenarien 1–8) → Blatt 05"),
    ("Q-003", "IfG; UGS; CTU Prag (2022): Energiespeicherung in Salzkavernen mit superkritischem CO₂ – ESsCO₂. Abschlussbericht, FKZ 03ET6154.",
     "Kap. 6.2: Rohrreibungszahl 0,015–0,017 → f = 0,016; Kap. 9.3.2.4: Restsole ≤ 1 %"),
    ("Q-028", "Buzogany, R.; Kruck, O. (2022): CO₂ – Storage in Salt Caverns: Thermodynamic Analysis and Challenges. SMRI Fall 2022.",
     "Tab. 1: Kavernendaten (Volumen, p_min, Bezugsteufe); Kap. 3 (S. 5), Fig. 3: Gebirge 10 °C + 0,03 K/m, Gassäulen; S. 10 "
     "(Kap. 4): Betriebsszenario mit Stillstand 21,67 d, Kalibrierung UA; S. 11: Mehrphasen-Bohrlochmodell nötig"),
    ("Q-029", "DBI für Uniper (2024): CCUS-P2 – CO₂-Storage for CCU/S, Phase 2 – Detailing Study, Final Report V08.",
     "Management Summary (PDF-S. 3), Kap. 5 (PDF-S. 41–43): 4 Zyklen/a (Approach 2) bzw. 1 Zyklus/a (Approach 8)"),
    ("Q-030", "DBI für Uniper (2023): CCUS-P1 – CO₂-Storage for CCU/S, Phase 1 – Initial Concept Study, Final Report.",
     "Tab. 1: Teufe 1200 m (Option 2), Förderstrang 8 5/8\"; Kap. 1.2.1: 10 bar/d, 20 m/s; S. 15: Einspeisung mit konstanter "
     "Kopftemperatur; S. 17: flache Option verworfen, 160 bar Mindestdruck (Option 2)"),
    ("Q-035", "Uniper Energy Storage (2024): Kavernenspeicher – Technische Kennlinien, Kapazitäten, Restriktionen (Onboarding, intern).",
     "Folie 5: Druckbereich Epe; Folie 31: unterer Druckbereich höchstens 60 Tage je Jahr; Folie 36: Gastemperatur am Kopf ≥ 5 °C"),
    ("Q-089", "UGS GmbH (Frommhold, Schwabe) für DGMK (2024): Konzeptstudie zu Reinheitsanforderungen H₂-Netz.",
     "Tab. 2: 10 bar/d; Abschn. 6.3–6.5, Tab. 10, 12, 16, 19, 20: Einträge in das Speichergas → Blatt 06"),
    ("OGE", "Gesprächsnotiz Austausch OGE / Uniper, 16.09.2026 (Notizen S. Maxara).",
     "Rückgabe: mind. 91 bar, H₂O ≤ 30 ppm, H₂S < 0,5 ppm"),
    ("Vosbeck", "Besprechung mit Katrin Vosbeck (Uniper, Kavernen), 15.10.2026 (Notizen S. Maxara).",
     "nur tiefe Kaverne; Strang adiabat; Einspeisung gekoppelt an die Einspeicherung; Regelung über den Kopfdruck; Stillstand "
     "aus einer Quelle; Bandbreite Ausspeichergas"),
    ("—", "Span, R.; Wagner, W. (1996): A New Equation of State for Carbon Dioxide … J. Phys. Chem. Ref. Data 25(6), 1509–1596.",
     "Zustandsgleichung reines CO₂ (über CoolProp)"),
    ("—", "Kunz, O.; Wagner, W. (2012): The GERG-2008 Wide-Range Equation of State for Natural Gases and Other Mixtures. "
          "J. Chem. Eng. Data 57, 3032–3091.", "Mischungsregel für das Worst-Case-Gemisch"),
    ("—", "Bell, I. H. et al. (2014): Pure and Pseudo-pure Fluid Thermophysical Property Evaluation and the Open-Source "
          "Thermophysical Property Library CoolProp. Ind. Eng. Chem. Res. 53(6), 2498–2508.", "CoolProp 8.0.0"),
]


def blatt_quellen(wb):
    ws = blatt(wb, "Quellen", "", "A6A6A6", dict(B=12, C=80, D=70))
    kopf_blatt(ws, "Quellen", "Herkunft der Annahmen dieser Mappe (IDs wie in der Quellen- und Wissensmatrix), mit Fundstelle")
    kopfzeile(ws, 4, "BCD", ["ID", "Quelle", "Verwendet für (Fundstelle)"], 20)
    for i, (a, b, c) in enumerate(QUELLEN):
        r = 5 + i
        z(ws, f"B{r}", a, "textfett")
        z(ws, f"C{r}", b, "text")
        z(ws, f"D{r}", c, "text")
        ws.row_dimensions[r].height = 45


# ======================================================================================
def erzeuge(stoff):
    D = lade(stoff, "kaverne_stillstand.json")
    A = lade(stoff, "arbeitsgas.json")
    Z = lade(stoff, "kavernenzyklus.json")
    BK = lade(stoff, "bohrlochkopf.json")
    AG = lade(stoff, "zusammensetzung_nach_kaverne.json")
    V = lade("gem", "vergleich_rein_gemisch.json") if stoff == "gem" else None
    wb = Workbook()
    K = [("Kopfdruck im Stillstand, Kaverne voll", "=Kopf_voll", "bar", "0.0", "Blatt 01 – ruhende Säule"),
         ("Zieldruck der Pumpe im Betrieb S1 (100.000 Nm³/h)", "=Kopf_betr_S1", "bar", "0.0", "Blatt 01 – strömende Säule"),
         ("Zieldruck der Pumpe im Betrieb S2 (100.000 Nm³/h)", "=Kopf_betr_S2", "bar", "0.0", "Blatt 01 – Auslegung der Pumpe"),
         ("Kopfdruck im Stillstand, Kaverne leer", "=Kopf_leer", "bar", "0.0", "Blatt 01"),
         ("Inhalt bei 210 bar (TGV)", "=TGV", "t", "#,##0", "Blatt 02"),
         ("Arbeitsgas je Zyklus (S2, 100.000 Nm³/h, Stillstand 21,67 d)", "=AG_zyklus", "t", "#,##0", "Blatt 03"),
         ("Zyklusdauer / Zyklen pro Jahr", "=Zyklusdauer", "d", "0", "=\"→ \"&TEXT(Zyklen_a,\"0.0\")&\" Zyklen pro Jahr (Q-029: 4 bzw. 1)\""),
         ("Grenzrate: Kaverne bleibt beim Ausspeichern einphasig bis", "=Grenzrate", "kg/s", "0", "Blatt 03"),
         ("Kopfdruck beim Ausspeichern mindestens (100.000 Nm³/h)", "=Kopf_aus_min", "bar", "0.0",
          "Blatt 04 → Rückverdichtung auf 91 bar nötig"),
         ("Temperatur nach Drossel auf 30 bar mindestens", "=T_drossel_min", "°C", "0.0", "Blatt 04 → Vorwärmung"),
         ("Kavernen für den Bedarf nach Q-002 (unten / oben)", "=Kavernen_min", "–", "0",
          "=\"bis \"&TEXT(Kavernen_max,\"0\")&\" Kavernen bei 1,21 Mt (Blatt 05)\"")]
    uebersicht(wb, stoff, K)
    DD = Workbook().active   # Platzhalter, ersetzt unten
    DD = wb.create_sheet("Diagrammdaten")
    blatt_01(wb, stoff, D, DD)
    blatt_02(wb, stoff, A, DD)
    blatt_03(wb, stoff, Z, DD)
    blatt_04(wb, stoff, BK, DD)
    blatt_05(wb, stoff, Z)
    blatt_06(wb, stoff, AG)
    if stoff == "gem":
        blatt_vergleich(wb, V)
    blatt_quellen(wb)
    # Diagrammdaten ans Ende, Kopf
    wb.move_sheet("Diagrammdaten", offset=len(wb.sheetnames) - 1 - wb.sheetnames.index("Diagrammdaten"))
    DD.sheet_properties.tabColor = "D9D9D9"
    DD.sheet_view.showGridLines = False
    DD["B1"] = "Diagrammdaten – Hilfstabellen (nicht ändern)"
    DD["B1"].font = Font(name="Calibri", size=16, bold=True, color=BLAU)
    DD["B2"] = "Werte aus den JSON-Dateien der Skripte, Grundlage der Diagramme."
    DD["B2"].font = Font(name="Calibri", size=10, italic=True, color="595959")
    link(DD, "B3", "← zur Übersicht", "Übersicht")
    # Namen, die fehlen könnten (Blatt nicht gerechnet), als #N/A anlegen, damit die Übersicht nicht #NAME? zeigt
    for n in ("Kopf_voll", "Kopf_leer", "Kopf_betr_S1", "Kopf_betr_S2", "TGV", "AG_zyklus", "Zyklusdauer", "Zyklen_a",
              "Grenzrate", "Kopf_aus_min", "T_drossel_min", "Kavernen_min", "Kavernen_max"):
        if n not in wb.defined_names:
            wb.defined_names[n] = DefinedName(n, attr_text="NA()")
    wb.calculation.fullCalcOnLoad = True
    ziel = ORDNER[stoff] / DATEI[stoff]
    wb.save(ziel)
    print(f"gespeichert: {ziel}" + ("" if BK else "  (ohne Blatt 04 – bohrlochkopf.json fehlt)"))
    return ziel


if __name__ == "__main__":
    for s in (sys.argv[1:] or ["rein", "gem"]):
        erzeuge(s)
