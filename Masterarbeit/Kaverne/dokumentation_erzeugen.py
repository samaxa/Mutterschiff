# -*- coding: utf-8 -*-
"""
Erzeugt Dokumentation_Kaverne.docx aus den Ergebnissen der Schritte 01-06
============================================================================
Zahlen kommen aus den JSON-Dateien der Kette (also immer der letzte
Rechenstand), Annahmen und Begründungen aus den Tabellen unten. Formatvorlagen
wie die übrigen Dokumentationen (Vorlage: Einspeicherung/Begleitstoffe_CO2/
Dokumentation_S2_Gemisch.docx).

Aufruf (nach 01-06):  python dokumentation_erzeugen.py
"""
import datetime
import json

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm
from CoolProp import __version__ as COOLPROP_VERSION

import ausspeichergas as ag
import kaverne_bausteine as kb

VORLAGE = kb.ORDNER.parent / "Einspeicherung" / "Begleitstoffe_CO2" / "Dokumentation_S2_Gemisch.docx"
ZIEL = kb.ORDNER / "Dokumentation_Kaverne.docx"
R50, R100 = "50.000 Nm³/h", "100.000 Nm³/h"


def lade(name):
    with open(kb.ORDNER / name, encoding="utf-8") as f:
        return json.load(f)


S = lade("kaverne_stillstand.json")
A = lade("arbeitsgas.json")
Z = lade("kavernenzyklus.json")
K = lade("bohrlochkopf_ausspeicherung.json")
GR = lade("gemisch_vs_rein.json")
ZS = lade("zusammensetzung_nach_kaverne.json")


def de(x, nd=1):
    return f"{x:,.{nd}f}".replace(",", "X").replace(".", ",").replace("X", ".")


def kw(kurz, rate, i=0):
    """Kennwerte eines Zyklusabschnitts (0 = aus, 1 = still, 2 = ein, 3 = still), reines CO₂."""
    return Z[kurz]["raten"][rate]["kennwerte"][i]


def wgv_anteil(kurz, rate):
    a = kw(kurz, rate)
    return 100 * (a["m_start"] - a["m_ende"]) / a["m_start"]


def nm3h(m_dot):
    return m_dot * 3600 / kb.REIN.rho_norm()


# ---- Dokument-Bausteine --------------------------------------------------------------------
doc = Document(str(VORLAGE))
body = doc.element.body
for el in list(body):
    if el.tag != qn("w:sectPr"):
        body.remove(el)
zaehler = {"tab": 0, "abb": 0}


def absatz(text, stil="Body Text"):
    """Absatz; **fett** im Text wird fett gesetzt."""
    p = doc.add_paragraph(style=stil)
    teile = text.split("**")
    for i, t in enumerate(teile):
        if t:
            p.add_run(t).bold = (i % 2 == 1)
    return p


def ueberschrift(text, ebene=1):
    return doc.add_paragraph(text, style=f"Heading {ebene}")


def liste(punkte):
    for t in punkte:
        absatz("– " + t, "Compact")


def tabelle(kopf, zeilen, breiten_cm, titel):
    zaehler["tab"] += 1
    absatz(f"Tabelle {zaehler['tab']}: {titel}", "Table Caption")
    t = doc.add_table(rows=1 + len(zeilen), cols=len(kopf))
    t.style = doc.styles["Table"]
    for j, (h, b) in enumerate(zip(kopf, breiten_cm)):
        for i, inhalt in enumerate([h] + [z[j] for z in zeilen]):
            c = t.cell(i, j)
            c.width = Cm(b)
            p = c.paragraphs[0]
            p.style = doc.styles["Compact"]
            teile = str(inhalt).split("**")
            for k, s in enumerate(teile):
                if s:
                    p.add_run(s).bold = (i == 0) or (k % 2 == 1)
    tr = t.rows[0]._tr
    trPr = tr.get_or_add_trPr()
    kopfzeile = OxmlElement("w:tblHeader")
    kopfzeile.set(qn("w:val"), "on")
    trPr.append(kopfzeile)
    doc.add_paragraph(style="Body Text")
    return t


def abbildung(datei, titel, breite_cm=16.0):
    zaehler["abb"] += 1
    p = doc.add_paragraph(style="Captioned Figure")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(kb.ORDNER / datei), width=Cm(breite_cm))
    absatz(f"Abbildung {zaehler['abb']}: {titel}", "Image Caption")


# ---- Kennzahlen ------------------------------------------------------------------------------------
tief, flach = kb.TIEF, kb.FLACH
gr_flach = Z["flach"]["grenzrate"]
gr_tief = Z["tief"]["grenzrate"]
schwelle = K["p_lccs_strang_zweiphasig"]
kopf_bereich = {}
for kurz in ("tief", "flach"):
    pk = [z[m]["p"] for r in (R50, R100) for z in K[kurz][r] for m in ("adiabat", "gebirge") if z[m]]
    vk = [z["adiabat"]["v_max"] for r in (R50, R100) for z in K[kurz][r] if z["adiabat"]]
    kopf_bereich[kurz] = (min(pk), max(pk), max(vk))
d_tg = ZS["details"]["tief_gem"]


def rate_dpdt(kav, p_lccs, art):
    """Zulässige Rate aus 10 bar/d [kg/s] (wie Schritt 03)."""
    T = kav.T_mitte
    pm = kb.p_mitte(kb.REIN, kav, p_lccs, T)
    drho = kb.REIN.drho_dp(pm, T) if art == "isotherm" else 1e5 / kb.REIN.schall(pm, T) ** 2
    return kav.volumen * drho * kb.DPDT_MAX / 86400


r_iso_voll, r_ise_voll = rate_dpdt(kb.TIEF, kb.TIEF.p_max, "isotherm"), rate_dpdt(kb.TIEF, kb.TIEF.p_max, "isentrop")
fall_ausl = ZS["faelle"][ZS["auslegungsfall"]]["feucht"]


def zust_text(kurz, rate):
    a = kw(kurz, rate)
    if a["zweiphasig"]:
        return f"zweiphasig ab Tag {de(a['t_zweiphasig'])}"
    return f"einphasig, T_min {de(a['T_min'])} °C"


# ===== Titel ==========================================================================================
absatz("Dokumentation – Kaverne", "Title")
absatz("Stillstand, Arbeitsgas, Ein- und Ausspeicherrate, Bohrlochkopf und Ausspeichergas – "
       "tiefe und flache Kaverne, reines CO₂ und Worst-Case-Gemisch", "Subtitle")
absatz(f"Stand: {datetime.date.today():%d.%m.%Y} · CoolProp {COOLPROP_VERSION} · Skripte 01–06 im Ordner "
       f"Kaverne, erzeugt mit dokumentation_erzeugen.py (Zahlen aus den JSON-Ergebnissen der Kette)", "Date")
absatz("Die Kaverne selbst ist nicht Gegenstand der Arbeit, sie bestimmt aber, was die Obertageanlage "
       "leisten muss: den Kopfdruck beim Einspeichern, die Raten, den Kopfzustand beim Ausspeichern und "
       "die Zusammensetzung des Gases, das wieder herauskommt. Eine vollständige Simulation von Kaverne "
       "und Bohrung (Mehrphasenströmung, instationäre Wärmeleitung) ist mit den verfügbaren Werkzeugen "
       "nicht möglich. Die Skripte rechnen deshalb mit einfachen, nachvollziehbaren Modellen und "
       "Grenzfällen, die an den Quellen [Q-028], [Q-030] und [Q-089] geprüft sind.", "First Paragraph")

# ===== 1 Ergebnis in Kürze ==============================================================================
ueberschrift("1\tErgebnis in Kürze")
gr_flach_txt = (f"bis etwa {de(round(nm3h(gr_flach['rate']), -3), 0)} Nm³/h ({de(gr_flach['rate'])} kg/s)"
                if gr_flach["rate"] else "bei keiner Rate")
tabelle(["Frage", "Ergebnis"], [
    ["Welche Kaverne?",
     f"**Die tiefe Kaverne (V1, LCCS 1200 m).** Sie bleibt bei 50.000 und 100.000 Nm³/h beim Ausspeichern "
     f"einphasig (bei 100.000 Nm³/h kühlt sie bis {de(kw('tief', R100)['T_min'])} °C ab). Die flache Kaverne "
     f"(LCCS 700 m) hat mehr "
     f"Arbeitsgas, bleibt aber nur {gr_flach_txt} einphasig; bei 100.000 Nm³/h siedet ihr Inhalt "
     f"({zust_text('flach', R100)})."],
    ["Zieldruck der Einspeicherung (Kopf, voll)",
     f"tief: {de(S['tief']['saeulen']['voll']['p_kopf'], 2)} bar (rein) / "
     f"{de(GR['tief']['gem']['kopf_voll'], 2)} bar (Gemisch); flach: "
     f"{de(S['flach']['saeulen']['voll']['p_kopf'], 2)} / {de(GR['flach']['gem']['kopf_voll'], 2)} bar. "
     f"Identisch mit den Werten der Einspeicherung."],
    ["Arbeitsgas",
     f"tief: {de(100 * A['tief']['varianten'][0]['anteil_iso'], 0)} % des Inhalts (langsam) bis "
     f"{de(100 * A['tief']['varianten'][0]['anteil_ise'], 0)} % (schnell, ohne Wärmezufuhr); im Zyklus bei "
     f"100.000 Nm³/h {de(wgv_anteil('tief', R100), 0)} %. flach: {de(100 * A['flach']['varianten'][0]['anteil_iso'], 0)} % "
     f"(langsam)."],
    ["Ein- und Ausspeicherrate",
     f"Grenze ist die Druckänderung von 10 bar/d. Bei voller tiefer Kaverne erlaubt sie langsam "
     f"{de(r_iso_voll, 0)} kg/s, schnell nur {de(r_ise_voll, 0)} kg/s (100.000 Nm³/h = 54,9 kg/s); "
     f"bei 100.000 Nm³/h wird die Rate in "
     f"{de(100 * kw('tief', R100)['anteil_begrenzt'], 0)} % der Ausspeicherzeit gekürzt. Bei 650.000 m³ sind "
     f"100.000 Nm³/h für die tiefe Kaverne also möglich; bei 250.000 m³ (Uniper-Bestand) wären alle Raten aus 10 bar/d nur 38 % so groß."],
    ["Bohrlochkopf beim Ausspeichern",
     f"Kopfdruck tief {de(kopf_bereich['tief'][0], 0)}–{de(kopf_bereich['tief'][1], 0)} bar, flach "
     f"{de(kopf_bereich['flach'][0], 0)}–{de(kopf_bereich['flach'][1], 0)} bar. Im schnellen Grenzfall wird der "
     f"Förderstrang der tiefen Kaverne zweiphasig, sobald der Druck an der LCCS unter "
     f"{de(max(schwelle['tief'].values()), 0)} bar fällt – DBI hebt für dieselbe Teufe den Mindestdruck auf "
     f"160 bar an [Q-030]."],
    ["Ausspeichergas (Auslegungsfall)",
     f"Gesättigt mit Wasser: {de(1e6 * d_tg['y_h2o'], 0)} ppm = {de(ag.g_pro_Nm3(d_tg['y_h2o']), 2)} g/Nm³. "
     f"Dazu CH₄ aus Resterdgas {de(d_tg['ppm']['Resterdgas_CH4'], 0)} ppm (1. Zyklus), H₂S 10 ppm, "
     f"Spuren von Diesel und Verdichteröl (Tabelle {0})."],
    ["Was folgt für den Ausspeicherpfad",
     "Kopfdruck meist unter 91 bar → für dichte Übergabe Rückverdichtung. Drosseln auf 30 bar endet "
     "zweiphasig bei −5,6 °C → Vorwärmen vor der Drossel. Freies Wasser fällt am Kopf aus → Abscheider, "
     "Trocknung, Hydratschutz."],
], [4.0, 12.0], "Kernergebnisse (Zahlen aus den Skripten 01–06)")

# Platzhalter-Tabellennummer für die Zusammensetzung nachtragen (wird unten vergeben)
TAB_ZUSAMMENSETZUNG_PLATZHALTER = doc.tables[-1]

# ===== 2 Vorgehen =========================================================================================
ueberschrift("2\tVorgehen und Aufbau der Skripte")
absatz("Jeder Prozessschritt ist ein eigenes Skript. Die Skripte laufen in der Reihenfolge der Nummern und "
       "übergeben ihre Ergebnisse als JSON-Datei an die folgenden Schritte. Alle Annahmen stehen nur in "
       "kaverne_bausteine.py. Die Stoffmodelle (reines CO₂ nach Span & Wagner, Worst-Case-Gemisch nach "
       "GERG-2008) kommen unverändert aus Einspeicherung/Begleitstoffe_CO2, damit Ein- und Ausspeicherung "
       "mit demselben Stoff rechnen.", "First Paragraph")
tabelle(["Schritt", "Frage", "Übergabe an"], [
    ["01 Stillstand", "Was liegt in Kaverne und Bohrloch vor, wenn nichts strömt? Kopfdruck voll/leer.",
     "kaverne_stillstand.json → 02, Einspeicherung"],
    ["02 Arbeitsgas", "Wie viel Arbeitsgas, und warum hängt das an der Teufe?", "arbeitsgas.json"],
    ["03 Raten", "Welche Rate verträgt die Kaverne, was passiert bei 50.000/100.000 Nm³/h?",
     "kavernenzyklus.json → 04, 06"],
    ["04 Bohrlochkopf", "Was kommt beim Ausspeichern am Kopf an?", "bohrlochkopf_ausspeicherung.json → 06"],
    ["05 Gemisch", "Was ändert das Worst-Case-Gemisch in 01–04?", "gemisch_vs_rein.json → 06"],
    ["06 Zusammensetzung", "Mit welcher Zusammensetzung startet der Ausspeicherpfad?",
     "zusammensetzung_nach_kaverne.json → ausspeichergas.py"],
], [3.2, 8.3, 4.5], "Prozessschritte der Kaverne")

# ===== 3 Annahmen ===========================================================================================
ueberschrift("3\tAnnahmen")
absatz("Die beiden Kavernen haben dieselbe Form und dasselbe Volumen und unterscheiden sich nur in der Teufe. "
       "So zeigt der Vergleich allein den Einfluss der Teufe. Volumen und Geometrie folgen [Q-028], weil "
       "diese Quelle die einzige vollständige Thermodynamik-Simulation einer CO₂-Kaverne enthält und der "
       "Wärmeübergang daran kalibriert ist. Kennzeichnung: **Q** = aus einer Quelle übernommen, "
       "**A** = aus Quellen abgeleitet, **E** = eigene Annahme.", "First Paragraph")
tabelle(["Größe", "tief (V1)", "flach (V2)", "Herkunft"], [
    ["Teufe LCCS", "1200 m", "700 m", "Q – [Q-030] Tab. 1, Option 2 / Option 3"],
    ["Bezugsteufe Kaverneninhalt", f"{de(tief.z_mitte, 0)} m", f"{de(flach.z_mitte, 0)} m",
     "A – LCCS + 205 m wie [Q-028] Tab. 1 (LCCS 1150 m, Kavernentemperatur bei 1355 m)"],
    ["Gebirgstemperatur", f"{de(tief.T_mitte)} °C", f"{de(flach.T_mitte)} °C",
     "Q – 10 °C + 0,03 K/m [Q-028]"],
    ["p_max an der LCCS", f"{de(tief.p_max)} bar", f"{de(flach.p_max)} bar",
     "V1: Kaverne der Einspeicherung; V2: A – gleiche Gradienten (0,175 bar/m), [Q-030] Option 3: 125 bar"],
    ["p_min an der LCCS", f"{de(tief.p_min)} bar", f"{de(flach.p_min)} bar",
     "V1 wie Einspeicherung; V2: A – 0,0583 bar/m, [Q-030] Option 3: 40 bar"],
    ["Volumen", "650.000 m³", "650.000 m³",
     "Q – [Q-028] Tab. 1. Uniper-Bestandskavernen: 250.000 m³ [Q-030] Tab. 1 (siehe 3.1)"],
    ["Druckänderungsrate", "≤ 10 bar/d", "≤ 10 bar/d", "Q – [Q-030] Kap. 1.2.1, [Q-089] Tab. 2; vorsichtig 6 bar/d [Q-028]"],
    ["Förderstrang", "8 5/8\", Ø 198,8 mm", "8 5/8\", Ø 198,8 mm", "Q – [Q-030] Kap. 1.2.1"],
    ["Strömungsgeschwindigkeit", "≤ 20 m/s", "≤ 20 m/s", "Q – [Q-030] Kap. 1.2.1"],
    ["Rohrreibung (Darcy)", "0,016", "0,016", "Q – [Q-003] Kap. 6.2: 0,015–0,017 für neues Stahlrohr"],
    ["Wärmeübergang UA", "546 kW/K", "546 kW/K",
     "A – an [Q-028] kalibriert (Kavernenrechner, Nachrechnung in 03)"],
    ["Einspeisetemperatur an der LCCS", "30 °C", "30 °C", "Q – [Q-028] Tab. 1"],
    ["Stillstand zwischen Aus- und Einspeichern", "21,67 d", "21,67 d", "Q – [Q-028] Kap. 4"],
    ["Durchsatz", "50.000 / 100.000 Nm³/h", "50.000 / 100.000 Nm³/h", "wie Einspeicherung"],
    ["Sicherheitsabstand zu T_krit", "3 K", "3 K", "E – wie Kavernenrechner (nur Teufenfenster)"],
    ["Übergabedrücke Ausspeicherung", "91 / 30 bar", "91 / 30 bar", "E – wie Einspeicherung S1 / S2"],
], [3.8, 2.4, 2.4, 7.4], "Annahmen der Kavernenrechnung (kaverne_bausteine.py)")
ueberschrift("3.1\tWarum beide Kavernen 650.000 m³ haben", 2)
absatz("Ein Vergleich „kleine flache gegen große tiefe Kaverne“ würde Teufe und Volumen gleichzeitig ändern, "
       "die Ursache eines Unterschieds wäre dann nicht mehr zuzuordnen. Die meisten Ergebnisse hängen gar "
       "nicht vom Volumen ab: Phasenlage, Temperaturen, Kopfzustand, Arbeitsgas in Prozent und Wassergehalt. "
       "Linear mit dem Volumen skalieren nur die Massen in Tonnen und die zulässige Rate aus der "
       "Druckänderungsrate: Bei 250.000 m³ (Uniper-Bestandskavernen, [Q-030] Tab. 1 – dort für alle Teufen "
       "gleich) wären es 38 % der hier genannten Raten. Große Kavernen in geringer Teufe sind nicht "
       "ungewöhnlich: [Q-003] legt eine Kaverne mit 800.000 m³ in 800 m Teufe aus.", "First Paragraph")
absatz("Der Wärmeübergang skaliert mit der Oberfläche (∝ V^2/3); für 250.000 m³ wären es 289 kW/K. "
       "Offen: Volumen einer realen Kaverne mit Uniper abstimmen.")
ueberschrift("3.2\tModellgrenzen", 2)
liste(["Kaverne als ein ideal durchmischter Knoten. [Q-028] beschreibt dauerhafte Konvektion im Kavernenraum, "
       "das stützt die Durchmischung; Schichtung (Flüssigkeit im Sumpf) bildet das Modell nicht ab.",
       "Wärmeübergang als konstantes UA statt instationärer Wärmeleitung im Gebirge.",
       "Förderstrang in zwei Grenzfällen: adiabat (schnell) und mit Gebirgstemperatur (langsam). "
       "Der reale Zustand liegt dazwischen; genauer nur mit einem Bohrlochmodell mit Mehrphasenströmung [Q-028].",
       "Keine Sole im Modell (\"completely debrined\", [Q-028]); CO₂-Lösung in Restsole nur abgeschätzt (Kap. 9).",
       "Gemisch: im Zweiphasengebiet endet die Kavernenbilanz an der Phasengrenze, weil der schnelle "
       "Stoffwert-Aufruf dort nur metastabile Werte liefert."])

# ===== 4 Stillstand ==========================================================================================
ueberschrift("4\tStillstand (Schritt 01)")
absatz("Die Säule im Bohrloch wird von der LCCS nach oben integriert, dp/dz = ρ(p, T)·g, mit der "
       "Gebirgstemperatur. Drei Füllstände wie [Q-028] Fig. 3: voll, Grenzfall (Kopfdruck = kritischer Druck) "
       "und leer.", "First Paragraph")
zeilen = []
for kurz in ("tief", "flach"):
    s = S[kurz]["saeulen"]
    for fall, x in s.items():
        zeilen.append([f"{kurz}, {fall}", f"{de(x['p_lccs'])} bar", f"{de(x['p_kopf'], 2)} bar", x["phase_kopf"],
                       f"{de(x['rho_kaverne'], 0)} kg/m³, {x['phase_kaverne']}"])
tabelle(["Kaverne, Füllstand", "p LCCS", "p Kopf", "Phase Kopf (10 °C)", "Kaverneninhalt"], zeilen,
        [3.2, 2.2, 2.4, 3.6, 4.6], "Gassäule im Stillstand, reines CO₂")
absatz(f"Kontrolle: tief/voll ergibt {de(S['tief']['saeulen']['voll']['p_kopf'], 2)} bar wie in der Einspeicherung. "
       "In der flachen Kaverne bleibt der Kopf in jedem Füllstand unter dem kritischen Druck: Das CO₂ wechselt "
       "im Bohrloch immer die Phase. Bei leerer tiefer Kaverne liegt der obere Teil des Bohrlochs auf der "
       "Siedelinie (Nassdampf, [Q-028]: oberhalb etwa 250 m).")
abbildung("abb_K01_gassaeulen.png", "Gassäulen im Stillstand, Darstellung nach [Q-028] Fig. 3 (Skript 01)")

# ===== 5 Arbeitsgas ===========================================================================================
ueberschrift("5\tArbeitsgas und Teufe (Schritt 02)")
absatz("Arbeitsgas = Inhalt bei p_max minus Inhalt bei p_min, m = ρ(p, T)·V wie DBI [Q-030] Kap. 1.3. "
       "Zwei Grenzfälle für den Inhalt bei p_min: langsam (isotherm, das Gebirge heizt nach) und schnell "
       "(isentrop, keine Wärmezufuhr). Die Dichten von CoolProp stimmen mit DBI für Option 2 (= tiefe Kaverne) "
       "auf 0,1 % überein.", "First Paragraph")
zeilen = []
for kurz in ("tief", "flach"):
    for v in A[kurz]["varianten"]:
        zp = f", zweiphasig (q = {de(v['q_ise'], 2)})" if v["q_ise"] is not None else ""
        zeilen.append([f"{kurz}: {v['variante']}", f"{de(v['p_min'])} bar", f"{de(100 * v['anteil_iso'], 0)} %",
                       f"{de(100 * v['anteil_ise'], 0)} % ({de(v['T_ise'])} °C{zp})"])
tabelle(["Kaverne: p_min", "p_min LCCS", "Arbeitsgas langsam", "Arbeitsgas schnell (T Ende)"], zeilen,
        [4.2, 2.4, 3.0, 6.4], "Arbeitsgasanteil (Anteil am Inhalt bei p_max)")
tf = A["teufenfenster"]
absatz(f"**Teufenfenster.** Untere Grenze: Das Gebirge soll im Kaverneninhalt mindestens 3 K über der "
       f"kritischen Temperatur liegen, sonst bildet sich beim Abkühlen Flüssigkeit in der Kaverne – LCCS ≥ "
       f"{de(tf['z_lccs_min'], 0)} m. Obere Grenze: p_min soll unter dem kritischen Druck liegen, sonst bleibt "
       f"der Inhalt immer dicht und das Arbeitsgas liegt im flachen Teil der Dichtekurve – LCCS ≤ "
       f"{de(tf['z_lccs_max'], 0)} m. Beide Kavernen liegen im Fenster, die flache aber nur "
       f"{de(flach.z_lccs - tf['z_lccs_min'], 0)} m über der unteren Grenze. Das Fenster gilt für die ruhende "
       f"Kaverne; wie weit der Inhalt beim Ausspeichern abkühlt, hängt von der Rate ab (Kap. 6). "
       f"[Q-028] Kap. 5 kommt qualitativ zum selben Schluss: das größte Arbeitsgas im Übergang zwischen "
       f"überkritisch und gasförmig, flachere Kavernen mit dem Risiko einer Flüssigphase.")
abbildung("abb_K02_arbeitsgas.png", "Dichte über Druck und Arbeitsgas je Mindestdruck (Skript 02)")

# ===== 6 Raten =================================================================================================
ueberschrift("6\tEin- und Ausspeicherrate (Schritt 03)")
absatz("Die Rate einer Kaverne ist durch drei Dinge begrenzt: die zulässige Druckänderung an der LCCS, den "
       "Phasenzustand des Kaverneninhalts (er kühlt beim Ausspeichern ab) und die Strömungsgeschwindigkeit im "
       "Förderstrang (Kap. 7).", "First Paragraph")
ueberschrift("6.1\tGrenze aus der Druckänderungsrate", 2)
absatz("ṁ_max = V · (∂ρ/∂p) · (dp/dt)_max – die Masse, die je bar Druckänderung aus der Kaverne kommt "
       "([Q-028] Fig. 5), mal der zulässigen Druckänderung. Langsam gilt die isotherme, schnell die isentrope "
       "Ableitung (1/c²). Hohe Raten sind nur nahe dem kritischen Druck möglich, bei voller Kaverne ist die "
       "Rate am kleinsten.")
abbildung("abb_K03_ratengrenze.png", "Zulässige Rate aus 10 bar/d über dem Druck an der LCCS (Skript 03)")
ueberschrift("6.2\tKavernenbilanz und Validierung", 2)
v = Z["validierung_q028"]
absatz(f"Masse und innere Energie des Kaverneninhalts werden über einen Zyklus wie in [Q-028] bilanziert: "
       f"ausspeichern bis p_min, 21,67 d Stillstand, einspeichern bis p_max, 21,67 d Stillstand; die Rate wird "
       f"gekürzt, sobald der Druck schneller als 10 bar/d ändert. Nachrechnung von [Q-028] (650.000 m³, 55 °C, "
       f"135 → 70 bar, höchste Rate bei 6 bar/d): Ende {de(v['p_ende'])} bar / {de(v['T_ende'])} °C "
       f"(Quelle 77,6 bar / 33,1 °C), Raten {de(v['rate_min_th'], 0)}–{de(v['rate_max_th'], 0)} t/h "
       f"(Quelle etwa 90–600 t/h).")
zeilen = []
for kurz in ("tief", "flach"):
    for r in (R50, R100):
        a, st, e = kw(kurz, r, 0), kw(kurz, r, 1), kw(kurz, r, 2)
        zeilen.append([f"{kurz}, {r}", f"{de(a['dauer'], 0)} d", zust_text(kurz, r),
                       f"{de(wgv_anteil(kurz, r), 0)} %", f"{de(100 * a['anteil_begrenzt'], 0)} %",
                       f"{de(st['p_lccs_start'], 0)} → {de(st['p_lccs_ende'], 0)} bar"])
tabelle(["Kaverne, Rate", "Ausspeichern", "Kaverneninhalt", "Arbeitsgas", "Rate gekürzt",
         "Stillstand danach"], zeilen, [3.4, 2.2, 3.8, 2.0, 2.0, 2.6],
        "Ausspeichern im Zyklus (reines CO₂, 10 bar/d)")
gr_t = ("die Druckänderungsrate allein begrenzt" if not gr_tief["bindend"] else
        f"etwa {de(round(nm3h(gr_tief['rate']), -3), 0)} Nm³/h")
absatz(f"**Grenzrate** (höchste konstante Ausspeicherrate, bei der der Inhalt bis p_min einphasig bleibt): "
       f"tief – {gr_t}; flach – {gr_flach_txt}. In der flachen Kaverne fällt der Inhalt bei höheren Raten ins "
       f"Zweiphasengebiet; der Druck hängt dann nur noch an der Temperatur und ist kein Maß mehr für den "
       f"Inhalt [Q-028]. Nach dem Einspeichern fällt der Druck im Stillstand, weil der Inhalt abkühlt – "
       f"die Kaverne ist danach nicht mehr ganz voll.")
abbildung("abb_K03_zyklus.png", "Druck, Temperatur und Inhalt über einen Zyklus (Skript 03)")

# ===== 7 Bohrlochkopf ===========================================================================================
ueberschrift("7\tBohrlochkopf beim Ausspeichern (Schritt 04)")
absatz("Für Zeitpunkte während des Ausspeicherns wird der Förderstrang von der LCCS zum Kopf gerechnet: Schwere "
       "und Reibung, dazu entweder h + g·z = konstant (adiabat, schnell) oder die Gebirgstemperatur "
       "(langsam). Der Kopfzustand ist der Startpunkt des Ausspeicherpfads.", "First Paragraph")
zeilen = []
for kurz in ("tief", "flach"):
    for r in (R50, R100):
        pts = [z["adiabat"] for z in K[kurz][r] if z["adiabat"]]
        zeilen.append([f"{kurz}, {r}", f"{de(min(p['p'] for p in pts), 0)}–{de(max(p['p'] for p in pts), 0)} bar",
                       f"{de(min(p['T'] for p in pts), 0)}–{de(max(p['T'] for p in pts), 0)} °C",
                       f"{de(max(p['v_max'] for p in pts))} m/s",
                       f"< {de(schwelle[kurz][r], 0)} bar" if r in schwelle[kurz] else "–"])
tabelle(["Kaverne, Rate", "Kopfdruck (adiabat)", "Kopftemperatur", "v_max Strang", "Strang zweiphasig ab p_LCCS"],
        zeilen, [3.4, 3.2, 2.8, 2.6, 4.0], "Kopfzustand beim Ausspeichern, schneller Grenzfall")
absatz("Folgen für die Obertageanlage: Der Kopfdruck liegt fast immer unter 91 bar – für eine dichte Übergabe "
       "muss zurückverdichtet werden. Wird auf 30 bar gedrosselt, endet das CO₂ zweiphasig bei −5,6 °C "
       "(Sättigungstemperatur bei 30 bar) – vor der Drossel muss vorgewärmt werden. Die "
       "Strömungsgeschwindigkeit bleibt in der tiefen Kaverne unter 20 m/s; in der flachen Kaverne wird sie "
       "bei 100.000 Nm³/h gegen Ende überschritten. Ab welchem Druck der Strang zweiphasig wird, deckt sich "
       "mit [Q-030]: DBI hebt für Option 2 (1200 m) den Mindestdruck auf 160 bar an.")
abbildung("abb_K04_kopfzustand_pT.png", "Kopfzustände beim Ausspeichern im p-T-Diagramm (Skript 04)")

# ===== 8 Gemisch ===================================================================================================
ueberschrift("8\tWorst-Case-Gemisch gegenüber reinem CO₂ (Schritt 05)")
absatz("Das Gemisch (95 % CO₂, 2,4 % N₂, 1 % Ar, 1 % CH₄, 0,5 % H₂, 0,1 % CO) hat kein Siedelinie, sondern ein "
       "Zweiphasengebiet bis zur Cricondenbar 82,2 bar und Cricondentherm 28,1 °C (reines CO₂: kritischer "
       "Punkt 31,0 °C / 73,8 bar). Flüssigkeit entsteht also bei höherem Druck, aber erst bei tieferer "
       "Temperatur.", "First Paragraph")
zeilen = []
for kurz in ("tief", "flach"):
    r_, g_ = GR[kurz]["rein"], GR[kurz]["gem"]
    for name, x, y, nd, e in (("Kopfdruck voll", r_["kopf_voll"], g_["kopf_voll"], 2, "bar"),
                              ("Kopfdruck leer", r_["kopf_leer"], g_["kopf_leer"], 2, "bar"),
                              ("Inhalt voll", r_["TGV"] / 1e3, g_["TGV"] / 1e3, 0, "kt"),
                              ("Arbeitsgas langsam", r_["WGV"] / 1e3, g_["WGV"] / 1e3, 0, "kt")):
        zeilen.append([f"{kurz}: {name}", f"{de(x, nd)} {e}", f"{de(y, nd)} {e}"])
    for r in (R50, R100):
        def t(x):
            a = x["raten"][r]
            return (f"zweiphasig nach {de(a['kennwerte']['dauer'])} d" if a["zweiphasig"]
                    else f"einphasig, T_min {de(a['kennwerte']['T_min'])} °C")
        zeilen.append([f"{kurz}: Ausspeichern {r}", t(r_), t(g_)])
tabelle(["Kaverne: Größe", "reines CO₂", "Gemisch"], zeilen, [6.0, 5.0, 5.0], "Reines CO₂ und Gemisch im Vergleich")
abbildung("abb_K05_kaverne_huelle_pT.png", "Kavernenzustand beim Ausspeichern mit 100.000 Nm³/h gegen die "
          "Phasengrenzen (Skript 05)")

# ===== 9 Zusammensetzung =============================================================================================
ueberschrift("9\tZusammensetzung des Ausspeichergases (Schritt 06)")
absatz("Wie sich das CO₂ in der Kaverne verunreinigt, lässt sich nicht genau vorhersagen. Die Annahmen folgen "
       "deshalb der Konzeptstudie [Q-089] für Wasserstoff in einer Uniper-Musterkaverne (frühere "
       "Erdgaskaverne, Diesel-Blanket, ölgeschmierte Kolbenverdichter) und werden auf CO₂ übertragen. "
       "Dieselben Quellen gelten auch für CO₂, die Mengen ändern sich aber mit dem Stoff.", "First Paragraph")
d = ZS["details"]
tabelle(["Eintrag", "Ansatz bei H₂ [Q-089]", "Übertrag auf CO₂", "tief, Gemisch"], [
    ["Wasser", "Sättigung, voll 75 % (Sole), leer 100 %; 0,47–1,38 g/Nm³",
     "**Gesättigt** über reinem Wasser im wasserreichsten Zustand (Auslegung). CO₂ löst im dichten Zustand "
     "viel mehr Wasser als H₂ und als gasförmiges CO₂. Modell Spycher et al. (2003).",
     f"{de(1e6 * d['tief_gem']['y_h2o'], 0)} ppm (Q-089-Ansatz: {de(1e6 * d['tief_gem']['y_q089'], 0)} ppm)"],
    ["Resterdgas (CH₄)", "Finger 2000 m³ in 500.000 m³, homogen vermischt; +0,5 %",
     "Gleiches Volumenverhältnis; dichtes CO₂ enthält je m³ ~2,5-mal so viele Mole wie H₂ → kleinerer Anteil. "
     "Nur 1. Zyklus.", f"{de(d['tief_gem']['ppm']['Resterdgas_CH4'], 0)} ppm"],
    ["Diesel-Blanket", "Restblanket 0,08–1,76 m³; H₂ sättigt bei 162 kg",
     "Dichtes CO₂ ist ein Lösungsmittel (Extraktion, [Q-003] Kap. 8.2.2.1) → 1,76 m³ vollständig gelöst",
     f"{de(d['tief_gem']['ppm']['Diesel'], 2)} ppm"],
    ["Verdichteröl", "0,135–0,26 kg je 100.000 Nm³", "Oberer Wert, nur bei Verdichtern (S2-Pfad)",
     f"{de(d['tief_gem']['ppm']['Verdichteroel'], 2)} ppm"],
    ["H₂S (mikrobiell)", "10 ppm angenommen", "Übernommen; zulässig sind 10 ppm [Q-030] Kap. 1.4. Saure Sole "
     "und hoher Salzgehalt hemmen die Mikroben eher als bei H₂.", "10 ppm"],
    ["CH₄ (mikrobiell)", "20–50 ppm", "Nur Gemisch: Methanbildner brauchen H₂ (4 H₂ + CO₂ → CH₄ + 2 H₂O); "
     "50 ppm verbrauchen 200 ppm H₂. Reines CO₂ enthält keinen Wasserstoff.", "50 ppm"],
    ["CO₂ in Restsole", "–", "E – 1 % des Volumens Restsole, Löslichkeit nach Spycher (obere Schranke)",
     f"{de(d['tief_gem']['co2_sole_t'], 0)} t, < 0,1 % des Inhalts"],
], [2.6, 3.8, 6.6, 3.0], "Einträge in der Kaverne: Übertrag von [Q-089] auf CO₂")
tab_nr_zus = zaehler["tab"] + 1
zeilen = [[k.replace("CarbonMonoxide", "CO").replace("HydrogenSulfide", "H₂S").replace("Nitrogen", "N₂")
           .replace("Argon", "Ar").replace("Methane", "CH₄").replace("Hydrogen", "H₂").replace("Water", "H₂O")
           .replace("Verdichteroel", "Verdichteröl"),
           de(100 * x, 4)] for k, x in fall_ausl.items()]
tabelle(["Komponente", "mol-%"], zeilen, [6.0, 4.0],
        f"Ausspeichergas im Auslegungsfall ({ZS['auslegungsfall']}: tiefe Kaverne, Gemisch, 1. Zyklus, feucht)")
absatz("Für den Ausspeicherpfad stellt ausspeichergas.py die Zusammensetzung bereit: gemisch_trocken() als "
       "CoolProp-Zustand des getrockneten Gases (schwere Kohlenwasserstoffe nur als Spuren, ohne Stoffmodell) "
       "und wasser_im_co2(p, T) für Abscheider und Trocknung. Am Kopf fällt in allen Fällen freies Wasser aus, "
       "weil das Gas beim Aufsteigen abkühlt und sein Druck sinkt. Mit freiem Wasser kann unter etwa 10 °C "
       "CO₂-Hydrat entstehen (nach dem Drosseln auf 30 bar unter etwa 6,7 °C).")
abbildung("abb_K06_wasser.png", "Wassersättigung von CO₂ und Kavernen-/Kopfzustände (Skript 06)")
abbildung("abb_K06_eintraege.png", "Einträge je Quelle (Skript 06)")

# ===== 10 Offene Punkte ==================================================================================================
ueberschrift("10\tOffene Punkte und Quellen, die noch nachzutragen sind")
liste(["Volumen einer realen Kaverne mit Uniper abstimmen (hier 650.000 m³ nach [Q-028]; Bestand 250.000 m³).",
       "In die Quellenmatrix aufnehmen und nachvollziehen: Spycher, N.; Pruess, K.; Ennis-King, J. (2003): "
       "CO₂-H₂O mixtures in the geological sequestration of CO₂. I. Assessment and calculation of mutual "
       "solubilities from 12 to 100 °C and up to 600 bar. Geochimica et Cosmochimica Acta 67 (16), 3015–3031 "
       "– Wassersättigung und CO₂-Löslichkeit.",
       "In die Quellenmatrix aufnehmen und nachvollziehen: Sloan, E. D.; Koh, C. A. (2008): Clathrate Hydrates "
       "of Natural Gases, 3. Aufl., CRC Press – Quadrupelpunkte der CO₂-Hydratlinie (≈ 0 °C / 12,6 bar und "
       "≈ 9,8 °C / 44,9 bar). Die Hydratgrenze ist im Skript nur angenähert.",
       "Spycher et al. gilt ab 12 °C. Kopf- und Drosselzustände darunter sind extrapoliert und im Skript gekennzeichnet.",
       "Restsole 1 % des Volumens ist eine eigene Annahme (nur zur Größenordnung).",
       "Für eine genauere Aussage zum Kopfzustand wäre ein Bohrlochmodell mit Mehrphasenströmung nötig "
       "([Q-028] Kap. 5; DGMK-Projektskizze [Q-019] AP 4/5)."])

ueberschrift("11\tDateien")
tabelle(["Datei", "Inhalt"], [
    ["kaverne_bausteine.py", "Annahmen, Kavernen, Stoffschnittstelle, Gassäule, Kavernenbilanz, Förderstrang"],
    ["01_…06_….py", "Prozessschritte (Kap. 4–9)"],
    ["ausspeichergas.py", "Stoffdatenblatt für den Ausspeicherpfad (Wasser, Hydrat, Zusammensetzung)"],
    ["*.json", "Übergaben zwischen den Schritten"],
    ["abb_K0x_….png", "Abbildungen"],
    ["dokumentation_erzeugen.py", "erzeugt diese Datei"],
], [5.0, 11.0], "Dateien im Ordner Kaverne")

# Verweis auf die Zusammensetzungstabelle in Tabelle 1 nachtragen
for zelle in TAB_ZUSAMMENSETZUNG_PLATZHALTER.rows[6].cells[1:]:
    for r in zelle.paragraphs[0].runs:
        r.text = r.text.replace("(Tabelle 0)", f"(Tabelle {tab_nr_zus})")

doc.core_properties.title = "Dokumentation – Kaverne"
doc.save(str(ZIEL))
print(f"gespeichert: {ZIEL.name} ({zaehler['tab']} Tabellen, {zaehler['abb']} Abbildungen)")
