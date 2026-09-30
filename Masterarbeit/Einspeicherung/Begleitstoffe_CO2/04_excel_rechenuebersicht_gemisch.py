# -*- coding: utf-8 -*-
"""
Schritt 4 (Gemisch) - Excel-Rechenübersicht für das Worst-Case-Gemisch
============================================================================
Erzeugt aus der Mappe für reines CO2 (CO2_Zwischenspeicher_Rechenuebersicht2.xlsx)
die gleiche Mappe für das Worst-Case-Gemisch:

  - gleicher Aufbau, gleiche Formeln, gleiche Diagramme
  - alle grauen CoolProp-Werte neu gerechnet (excel_werte_gemisch.py)
  - Phasendiagramme: zusätzliches Zweiphasengebiet, Phasengrenze des Gemischs,
    Sicherheitsabstand ±3 bar, Sättigungslinie reines CO2 zum Vergleich
  - neues Blatt "Gemisch & Grenzen": Zusammensetzung, Kennpunkte der
    Phasengrenze, Tabelle Tau-/Blasendruck, neue Grenzen der Anlage
  - neue Verdichtung: Enddruck p_V = 90 bar (über Cricondenbar + 3 bar), dort in
    der dichten Phase auf 20 °C kühlen (kein Zweiphasengebiet), Pumpe bis
    Kopfdruck; 1 Stufe überschreitet 95 °C -> 2 Stufen
  - alte Verdichtung (wie reines CO2: 60 bar, verflüssigen) zum Vergleich:
    Pfad im Phasendiagramm und Block auf "Gemisch & Grenzen"

Aufruf:
    python 04_excel_rechenuebersicht_gemisch.py  ORIGINAL.xlsx  [ZIEL.xlsx]

Die Datei wird auf XML-Ebene bearbeitet (xlsx_xml.py), weil openpyxl beim
Speichern Diagramme und Formen der Mappe beschädigt. Excel rechnet beim
Öffnen alle Formeln neu (fullCalcOnLoad).
"""
import re
import sys
import uuid

import numpy as np
from xml.sax.saxutils import escape

import gemisch_worstcase as gw
from excel_werte_gemisch import werte, U_PG
from xlsx_xml import Paket, SharedStrings, Blatt, formelwerte_loeschen


def de(x, n=1):
    """Zahl im deutschen Format: 114.93 -> '114,9'."""
    return f"{x:.{n}f}".replace(".", ",")


QUELLE = sys.argv[1] if len(sys.argv) > 1 else "CO2_Zwischenspeicher_Rechenuebersicht2.xlsx"
ZIEL = sys.argv[2] if len(sys.argv) > 2 else "CO2_Zwischenspeicher_Rechenuebersicht_Gemisch.xlsx"

print("Rechne Gemischwerte ...")
W = werte()
A, K, S1, S2, S2v, Dv = W["A"], W["kav"], W["S1"], W["S2"], W["S2v"], W["Dv"]
pg, dia = W["pg"], W["dia"]
p_kopf = K["voll"]["p_kopf"]
p_cb = pg["cricondenbar"][1]
T_kG, p_kG = pg["krit"]

# abgeleitete Ergebnisse (wie die Formeln der Mappe)
rho_n = S1["rho_n"]
m_min, m_max = A["V_min"] * rho_n / 3600, A["V_max"] * rho_n / 3600
w_S2 = S2["w1"] + S2["wP"]
w_S2v = S2v["w1"] + S2v["w2"] + S2["wP"]
w_Dv = Dv["w1"] + Dv["w2"]
ersparnis_v = 1 - w_S2v / w_S2
vorteil_S2 = 1 - w_S2v / w_Dv     # Gemisch: 2 Stufen sind Pflicht, deshalb Variante vergleichen
REIN = dict(p_kopf=107.603219569636, p_grenz=171.6522772915801, p_leer=44.99299418696985,
            tief_leer=236, rho_unten=817.604266524432, psat_TK=57.29052581475148,
            T1a_1stufe=73.15539245907911, w_S1=2.3473871252743947, w_S2=46.42488823757327)

pk = Paket(QUELLE)
sst = SharedStrings(pk)
B = {n: Blatt(pk, f"xl/worksheets/sheet{i}.xml", sst) for i, n in enumerate(
    ["Übersicht", "Kaverne", "Medien", "S1", "S2", "Phasendiagramm", "Quellen", "Diagrammdaten"], 1)}

# =============================================================================
# 1) Übersicht
# =============================================================================
u = B["Übersicht"]
u.text("B1", "CO₂-Zwischenspeicherung in Salzkavernen – Obertageanlage · Worst-Case-Gemisch")
u.text("B3", "Stand: 30.09.2026 · Stoffdaten: CoolProp 8.0.0, Mehrfluid-Helmholtz-Modell "
             "(GERG-2008-Mischungsregel, Reinstoffe u. a. Span & Wagner 1996) · Worst-Case-Gemisch "
             "95 % CO₂, 2,4 % N₂, 1 % Ar, 1 % CH₄, 0,5 % H₂, 0,1 % CO (mol-%)")
u.text("B6", "Diese Mappe zeigt Schritt für Schritt, wie die Obertageanlage für die Einspeicherung "
             "berechnet wird:\n1. Welcher Druck muss am Bohrlochkopf anliegen? → das gibt die Kaverne über "
             "die Gassäule vor.\n2. Wie bringt die Anlage das CO₂ auf diesen Druck? → Pumpe (dichte "
             "Anlieferung, S1) bzw. Verdichter, Kühler und Pumpe (gasförmige Anlieferung, S2).\n"
             "3. Wo liegen alle Zustände im Phasendiagramm – und wo liegt das Betriebsfenster von Bohrloch "
             "und Kaverne?\nDiese Mappe rechnet das Worst-Case-Gemisch mit 5 % Begleitstoffen – gleicher "
             "Aufbau wie die Mappe für reines CO₂. Was sich an den Grenzen ändert und warum: Blatt "
             "„Gemisch & Grenzen“.")
u.zahl("J7", round(S1["h1"], 1))
u.zahl("J8", round(m_min * S1["w"], 1))
u.zahl("J9", round(p_kopf, 1))
u.text("E16", "Anlieferung gasförmig: Verdichtung (2 Stufen), Zwischenkühlung, Kühlung in der dichten Phase, Pumpe.")
# Inhaltsverzeichnis: neues Blatt
u.text("B20", "→ Gemisch & Grenzen", stil=u.stil("B19"))
u.text("E20", "Was die Begleitstoffe ändern: Zweiphasengebiet, Unsicherheit ±3 bar, neue Grenzen der Anlage.",
       stil=u.stil("E19"))
u.ersetze('<mergeCell ref="B19:D19"/>', '<mergeCell ref="B19:D19"/><mergeCell ref="B20:D20"/><mergeCell ref="E20:O20"/>')
u.xml = re.sub(r'<mergeCells count="(\d+)"', lambda m_: f'<mergeCells count="{int(m_.group(1)) + 2}"', u.xml, count=1)
u.ersetze('</hyperlinks>', '<hyperlink ref="B20" location="\'Gemisch &amp; Grenzen\'!A1" '
                            'display="→ Gemisch &amp; Grenzen"/></hyperlinks>')
# Annahme Verflüssigungsdruck
u.zahl("F50", A["p_V"])
u.zahl("M50", A["p_V"])
u.text("B50", "Enddruck Verdichtung = Kühlerdruck")
u.text("H50", f"Gemisch: ≥ Cricondenbar {de(p_cb)} + 3 bar → Kühlen ohne Zweiphasengebiet (rein: 60 bar, verflüssigen)")
u.text("H51", f"Gemisch: ≤ Taudruck(T_K) − 3 bar = {de(Dv['p_zw_max'])} bar; Stufe 2 ≤ 95 °C ab "
              f"{de(Dv['p_zw_min'])} bar → zulässig ca. {Dv['p_zw_min']:.0f}–{Dv['p_zw_max']:.0f} bar")
u.text("H76", f"Gemisch: oberes Bohrloch 0–{K['leer']['tief_2ph']:.0f} m zweiphasig (rein: 45,0 bar, 0–236 m)")
u.text("B80", "S2 Basisfall 1 Stufe: spezifische Arbeit gesamt")
u.text("H80", f"Verdichter + Pumpe; 1 Stufe erreicht {de(S2['T1a'])} °C > 95 °C (Q-016)")
u.text("H82", f"Pflicht beim Gemisch (1 Stufe > 95 °C); spart außerdem rund {ersparnis_v:.0%}")
u.text("H83", f"S2 (2 Stufen) braucht rund {vorteil_S2:.0%} weniger Arbeit – deshalb dicht kühlen und pumpen")
u.speichere()

# =============================================================================
# 2) Kaverne & Gassäule
# =============================================================================
k = B["Kaverne"]
k.text("B2", "Einspeicherung, Kaverne V1 (tief) · Stillstand: Temperatur = Gebirgstemperatur · Worst-Case-Gemisch")
k.text("B4", "WAS WIRD BERECHNET?  Die Kaverne legt fest, welcher Druck unten an der LCCS (letzte zementierte "
             "Rohrtour) höchstens anliegen darf: 210 bar. Gesucht ist der Druck, den die Anlage dafür oben am "
             "Bohrlochkopf liefern muss.\nWARUM?  Zwischen Kopf und Kaverne steht eine 1200 m hohe Säule aus dem "
             "CO₂-Gemisch. Ihr Gewicht erhöht den Druck nach unten: dp/dz = ρ(p,T) · g.\nWIE WIRD GERECHNET?  "
             "Die Dichte ρ hängt selbst von Druck und Temperatur ab – deshalb wird die Säule in 600 Schritten à "
             "2 m gerechnet, in jedem Schritt mit neuer Dichte aus CoolProp (Gemisch). Temperatur = "
             "Gebirgstemperatur T(z) = T₀ + grad T · z (Stillstand). Gerechnet wird von unten (bekannter "
             "Kavernendruck) nach oben – das Ergebnis ist direkt der Kopfdruck.\nGRENZFALL:  Mit dem "
             f"Sekantenverfahren wird der Kavernendruck gesucht, bei dem der Kopfdruck genau die Cricondenbar "
             f"des Gemischs ist ({de(p_cb)} bar = höchster Druck des Zweiphasengebiets; bei reinem CO₂ war es "
             f"p_krit = 73,8 bar). Darüber bleibt das ganze Bohrloch einphasig, darunter kann sich im oberen "
             f"Bohrloch eine Gasphase bilden.\nERGEBNIS:  Für 210 bar unten sind {de(p_kopf)} bar am Kopf nötig – "
             f"die Säule liefert rund {210 - p_kopf:.0f} bar (reines CO₂: 107,6 bar bzw. 102 bar). Das Gemisch "
             f"ist leichter, deshalb muss die Anlage oben mehr Druck aufbringen.")
for zeile, name in ((18, "voll"), (19, "grenz"), (20, "leer")):
    e = K[name]
    k.zahl(f"C{zeile}", e["p_unten"])
    k.zahl(f"D{zeile}", e["p_kopf"])
    k.zahl(f"E{zeile}", e["T_kopf"])
    k.zahl(f"F{zeile}", e["rho_kopf"])
    k.zahl(f"G{zeile}", e["rho_unten"])
k.text("I19", f"Kopfdruck = Cricondenbar ({de(p_cb)} bar): ab hier ist das Bohrloch ganz einphasig")
k.text("I20", f"Kopf {de(K['leer']['p_kopf'])} bar → zweiphasig 0–{K['leer']['tief_2ph']:.0f} m "
              f"(reines CO₂: 45,0 bar, 0–236 m)")
stil_normal, stil_schraffur = k.stil("I49"), k.stil("I50")
for i, zz in enumerate(range(0, 1201, 50)):
    r = 49 + i
    T, pv, rv, _ = K["voll"]["profil"][zz]
    _, pgz, rg, _ = K["grenz"]["profil"][zz]
    _, pl, rl, zwei = K["leer"]["profil"][zz]
    k.zahl(f"C{r}", T)
    k.zahl(f"D{r}", pv), k.zahl(f"E{r}", rv)
    k.zahl(f"F{r}", pgz), k.zahl(f"G{r}", rg)
    stil = stil_schraffur if zwei else stil_normal
    k.zahl(f"H{r}", pl, stil=stil), k.zahl(f"I{r}", rl, stil=stil)
k.text("B74", f"Schraffierte Felder: Zweiphasengebiet des Gemischs (Q-028: „multiphase region“). Beim leeren "
              f"Fall liegt das Bohrloch von 0 bis ca. {K['leer']['tief_2ph']:.0f} m zwischen Tau- und Blasenlinie. "
              f"Anders als bei reinem CO₂ ist der Druck dort nicht auf eine Linie festgelegt – ein Gemisch "
              f"verdampft über ein Druckband. Die Dichte ist die Gleichgewichtsdichte bei homogen verteilten "
              f"Phasen; im Stillstand würden sich Gas und Flüssigkeit trennen, die Werte sind dort Richtwerte.")
k.speichere()

# =============================================================================
# 3) Medienvergleich
# =============================================================================
m = B["Medien"]
rho_voll = [K["voll"]["profil"][zz][2] for zz in range(0, 1201, 50)]
m.text("B4", "IDEE:  Je schwerer die Säule im Bohrloch, desto mehr Druck liefert sie und desto weniger muss die "
             f"Anlage oben aufbringen.\nDas CO₂-Gemisch ist unter diesen Bedingungen dicht ({min(rho_voll):.0f}–"
             f"{max(rho_voll):.0f} kg/m³; reines CO₂ 820–930 kg/m³), Erdgas und Wasserstoff bleiben Gase (Erdgas "
             f"~150 kg/m³, Wasserstoff ~15 kg/m³).\nFOLGE:  Für das CO₂-Gemisch reicht eine Pumpe auf "
             f"~{p_kopf:.0f} bar (reines CO₂ ~108 bar). Erdgas- und Wasserstoffspeicher brauchen Verdichter fast "
             "bis zum vollen Kavernendruck.\nRechnung wie auf Blatt „Kaverne & Gassäule“: dp/dz = ρ(p,T) · g, "
             "von unten nach oben, Dichte aus CoolProp (CO₂: Worst-Case-Gemisch, Methan als Erdgas-Ersatz, "
             "reiner Wasserstoff). Der Zieldruck unten bleibt für alle Teufen bewusst bei 210 bar, damit nur "
             "der Effekt der Säule sichtbar wird.")
m.text("C7", "CO₂-Gemisch [bar]")
for i, zz in enumerate((600, 800, 1000, 1200, 1400, 1600)):
    m.zahl(f"C{8 + i}", W["mv_kopf"][zz])
m.text("B14", f"CO₂-Gemisch: {W['mv_kopf'][600] - W['mv_kopf'][1600]:.0f} bar weniger Kopfdruck zwischen 600 "
              "und 1600 m (reines CO₂: 74 bar) – Erdgas und Wasserstoff profitieren kaum von der Teufe.")
m.text("B18", "CO₂-Gemisch")
m.zahl("C18", p_kopf)
m.zahl("E18", K["voll"]["rho_kopf"])
m.zahl("F18", K["voll"]["rho_unten"])
m.text("G18", f"Pumpe (91 → {p_kopf:.0f} bar)")
m.text("C24", "p CO₂-Gemisch [bar]")
m.text("D24", "ρ CO₂-Gemisch [kg/m³]")
for i, zz in enumerate(range(0, 1201, 50)):
    _, pv, rv, _ = K["voll"]["profil"][zz]
    m.zahl(f"C{25 + i}", pv), m.zahl(f"D{25 + i}", rv)
m.speichere()

# =============================================================================
# 4) Einspeicherung S1
# =============================================================================
s1 = B["S1"]
s1.text("B1", "Einspeicherung S1 – Anlieferung in dichter Phase: Pumpe (Worst-Case-Gemisch)")
a_bl15 = A["p_S1"] - gw.blasendruck(A["T_S1"])
s1.text("B4", "AUSGANGSLAGE:  Das CO₂-Gemisch kommt dicht (flüssigkeitsähnlich) aus der Fernleitung: 91 bar, "
              f"15 °C. Es liegt {de(a_bl15, 0)} bar über der Blasenlinie und {de(A['p_S1'] - p_cb)} bar über der "
              f"Cricondenbar – sicher außerhalb der Unsicherheit der Phasengrenze (±3 bar).\nZIEL:  {de(p_kopf)} bar "
              "am Bohrlochkopf – vorgegeben durch die Kaverne (Blatt „Kaverne & Gassäule“).\nWARUM EINE PUMPE?  "
              "Das Gemisch ist bereits dicht und kaum kompressibel. Eine Pumpe genügt, ein Verdichter wäre nur für "
              "Gas nötig.\nRECHENWEG:  (1) Zustand am Eintritt → Enthalpie h₁, Entropie s₁.  (2) Idealfall: "
              "isentrope Druckerhöhung (Entropie bleibt gleich) → h₂s.  (3) Realfall mit Verlusten über den "
              "isentropen Wirkungsgrad: h₂ = h₁ + (h₂s − h₁) / η.  (4) Aus p₂ und h₂ → Austrittstemperatur T₂.  "
              "(5) Spezifische Arbeit w = h₂ − h₁, Leistung P = ṁ · w.\nBEGRIFF:  isentrop = Entropie konstant "
              "(Pumpe, Verdichter). Nicht verwechseln mit isenthalp = Enthalpie konstant (Drossel bei der "
              "Ausspeicherung).\nANNAHMEN:  η = 0,80 (eigene Annahme, vgl. Q-016: Verdichter 84–76 %, "
              "Glykolpumpe 70 %); kein Druckverlust in Messtechnik und Filtration; Worst-Case-Gemisch (Enthalpien "
              "haben einen anderen Nullpunkt als bei reinem CO₂ – vergleichbar sind nur Differenzen).")
s1.zahl("E10", S1["h1"]), s1.zahl("E11", S1["s1"]), s1.zahl("E13", S1["h2s"]), s1.zahl("E16", S1["T2"])
s1.text("E17", S1["phase2"])
s1.text("C19", "Normdichte Gemisch (0 °C; 1,01325 bar)")
s1.zahl("E19", rho_n)
s1.zahl("F28", S1["h1"]), s1.text("G28", S1["phase1"])
s1.zahl("D29", p_kopf), s1.zahl("E29", S1["T2"]), s1.zahl("F29", S1["h2"]), s1.text("G29", S1["phase2"])
s1.speichere()

# =============================================================================
# 5) Einspeicherung S2
# =============================================================================
s2 = B["S2"]
s2.text("B1", "Einspeicherung S2 – gasförmige Anlieferung: verdichten, dicht kühlen, pumpen (Worst-Case-Gemisch)")
s2.text("B4", "AUSGANGSLAGE:  Das CO₂-Gemisch kommt gasförmig an: 30 bar, 15 °C (Annahme).\nQ-016:  Bielka, "
              "Kuczyński, Nagy (2023): CO₂ Compression and Dehydration for Transport and Geological Storage. "
              "Energies 16, 1804 (AGH Krakau). Übersichtsarbeit zu Abscheidung, Trocknung, Transport und "
              "Speicherstandortwahl – plus eine eigene Prozesssimulation (BR&E ProMax, Peng-Robinson): CO₂ aus "
              "einer Rauchgasabscheidung (1,51 bar, 35 °C, 2,45 Mio. t/a ≈ 78 kg/s) wird in 5 Stufen auf 100 bar "
              "verdichtet, mit TEG getrocknet und auf 20 °C gekühlt – flüssig für den Pipelinetransport zu einer "
              "Speicherstätte in 30 km Entfernung (S. 5: 74 bar sind das Minimum, üblich 100 bar als Reserve für "
              "Druckverluste).\nAnnahmen aus Q 16 (S. 4, Kap. 2.3):  Kühlung auf 20 °C nach jeder Stufe · höchstens "
              "95 °C nach jeder Stufe · Wirkungsgrad 84 % in Stufe 1, je weitere Stufe −2 % · zum Schluss Kühlung "
              f"auf 20 °C → flüssig. Q-016 verdichtet für die Pipeline, hier für den Bohrlochkopf – das Druckniveau "
              f"(100 bzw. {de(p_kopf)} bar) und der Massenstrom (78 bzw. 27–54 kg/s) sind vergleichbar. NICHT "
              "übernommen: 45 bar Zwischendruck (für die TEG-Trocknung gewählt, Glykol schließt OGE aus).\n"
              f"NEUE VERDICHTUNG BEIM GEMISCH:  Reines CO₂ wird bei 60 bar und 20 °C verflüssigt (Dampfdruck "
              f"57,3 bar). Das Gemisch hat bei 20 °C ein Zweiphasengebiet von {de(S2['p_tau_TK'])} bar (Taulinie) bis "
              f"{de(S2['p_bl_TK'])} bar (Blasenlinie) – bei 60 bar bliebe es gasförmig (alte Verdichtung, Blatt "
              f"„Gemisch & Grenzen“). Deshalb wird über die Cricondenbar ({de(p_cb)} bar = höchster Druck des "
              f"Zweiphasengebiets) + 3 bar Unsicherheit verdichtet: p_V = {de(A['p_V'], 0)} bar. Dort kühlt der "
              "Kühler von der Verdichteraustrittstemperatur auf 20 °C, ohne das Zweiphasengebiet zu berühren – das "
              "Gemisch geht stetig in die dichte Phase über, es wird nichts kondensiert (wie Q-016: Endstufe 100 bar, "
              f"dann 20 °C). Danach bringt die Pumpe das dichte Gemisch auf {de(p_kopf)} bar.\nZAHL DER STUFEN:  Eine "
              f"Stufe bis {de(A['p_V'], 0)} bar erreicht {de(S2['T1a'])} °C und überschreitet die 95-°C-Grenze aus Q-016 "
              "(reines CO₂ bis 60 bar: 73,2 °C). Die zweistufige Variante unten ist deshalb Pflicht – sie spart außerdem "
              f"{ersparnis_v:.0%} Arbeit.\nRECHENWEG:  Verdichter und Pumpe wie in S1: isentrop + Wirkungsgrad, "
              "h_aus = h_ein + (h_s − h_ein) / η. Kühler: abgeführte Wärme je kg q = h_vor − h_nach.\n"
              f"GRENZEN:  20 °C setzen Kühlwasser voraus. Mit Luftkühlung (ca. 35 °C) läge die Kühltemperatur "
              f"über der kritischen Temperatur des Gemischs ({de(T_kG)} °C): das Gemisch wäre dann vor der Pumpe "
              "überkritisch und deutlich weniger dicht.")
s2.zahl("E12", S2["h1"]), s2.zahl("E14", S2["h1s"]), s2.zahl("E16", S2["T1a"])
s2.text("B2", "Netzübergabe → Messtechnik/Filtration → Verdichter → Kühler (dicht) → Pumpe → Bohrlochkopf")
s2.text("B19", "Kühler (dichte Phase)")
s2.text("C11", "Austrittsdruck = Kühlerdruck p_V")
s2.text("C20", "Cricondenbar Gemisch")
s2.text("D20", "CoolProp: höchster Druck der Phasengrenze")
s2.zahl("E20", p_cb)
s2.text("C21", "Abstand p_V − Cricondenbar")
s2.text("D21", "muss ≥ 3 bar sein, dann kein Zweiphasengebiet beim Kühlen")
s2.text("C22", "Blasendruck bei T_K")
s2.text("D22", f"CoolProp: p(Q = 0, T_K) – Taudruck {de(S2['p_tau_TK'])} bar")
s2.zahl("E22", S2["p_bl_TK"])
s2.text("C23", "Enthalpie nach Kühler")
s2.text("C24", "Phase nach Kühler")
s2.zahl("E23", S2["hV"]), s2.text("E24", S2["phaseV"])
s2.text("C36", "Kühlleistung Kühler, oberer Fall")
s2.zahl("E28", S2["hPs"]), s2.zahl("E30", S2["TPa"])
s2.text("I24", f"S2 (2 Stufen) braucht rund {w_S2v / S1['w']:.0f}-mal so viel Arbeit wie S1 – und rund {vorteil_S2:.0%} "
               f"weniger als Durchverdichten. Beim Gemisch ist die zweite Verdichterstufe Pflicht "
               f"(1 Stufe: {de(S2['T1a'])} °C > 95 °C) und spart gegenüber einer Stufe {ersparnis_v:.0%}.")
s2.text("B40", f"AUFBAU:  Stufe 1 30 → {de(S2v['p_zw'])} bar, Zwischenkühlung auf 20 °C (das Gemisch bleibt "
               f"gasförmig: Taudruck bei 20 °C = {de(S2['p_tau_TK'])} bar), Stufe 2 → {de(A['p_V'], 0)} bar. Zwischendruck mit "
               "gleichem Druckverhältnis in beiden Stufen: p_zw = √(p_ein · p_V). Wirkungsgrade 84 / 82 % (Q-016). "
               "Kühler und Pumpe wie im Basisfall. Beim Gemisch ist diese Variante Pflicht (95-°C-Grenze).")
s2.text("B39", "Variante: zweistufige Verdichtung bis zum Kühlerdruck p_V (sonst wie Basisfall)")
s2.text("B58", "Kühler und Pumpe wie Basisfall")
s2.text("C59", "Abgeführte Wärme Kühler")
s2.zahl("E45", S2v["h1s"]), s2.zahl("E47", S2v["T1a"]), s2.zahl("E50", S2v["hZK"])
s2.zahl("E54", S2v["h2s"]), s2.zahl("E56", S2v["T2a"])
s2.text("B68", f"FRAGE:  Warum nicht einfach gasförmig bis {de(p_kopf)} bar durchverdichten, statt dicht zu kühlen "
               "und zu pumpen?\nAUFBAU nach denselben Regeln aus Q-016:  Zwischenkühlung auf 20 °C, höchstens "
               "95 °C je Stufe, Wirkungsgrade 84 / 82 %. Der Zwischendruck muss mit 3 bar Abstand unter dem "
               f"Taudruck des Gemischs bei 20 °C bleiben ({de(S2['p_tau_TK'])} − 3 = {de(Dv['p_zw_max'])} bar; "
               "reines CO₂: 57,3 bar), sonst beginnt die Kondensation schon im Zwischenkühler; und die zweite "
               f"Stufe darf 95 °C nicht überschreiten (ab {de(Dv['p_zw_min'])} bar). Daraus folgt ein zulässiger "
               f"Bereich von ca. {Dv['p_zw_min']:.0f}–{Dv['p_zw_max']:.0f} bar – gewählt wie bei reinem CO₂ "
               "50 bar.\nFAIRER VERGLEICH:  Am Ende kühlt ein Nachkühler auf denselben Endzustand wie S2 "
               "(gleicher Druck, gleiche Temperatur).")
s2.zahl("E75", Dv["h1s"]), s2.zahl("E77", Dv["T1a"]), s2.zahl("E81", Dv["hZK"])
s2.zahl("E86", Dv["h2s"]), s2.zahl("E88", Dv["T2a"]), s2.zahl("E93", Dv["hN"])
s2.text("D80", "p_zw ≤ p_Tau(T_K) − 3 bar ?")
s2.formel("E80", 'IF(p_zw<=p_tau_TK-U_PG,"✔ gasförmig (≥ 3 bar unter der Taulinie)",'
                 '"⚠ zu nah an der Taulinie – Zwischendruck senken")', ist_text=True)
s2.text("E102", "S2 Basisfall: 1 Stufe (> 95 °C!) + Kühler + Pumpe")
s2.text("F102", "S2 Variante: 2 Stufen + Kühler + Pumpe")
s2.formel("B109", '"Ergebnis: Die zweite Verdichterstufe ist beim Gemisch Pflicht (1 Stufe > 95 °C) und spart "'
                  '&TEXT(1-Szen2v_w_ges/Szen2_w_ges,"0%")&" gegenüber einer Stufe. Dicht kühlen und pumpen (2 Stufen) '
                  'braucht "&TEXT(1-Szen2v_w_ges/Dv_w_ges,"0%")&" weniger Arbeit als Durchverdichten."', ist_text=True)
s2.speichere()

# =============================================================================
# 6) Phasendiagramm (Texte)
# =============================================================================
ph = B["Phasendiagramm"]
ph.text("B1", "Phasendiagramm – Prozesspfade der Einspeicherung und Betriebsfenster von Bohrloch und Kaverne "
              "(Worst-Case-Gemisch)")
ph.text("B2", "Worst-Case-Gemisch (95 % CO₂ + 5 % Begleitstoffe) · Stoffdaten CoolProp (GERG-2008-Mischungsregel) "
              "· Fest-Gebiet, Sublimations- und Schmelzlinie: Näherung reines CO₂ (Span & Wagner 1996)")
ph.text("B36", "DIAGRAMM 1:  Farbflächen = Phasengebiete des Gemischs, rot = Zweiphasengebiet (reines CO₂ hat nur "
               "eine Linie, grau zum Vergleich). Rot gestrichelt = ±3 bar Unsicherheit der Phasengrenze. Grün = S1 "
               f"(Pumpe). Blau = S2 Basisfall (1 Stufe bis {de(A['p_V'], 0)} bar, {de(S2['T1a'])} °C > 95 °C, Kühler dicht auf "
               "20 °C, Pumpe). Grau gestrichelt = alte Verdichtung wie reines CO₂ (60 bar) – bleibt beim Gemisch "
               "gasförmig, die Pumpe könnte nicht fördern. Hellblau gepunktet = "
               "S2-Variante mit 2 Stufen (beim Gemisch Pflicht). Violett gestrichelt = Durchverdichten (2 Stufen bis "
               f"{de(p_kopf)} bar, dann Nachkühler).\nENDPUNKT:  Alle Pfade enden am Bohrlochkopf bei {de(p_kopf)} bar "
               "(reines CO₂: 107,6 bar). Diagramm 2 zeigt Bohrloch und Kaverne – nur im Stillstand.")
ph.text("B72", "DIAGRAMM 2:  Jede gestrichelte Linie ist eine Gassäule – ein Füllstand der Kaverne im Stillstand. Sie "
               "läuft vom Bohrlochkopf (grauer Balken, 0 m, 10 °C) bis zur Kaverne (dunkelblauer Balken, 1200 m, "
               "46 °C). Jeder Punkt auf der Linie ist eine Teufe: oben kalt und niedriger Druck, unten warm und "
               "hoher Druck.\nBETRIEBSFENSTER IM STILLSTAND:  Der Bereich zwischen der Linie „voll“ (oben) und "
               "der Linie „leer“ (unten), links begrenzt durch den Kopf-Balken und rechts durch den "
               "Kavernen-Balken.\nWICHTIG BEI LEERER KAVERNE:  70 bar unten ergeben oben "
               f"{de(K['leer']['p_kopf'])} bar (reines CO₂: 45,0 bar). Das obere Bohrloch "
               f"(0–{K['leer']['tief_2ph']:.0f} m) liegt dann im Zweiphasengebiet des Gemischs. Erst ab ca. "
               f"{de(K['grenz']['p_unten'])} bar Kavernendruck (Grenzfall: Kopf = Cricondenbar {de(p_cb)} bar) "
               "bleibt das ganze Bohrloch einphasig.\nKAVERNE SELBST:  Bei 46 °C liegt sie immer über der "
               f"Cricondentherm des Gemischs ({de(pg['cricondentherm'][0])} °C) – dort gibt es keine Flüssigphase. "
               f"Über {de(p_kG)} bar ist das Gemisch überkritisch, bei 70 bar gasförmig.\nSTILLSTAND VS. BETRIEB:  "
               "Die Gassäulen sind mit Gebirgstemperatur gerechnet (Kopf 10 °C) und gelten nur, solange nichts "
               f"fließt. Beim Einspeichern kommt das Gemisch wärmer aus der Pumpe (S1 {de(S1['T2'])} °C, S2 "
               f"{de(S2['TPa'])} °C).")
ph.speichere()

# =============================================================================
# 7) Quellen (neue Zeilen)
# =============================================================================
q = B["Quellen"]
st_id, st_q, st_v = q.stil("B15"), q.stil("C15"), q.stil("D15")
neue_quellen = [
    ("—", "Kunz, O.; Wagner, W. (2012): The GERG-2008 Wide-Range Equation of State for Natural Gases and Other "
          "Mixtures: An Expansion of GERG-2004. J. Chem. Eng. Data 57, 3032–3091.",
     "Mischungsregel (Mehrfluid-Helmholtz-Modell) für das Worst-Case-Gemisch; Paarparameter CO₂–CH₄, CO₂–H₂"),
    ("—", "Gernert, J.; Span, R. (2016): EOS–CG: A Helmholtz energy mixture model for humid gases and CCS "
          "mixtures. J. Chem. Thermodyn. 93, 274–293 (Parameter aus Gernert 2013, Diss. RUB).",
     "Paarparameter CO₂–N₂, CO₂–Ar, CO₂–CO in CoolProp"),
    ("Gemisch", "Vorschlag Worst-Case-Gemisch Anlieferung: 95 % CO₂, 2,4 % N₂, 1 % Ar, 1 % CH₄, 0,5 % H₂, 0,1 % CO "
                "(mol-%); Belege DIN EN 18397, ISO 27913, Porthos-Grenzwerte, DVGW C 260.",
     "Zusammensetzung des Gemischs (Blatt „Gemisch & Grenzen“)"),
    ("Doku", "Dokumentation_Stoffmodelle_Validierung.docx (eigene Arbeit): Gleichungen, Modellvergleich mit "
             "thermopack (GERG-2008, Peng-Robinson, SRK), Unsicherheit der Phasengrenze.",
     "Unsicherheit der Phasengrenze ±3 bar; Validierung der Stoffwerte"),
]
for i, (qid, text, wofuer) in enumerate(neue_quellen):
    r = 16 + i
    q.text(f"B{r}", qid, stil=st_id), q.text(f"C{r}", text, stil=st_q), q.text(f"D{r}", wofuer, stil=st_v)
q.ersetze('<dimension ref="B1:D15"/>', '<dimension ref="B1:D19"/>')
q.speichere()

# =============================================================================
# 8) Diagrammdaten
# =============================================================================
d = B["Diagrammdaten"]
d.text("B2", "Grundlage für das Phasendiagramm. Farbflächen: gestapelte Druckbereiche je Temperatur (Gemisch: "
             "gasförmig bis Taulinie, zweiphasig bis Blasenlinie, darüber flüssig; oberhalb des kritischen Punkts "
             "des Gemischs überkritisch; fest als Näherung reines CO₂). Linien: Punkte mit T, p und Position im "
             "Diagramm.")
st_kopf = d.stil("G7")
d.text("H7", "zweiphasig [bar]", stil=st_kopf)
# Temperaturachse bis 120 °C verlängern (S2 Basisfall 1 Stufe erreicht 111,6 °C)
for i, (T, gas, zwei, flue, sup, fest) in enumerate(dia["flaechen"]):
    r = 8 + i
    if r > 368:
        d.zahl(f"B{r}", T)
        d.formel(f"C{r}", f'IF(MOD(B{r},20)=0,TEXT(B{r},"0"),"")', ist_text=True)
    d.zahl(f"D{r}", gas), d.zahl(f"E{r}", flue), d.zahl(f"F{r}", sup), d.zahl(f"G{r}", fest)
    d.zahl(f"H{r}", zwei)
letzte = 8 + len(dia["flaechen"]) - 1
assert letzte == 408, letzte
d.ersetze('<col min="8" max="8" width="2.7109375" customWidth="1"/>',
          '<col min="8" max="8" width="12.7109375" customWidth="1"/>')
# Phasengrenze Gemisch statt Sättigungslinie (Zeilen 8-77)
d.text("I8", "Phasengrenze Gemisch (Tau-/Blasenlinie)")
for i, (T, p) in enumerate(dia["huelle"]):
    d.zahl(f"J{8 + i}", T), d.zahl(f"K{8 + i}", p)
d.text("I123", "Kritischer Punkt Gemisch")
d.text("I128", "S2: verdichten → dicht kühlen → pumpen")
d.zahl("J123", T_kG), d.zahl("K123", p_kG)
# Beschriftung Phasengebiete E (+ neu: zweiphasig in Zeile 148)
for r, (T, p) in zip(range(144, 149), [(-71, 70), (-30, 105), (72, 12), (85, 120), (-22, 34)]):
    d.zahl(f"J{r}", T), d.zahl(f"K{r}", p)
d.formel("L148", "(J148-Diagramm_Tmin)/Diagramm_dT+1")
# Beschriftung Punkte E
pos_E = [(1, 91), (15, 25), (S2["T1a"] - 12, A["p_V"] + 5), (45, A["p_V"] + 4), (23, p_kopf + 6),
         (Dv["T1a"] + 9, A["p_zw"] - 4), (36, A["p_zw"] + 3), (Dv["T2a"] + 9, p_kopf - 4)]
for r, (T, p) in zip(range(149, 157), pos_E):
    d.zahl(f"J{r}", T), d.zahl(f"K{r}", p)
# Beschriftungen Diagramm 2
d.text("I184", "Gassäule Grenzfall (Kopf = Cricondenbar)")
d.text("I239", f"Bohrlochkopf (0 m): {de(K['leer']['p_kopf'])}–{de(p_kopf)} bar")
for r, (T, p) in zip(range(242, 247), [(-71, 150), (-12, 170), (78, 25), (78, 160), (-25, 32)]):
    d.zahl(f"J{r}", T), d.zahl(f"K{r}", p)
d.formel("L246", "(J246-Diagramm_Tmin)/Diagramm_dT+1")
pos_K = [(60, 216), (60, 172), (58, 64), (-6, p_kopf + 4), (-8, p_cb + 1), (-11, K["leer"]["p_kopf"] - 6), (29, 100)]
for r, (T, p) in zip(range(247, 254), pos_K):
    d.zahl(f"J{r}", T), d.zahl(f"K{r}", p)
# neue Linien: reines CO2, sicher flüssig (+3 bar), sicher gasförmig (-3 bar)
ALT = W["alt"]
NEUE_LINIEN = [("Sättigungslinie reines CO₂ (Vergleich)", dia["rein"], 256),
               ("Blasenlinie + 3 bar (sicher flüssig)", dia["sicher_fl"], 298),
               ("Taulinie − 3 bar (sicher gasförmig)", dia["sicher_gas"], 329),
               ("Alte Verdichtung (60 bar wie reines CO₂) – bleibt gasförmig",
                [(A["T_S2"], A["p_S2"]), (ALT["T1a"], ALT["p_V"]), (A["T_K"], ALT["p_V"])], 360)]
st_I = d.stil("I8")
for name, pts, r0 in NEUE_LINIEN:
    d.text(f"I{r0}", name, stil=st_I)
    for i, (T, p) in enumerate(pts):
        r = r0 + i
        d.zahl(f"J{r}", T), d.zahl(f"K{r}", p)
        d.formel(f"L{r}", f"(J{r}-Diagramm_Tmin)/Diagramm_dT+1")
d.ersetze('<dimension ref="B1:L368"/>', '<dimension ref="B1:L408"/>')
d.speichere()

# =============================================================================
# 9) Diagramme 8 und 9 (Phasendiagramme)
# =============================================================================
FARBE_ZWEI = "F4B6B0"


def chart_bearbeiten(name, labels_punkte, mit_bereich_label_zeile, titel_zusatz, linien):
    xml = pk.text(name)
    # Bereiche bis Zeile 408 verlängern (120 °C)
    xml = xml.replace("$368", "$408")
    # Zahlen-Caches entfernen (Excel liest die Zellen beim Öffnen neu)
    xml = re.sub(r"<c:numCache>.*?</c:numCache>", "", xml, flags=re.S)
    xml = re.sub(r"<c:strCache>.*?</c:strCache>", "", xml, flags=re.S)
    # Reihenfolge: nach der Gasfläche eine neue Fläche "zweiphasig" einfügen
    xml = re.sub(r'<c:order val="(\d+)"/>',
                 lambda m_: f'<c:order val="{int(m_.group(1)) + (1 if int(m_.group(1)) >= 1 else 0)}"/>', xml)
    area = re.search(r"<c:areaChart>.*?</c:areaChart>", xml, re.S).group(0)
    sers = re.findall(r"<c:ser>.*?</c:ser>", area, re.S)
    vorlage = [s_ for s_ in sers if "$E$7" in s_][0]
    neu = vorlage.replace('<c:idx val="1"/>', '<c:idx val="40"/>').replace('<c:order val="2"/>', '<c:order val="1"/>')
    neu = neu.replace("$E$", "$H$")
    neu = re.sub(r'(<c:spPr>.*?<a:srgbClr val=")[0-9A-Fa-f]{6}', rf"\g<1>{FARBE_ZWEI}", neu, count=1, flags=re.S)
    area_neu = area.replace(sers[0], sers[0] + neu, 1)
    xml = xml.replace(area, area_neu)
    # Linie 4: Sättigungslinie -> Phasengrenze Gemisch
    xml = xml.replace("<c:v>Sättigungslinie</c:v>", "<c:v>Phasengrenze Gemisch (Tau-/Blasenlinie)</c:v>")
    xml = xml.replace("<c:v>Kritischer Punkt</c:v>", "<c:v>Kritischer Punkt Gemisch</c:v>")
    # neue Linien: Kopie der Sättigungslinie mit neuen Bereichen und Stil
    scat = re.search(r"<c:scatterChart>.*?</c:scatterChart>", xml, re.S).group(0)
    s_sat = [s_ for s_ in re.findall(r"<c:ser>.*?</c:ser>", scat, re.S) if "$L$8:$L$77" in s_][0]
    zus = ""
    for j, (lname, pts, r0) in enumerate(linien):
        r1 = r0 + len(pts) - 1
        s_ = s_sat.replace("$L$8:$L$77", f"$L${r0}:$L${r1}").replace("$K$8:$K$77", f"$K${r0}:$K${r1}")
        s_ = re.sub(r'<c:idx val="\d+"/>', f'<c:idx val="{41 + j}"/>', s_, count=1)
        s_ = re.sub(r'<c:order val="\d+"/>', f'<c:order val="{41 + j}"/>', s_, count=1)
        s_ = s_.replace("<c:v>Phasengrenze Gemisch (Tau-/Blasenlinie)</c:v>", f"<c:v>{escape(lname)}</c:v>")
        farbe, breite, strich = [("808080", "12700", None), ("CA220E", "12700", "dash"),
                                 ("CA220E", "12700", "dash"), ("7F7F7F", "22225", "sysDash")][j]
        s_ = re.sub(r"<c:spPr>.*?</c:spPr>",
                    f'<c:spPr><a:ln w="{breite}" cap="rnd"><a:solidFill><a:srgbClr val="{farbe}"/></a:solidFill>'
                    + (f'<a:prstDash val="{strich}"/>' if strich else "") + '<a:round/></a:ln></c:spPr>',
                    s_, count=1, flags=re.S)
        zus += s_
    scat_neu = scat.replace(s_sat, s_sat + zus, 1)
    xml = xml.replace(scat, scat_neu)
    # Phasengrenze Gemisch in Rot
    s_sat_neu = re.sub(r'(<c:spPr>.*?<a:srgbClr val=")[0-9A-Fa-f]{6}', r"\g<1>CA220E", s_sat, count=1, flags=re.S)
    xml = xml.replace(s_sat, s_sat_neu, 1)
    # Beschriftungen
    for alt, neu_t in labels_punkte:
        xml = xml.replace(f"<a:t>{escape(alt)}</a:t>", f"<a:t>{escape(neu_t)}</a:t>")
    # Bereichsbeschriftung "zweiphasig" (fünfter Punkt)
    r_lab = mit_bereich_label_zeile
    s_lab = [s_ for s_ in re.findall(r"<c:ser>.*?</c:ser>", xml, re.S) if "<a:t>überkritisch</a:t>" in s_][0]
    d_ueber = [dl for dl in re.findall(r"<c:dLbl>.*?</c:dLbl>", s_lab, re.S) if "überkritisch" in dl][0]
    d_zwei = d_ueber.replace('<c:idx val="3"/>', '<c:idx val="4"/>').replace("<a:t>überkritisch</a:t>",
                                                                            "<a:t>zweiphasig</a:t>")
    d_zwei = re.sub(r'(<a:solidFill><a:srgbClr val=")[0-9A-Fa-f]{6}', r"\g<1>CA220E", d_zwei)
    s_lab_neu = s_lab.replace(d_ueber, d_ueber + d_zwei)
    s_lab_neu = re.sub(r"\$(L|K)\$(\d+):\$(L|K)\$(\d+)",
                       lambda m_: f"${m_.group(1)}${m_.group(2)}:${m_.group(3)}${r_lab}", s_lab_neu)
    xml = xml.replace(s_lab, s_lab_neu)
    # Reihen durchnummerieren (idx = order = Position) und Legende neu aufbauen:
    # Flächen und Beschriftungsreihen erscheinen nicht in der Legende
    reihen = re.findall(r"<c:ser>.*?</c:ser>", xml, re.S)
    ausblenden = []
    for i, r_ in enumerate(reihen):
        r_neu = re.sub(r'<c:idx val="\d+"/>', f'<c:idx val="{i}"/>', r_, count=1)
        r_neu = re.sub(r'<c:order val="\d+"/>', f'<c:order val="{i}"/>', r_neu, count=1)
        xml = xml.replace(r_, r_neu, 1)
        if i < 5 or "Beschriftung" in r_:
            ausblenden.append(i)
    leg = re.search(r"<c:legend>.*?</c:legend>", xml, re.S).group(0)
    leg_neu = re.sub(r"<c:legendEntry>.*?</c:legendEntry>", "", leg, flags=re.S)
    eintraege = "".join(f'<c:legendEntry><c:idx val="{i}"/><c:delete val="1"/></c:legendEntry>' for i in ausblenden)
    leg_neu = re.sub(r"(<c:legendPos[^>]*/>)", lambda m_: m_.group(1) + eintraege, leg_neu, count=1)
    xml = xml.replace(leg, leg_neu)
    # Titel
    titel = re.search(r"<c:title>.*?</c:title>", xml, re.S).group(0)
    t_alt = re.findall(r"<a:t>([^<]*)</a:t>", titel)[0]
    xml = xml.replace(f"<a:t>{t_alt}</a:t>", f"<a:t>{t_alt}{titel_zusatz}</a:t>", 1)
    # kopierte Reihen/Beschriftungen brauchen eigene Kennungen (sonst "Reparatur" in Excel)
    gesehen = set()

    def neue_id(m_):
        v = m_.group(1)
        if v in gesehen:
            v = "{" + str(uuid.uuid4()).upper() + "}"
        gesehen.add(v)
        return f'uniqueId val="{v}"'
    xml = re.sub(r'uniqueId val="([^"]*)"', neue_id, xml)
    pk.setze(name, xml)


chart_bearbeiten("xl/charts/chart8.xml",
                 [("Bohrlochkopf 107,6 bar", f"Bohrlochkopf {de(p_kopf)} bar"),
                  ("Verdichter", "Verdichter 1 Stufe"), ("Verflüssiger", "Kühler (dicht)")],
                 148, " – Worst-Case-Gemisch", NEUE_LINIEN)
chart_bearbeiten("xl/charts/chart9.xml",
                 [("Grenzfall 171,7 bar", f"Grenzfall {de(K['grenz']['p_unten'])} bar"),
                  ("Kopf 107,6 bar", f"Kopf {de(p_kopf)} bar"),
                  ("Kopf 73,8 bar = p_krit", f"Kopf {de(p_cb)} bar = Cricondenbar"),
                  ("Kopf 45 bar = p_sat(10 °C)", f"Kopf {de(K['leer']['p_kopf'])} bar – zweiphasig"),
                  ("Gassäule Grenzfall (Kopf = p_krit)", "Gassäule Grenzfall (Kopf = Cricondenbar)"),
                  ("Bohrlochkopf (0 m): 45–107,6 bar",
                   f"Bohrlochkopf (0 m): {de(K['leer']['p_kopf'])}–{de(p_kopf)} bar")],
                 246, " – Worst-Case-Gemisch", NEUE_LINIEN[:3])
# Serienname Diagramm 9 steht als Literal im Diagramm
x9 = pk.text("xl/charts/chart9.xml")
x9 = x9.replace("<c:v>Gassäule Grenzfall (Kopf = p_krit)</c:v>", "<c:v>Gassäule Grenzfall (Kopf = Cricondenbar)</c:v>")
x9 = x9.replace("<c:v>Bohrlochkopf (0 m): 45–107,6 bar</c:v>",
                f"<c:v>Bohrlochkopf (0 m): {de(K['leer']['p_kopf'])}–{de(p_kopf)} bar</c:v>")
pk.setze("xl/charts/chart9.xml", x9)
x8 = pk.text("xl/charts/chart8.xml")
pk.setze("xl/charts/chart8.xml", x8.replace("S2: verdichten → verflüssigen → pumpen", "S2: verdichten → dicht kühlen → pumpen"))
x7 = pk.text("xl/charts/chart7.xml")
x7 = x7.replace("S2 Basisfall: 1 Stufe + Verflüssiger + Pumpe", "S2 Basisfall: 1 Stufe (&gt; 95 °C!) + Kühler + Pumpe")
x7 = x7.replace("S2 Variante: 2 Stufen + Verflüssiger + Pumpe", "S2 Variante: 2 Stufen + Kühler + Pumpe")
pk.setze("xl/charts/chart7.xml", x7)
# Medienvergleich: Reihenname CO2
for c in ("xl/charts/chart4.xml", "xl/charts/chart5.xml"):
    pk.setze(c, pk.text(c).replace("<c:v>CO₂</c:v>", "<c:v>CO₂-Gemisch</c:v>"))

# =============================================================================
# 10) Prozesskette (Formen auf der Übersicht)
# =============================================================================
dr = pk.text("xl/drawings/drawing1.xml")
for alt, neu_t in [("→ 107,6 bar · 16,8 °C", f"→ {de(p_kopf)} bar · {de(S1['T2'])} °C"),
                   ("→ 60 bar · 73 °C", f"→ {A['p_V']:.0f} bar · {S2v['T2a']:.0f} °C (2 Stufen)"),
                   ("60 bar · 20 °C", f"{A['p_V']:.0f} bar · 20 °C (dicht)"),
                   ("→ 107,6 bar · 27,5 °C", f"→ {de(p_kopf)} bar · {de(S2['TPa'])} °C"),
                   ("Kopf: 107,6 bar", f"Kopf: {de(p_kopf)} bar"),
                   ("+102 bar aus der Säule", f"+{210 - p_kopf:.0f} bar aus der Säule"),
                   ("Verflüssiger", "Kühler (dicht)"),
                   ("Kühlung auf 20 °C und Stufenregel nach Q-016",
                    "Kühlung auf 20 °C · Gemisch: 2 Stufen nötig (1 Stufe > 95 °C)")]:
    assert f"<a:t>{escape(alt)}</a:t>" in dr, alt
    dr = dr.replace(f"<a:t>{escape(alt)}</a:t>", f"<a:t>{escape(neu_t)}</a:t>")
pk.setze("xl/drawings/drawing1.xml", dr)

# =============================================================================
# 11) Neues Blatt "Gemisch & Grenzen"
# =============================================================================
ST = {"titel": u.stil("B1"), "unter": u.stil("B2"), "abschnitt": u.stil("B5"), "link": u.stil("B13"),
      "text": u.stil("E13"), "kopf": u.stil("B41"), "kopf2": u.stil("E41"), "zelle": u.stil("B43"),
      "zelle_m": u.stil("E43"), "grau": u.stil("M43"), "blau": u.stil("F43"), "ergebnis": u.stil("F75"),
      "formel": u.stil("F70"), "hinweis": u.stil("B38"), "gruppe": u.stil("B42")}
zeilen = {}


def zelle(ref, wert, stil, art=None):
    sp, z = re.match(r"([A-Z]+)(\d+)", ref).groups()
    s_attr = f' s="{ST[stil]}"' if ST.get(stil) else ""
    if art == "f":
        inhalt = f'<c r="{ref}"{s_attr}><f>{escape(wert)}</f></c>'
    elif isinstance(wert, str):
        inhalt = f'<c r="{ref}"{s_attr} t="s"><v>{sst.index(wert)}</v></c>'
    elif wert is None:
        inhalt = f'<c r="{ref}"{s_attr}/>'
    else:
        inhalt = f'<c r="{ref}"{s_attr}><v>{repr(float(wert))}</v></c>'
    zeilen.setdefault(int(z), []).append((sp, inhalt))


zelle("B1", "Gemisch & Grenzen – was die Begleitstoffe am Phasendiagramm ändern", "titel")
zelle("B2", "Worst-Case-Gemisch Anlieferung · Stoffdaten CoolProp 8.0.0 (GERG-2008-Mischungsregel) · Vergleich mit "
            "der Mappe für reines CO₂", "unter")
zelle("B3", "← zur Übersicht", "link")
zelle("B5", "Kurz gesagt", "abschnitt")
zelle("B6", "Reines CO₂ hat eine Dampfdrucklinie: bei jeder Temperatur genau ein Druck, bei dem es siedet. Das Gemisch "
            "hat stattdessen ein Zweiphasengebiet zwischen Taulinie (erster Tropfen) und Blasenlinie (letzte Blase). "
            "Die leichtflüchtigen Begleitstoffe (vor allem H₂ und N₂) schieben die Blasenlinie nach oben – für "
            "„sicher flüssig“ braucht man mehr Druck. Weil die Phasengrenze eines Gemischs unsicherer ist als die "
            "von reinem CO₂, gilt ein Sicherheitsabstand von ±3 bar (Modellvergleich, Dokumentation Kap. 9).", "text")

zelle("B8", "Zusammensetzung (mol-%)", "abschnitt")
for sp, t in zip("BCD", ["Komponente", "mol-%", "Beleg"]):
    zelle(f"{sp}9", t, "kopf" if sp == "B" else "kopf2")
for i, (name, _, x, beleg) in enumerate(gw.ZUSAMMENSETZUNG):
    r = 10 + i
    zelle(f"B{r}", name, "zelle"), zelle(f"C{r}", x, "blau"), zelle(f"D{r}", beleg, "zelle")
zelle("B16", "Σ Begleitstoffe", "zelle"), zelle("C16", "SUM(C11:C15)", "formel", "f"), zelle("D16", "= Grenze DIN EN 18397", "zelle")

zelle("F8", "Kennpunkte der Phasengrenze", "abschnitt")
for sp, t in zip("FGHIJ", ["Kennpunkt", "T [°C]", "p [bar]", "reines CO₂", "Bedeutung"]):
    zelle(f"{sp}9", t, "kopf" if sp == "F" else "kopf2")
kenn = [("Kritischer Punkt", T_kG, p_kG, "31,0 °C / 73,8 bar", "Tau- und Blasenlinie treffen sich"),
        ("Cricondenbar", pg["cricondenbar"][0], p_cb, "= krit. Punkt", "darüber bei jeder Temperatur einphasig"),
        ("Cricondentherm", pg["cricondentherm"][0], pg["cricondentherm"][1], "= krit. Punkt",
         "darüber bei jedem Druck einphasig")]
for i, (n, T, p, rein, bed) in enumerate(kenn):
    r = 10 + i
    zelle(f"F{r}", n, "zelle"), zelle(f"G{r}", T, "grau"), zelle(f"H{r}", p, "grau")
    zelle(f"I{r}", rein, "zelle_m"), zelle(f"J{r}", bed, "zelle")
zelle("F14", "Unsicherheit der Phasengrenze", "zelle"), zelle("H14", U_PG, "blau")
zelle("I14", "bar", "zelle_m"), zelle("J14", "Dokumentation Kap. 9: Modellstreuung 1 bar + Modellfehler 1–3 %", "zelle")
zelle("F15", "Blasendruck bei T_K", "zelle"), zelle("G15", A["T_K"], "grau"), zelle("H15", S2["p_bl_TK"], "grau")
zelle("I15", "57,3 (p_sat)", "zelle_m"), zelle("J15", "darüber vollständig flüssig", "zelle")
zelle("F16", "Taudruck bei T_K", "zelle"), zelle("G16", A["T_K"], "grau"), zelle("H16", S2["p_tau_TK"], "grau")
zelle("I16", "57,3 (p_sat)", "zelle_m"), zelle("J16", "darunter vollständig gasförmig", "zelle")

zelle("B18", "Tau- und Blasendruck über der Temperatur", "abschnitt")
kopf = ["T [°C]", "p_sat reines CO₂ [bar]", "Taudruck Gemisch [bar]", "Blasendruck Gemisch [bar]",
        "sicher gasförmig unter [bar]", "sicher flüssig über [bar]"]
for sp, t in zip("BCDEFG", kopf):
    zelle(f"{sp}19", t, "kopf" if sp == "B" else "kopf2")
from CoolProp.CoolProp import PropsSI
T_liste = [-50, -40, -30, -20, -10, 0, 5, 10, 15, 20, 25, 27]
for i, T in enumerate(T_liste):
    r = 20 + i
    zelle(f"B{r}", T, "zelle_m")
    zelle(f"C{r}", PropsSI("P", "T", T + 273.15, "Q", 0, "CO2") / 1e5, "grau")
    zelle(f"D{r}", gw.taudruck(T), "grau"), zelle(f"E{r}", gw.blasendruck(T), "grau")
    zelle(f"F{r}", f"D{r}-U_PG", "formel", "f"), zelle(f"G{r}", f"E{r}+U_PG", "formel", "f")
r_hinweis = 20 + len(T_liste)
zelle(f"B{r_hinweis}", "Unter −56,6 °C bildet sich Trockeneis (nicht im Modell). Oberhalb von "
                       f"{de(pg['cricondentherm'][0])} °C gibt es beim Gemisch kein Zweiphasengebiet mehr.", "hinweis")

r0 = r_hinweis + 2
zelle(f"B{r0}", "Neue Grenzen der Einspeicherung – wo sie sich ergeben", "abschnitt")
for sp, t in zip("BCDEFG", ["Grenze", "", "reines CO₂", "Gemisch", "Folge für die Anlage", "Blatt"]):
    zelle(f"{sp}{r0 + 1}", t or None, "kopf" if sp == "B" else "kopf2")
grenzen = [
    ("Mindestdruck dichte Anlieferung bei 15 °C", "50,9 bar", f"{de(gw.blasendruck(15))} + 3 bar",
     f"S1 mit 91 bar sicher dicht ({de(A['p_S1'] - p_cb)} bar über der Cricondenbar)", "Einspeicherung S1"),
    ("Enddruck Verdichtung = Kühlerdruck bei 20 °C", "> 57,3 bar (gewählt 60, verflüssigen)",
     f"≥ {de(p_cb + U_PG)} bar (gewählt {A['p_V']:.0f})",
     "über der Cricondenbar kühlen: kein Zweiphasengebiet, Gemisch wird stetig dicht", "Einspeicherung S2"),
    ("Verdichtung 30 bar → p_V in einer Stufe", "73,2 °C ✔", f"{de(S2['T1a'])} °C ✘ (> 95 °C)",
     "zwei Verdichterstufen mit Zwischenkühlung nötig", "Einspeicherung S2"),
    ("Zwischendruck (gasförmig bei 20 °C)", "< 57,3 bar", f"≤ {de(Dv['p_zw_max'])} bar",
     f"zulässig ca. {Dv['p_zw_min']:.0f}–{Dv['p_zw_max']:.0f} bar (95-°C-Grenze Stufe 2)", "Einspeicherung S2"),
    ("Kopfdruck bei voller Kaverne", "107,6 bar", f"{de(p_kopf)} bar",
     f"Pumpe muss {de(p_kopf - REIN['p_kopf'])} bar mehr liefern; Gemisch ist leichter", "Kaverne & Gassäule"),
    ("Bohrloch ganz einphasig (Grenzfall)", "Kopf ≥ 73,8 bar (p_krit)", f"Kopf ≥ {de(p_cb)} bar (Cricondenbar)",
     f"Kavernendruck ab {de(K['grenz']['p_unten'])} bar (rein: 171,7 bar)", "Kaverne & Gassäule"),
    ("Leere Kaverne (70 bar unten)", "Kopf 45,0 bar, 0–236 m zweiphasig",
     f"Kopf {de(K['leer']['p_kopf'])} bar, 0–{K['leer']['tief_2ph']:.0f} m zweiphasig",
     "Zweiphasengebiet im oberen Bohrloch kleiner, aber vorhanden", "Kaverne & Gassäule"),
    ("Kritische Temperatur", "31,0 °C", f"{de(T_kG)} °C",
     "Kühlung unter 28 °C nötig, sonst ist das Gemisch vor der Pumpe überkritisch", "Einspeicherung S2"),
    ("Dichte am Kavernenboden (210 bar, 46 °C)", f"{de(REIN['rho_unten'])} kg/m³", f"{de(K['voll']['rho_unten'])} kg/m³",
     f"{de(100 * (K['voll']['rho_unten'] / REIN['rho_unten'] - 1))} % Masse je m³ Kavernenvolumen", "Kaverne & Gassäule"),
]
for i, (g_, rein, gem, folge, blatt) in enumerate(grenzen):
    r = r0 + 2 + i
    zelle(f"B{r}", g_, "zelle"), zelle(f"D{r}", rein, "zelle_m"), zelle(f"E{r}", gem, "zelle_m")
    zelle(f"F{r}", folge, "zelle"), zelle(f"G{r}", blatt, "zelle")
r_end = r0 + 2 + len(grenzen)

# Alte vs. neue Verdichtung (S2) - gleiche Rechnung wie Blatt "Einspeicherung S2"
ALT = W["alt"]
r1 = r_end + 1
zelle(f"B{r1}", "Alte und neue Verdichtung (S2, gasförmige Anlieferung 30 bar / 15 °C)", "abschnitt")
for sp, t in zip("BCDEFG", ["", "", "alte Verdichtung (wie reines CO₂)", "neue Verdichtung (Gemisch)",
                            "Warum", ""]):
    zelle(f"{sp}{r1 + 1}", t or None, "kopf" if sp == "B" else "kopf2")
vergleich = [
    ("Enddruck Verdichtung = Druck beim Kühlen auf 20 °C", f"{ALT['p_V']:.0f} bar", f"{A['p_V']:.0f} bar",
     f"Cricondenbar {de(p_cb)} bar + 3 bar Unsicherheit = {de(p_cb + U_PG)} bar Minimum"),
    ("Verdichterstufen", "1 Stufe", "2 Stufen (Zwischenkühlung auf 20 °C)",
     f"1 Stufe bis {A['p_V']:.0f} bar: {de(S2['T1a'])} °C > 95 °C (Q-016)"),
    ("Austrittstemperatur Verdichter", f"{de(ALT['T1a'])} °C", f"{de(S2v['T1a'])} / {de(S2v['T2a'])} °C",
     "höchstens 95 °C je Stufe (Q-016)"),
    ("Zwischendruck", "–", f"{de(S2v['p_zw'])} bar",
     f"gasförmig: ≤ Taudruck(20 °C) − 3 = {de(Dv['p_zw_max'])} bar"),
    ("Zustand nach Kühlung auf 20 °C", f"{ALT['phase']} ✘", f"{S2['phaseV']} ✔",
     f"alt: {de(ALT['abstand_tau'])} bar unter der Taulinie ({de(ALT['p_tau_TK'])} bar) – nichts kondensiert, "
     f"die Pumpe kann Gas nicht fördern; neu: {de(A['p_V'] - p_cb)} bar über der Cricondenbar"),
    ("Spezifische Arbeit Verdichter + Pumpe", "– (Kette funktioniert nicht)",
     f"{de(w_S2v)} kJ/kg", f"reines CO₂ (alt, 1 Stufe): {de(REIN['w_S2'])} kJ/kg; Durchverdichten Gemisch: "
                           f"{de(w_Dv)} kJ/kg"),
    ("Druck am Bohrlochkopf", "–", f"{de(p_kopf)} bar (Pumpe, {de(S2['TPa'])} °C)",
     "wird mit der neuen Kette erreicht"),
]
for i, (g_, alt, neu_, warum) in enumerate(vergleich):
    r = r1 + 2 + i
    zelle(f"B{r}", g_, "zelle"), zelle(f"D{r}", alt, "zelle_m"), zelle(f"E{r}", neu_, "zelle_m")
    zelle(f"F{r}", warum, "zelle")
r_vgl = (r1 + 1, r1 + 2 + len(vergleich))
r_end = r_vgl[1]

rows_xml = ""
for z in sorted(zeilen):
    cells = sorted(zeilen[z], key=lambda t: (len(t[0]), t[0]))
    if z == 1:
        ht = ' ht="30" customHeight="1"'
    elif z == 6:
        ht = ' ht="48" customHeight="1"'
    elif r_vgl[0] + 1 <= z <= r_vgl[1]:
        ht = ' ht="45" customHeight="1"'
    elif z in (19, r0 + 1, r_vgl[0]) or r0 + 2 <= z < r_vgl[0] - 1:
        ht = ' ht="30" customHeight="1"'
    else:
        ht = ""
    rows_xml += f'<row r="{z}"{ht}>' + "".join(c for _, c in cells) + "</row>"
sheet9 = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
          '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" '
          'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
          '<sheetPr><tabColor rgb="FFCA220E"/><pageSetUpPr fitToPage="1"/></sheetPr>'
          f'<dimension ref="B1:J{r_end}"/><sheetViews><sheetView showGridLines="0" workbookViewId="0"/></sheetViews>'
          '<sheetFormatPr defaultRowHeight="15"/>'
          '<cols><col min="1" max="1" width="2.7109375" customWidth="1"/>'
          '<col min="2" max="2" width="36.7109375" customWidth="1"/>'
          '<col min="3" max="3" width="14.7109375" customWidth="1"/>'
          '<col min="4" max="5" width="26.7109375" customWidth="1"/>'
          '<col min="6" max="6" width="30.7109375" customWidth="1"/>'
          '<col min="7" max="9" width="14.7109375" customWidth="1"/>'
          '<col min="10" max="10" width="44.7109375" customWidth="1"/></cols>'
          f'<sheetData>{rows_xml}</sheetData>'
          f'<mergeCells count="{3 + r_vgl[1] - r_vgl[0] + 1}"><mergeCell ref="B6:J6"/>'
          f'<mergeCell ref="B{r_hinweis}:J{r_hinweis}"/><mergeCell ref="F14:G14"/>'
          + "".join(f'<mergeCell ref="F{r}:J{r}"/>' for r in range(r_vgl[0], r_vgl[1] + 1))
          + '</mergeCells>'
          '<hyperlinks><hyperlink ref="B3" location="\'Übersicht\'!A1" display="← zur Übersicht"/></hyperlinks>'
          '<pageMargins left="0.5" right="0.5" top="0.6" bottom="0.6" header="0.3" footer="0.3"/>'
          '<pageSetup paperSize="9" orientation="landscape" fitToHeight="0"/></worksheet>')
pk.setze("xl/worksheets/sheet9.xml", sheet9)

# Mappe: Blatt anmelden (nach "Übersicht" einsortiert), Namen, Neuberechnung
wb = pk.text("xl/workbook.xml")
wb = wb.replace('<sheet name="Kaverne &amp; Gassäule" sheetId="2" r:id="rId2"/>',
                '<sheet name="Gemisch &amp; Grenzen" sheetId="9" r:id="rId20"/>'
                '<sheet name="Kaverne &amp; Gassäule" sheetId="2" r:id="rId2"/>')
namen = {"p_blase_TK": "'Gemisch &amp; Grenzen'!$H$15", "p_tau_TK": "'Gemisch &amp; Grenzen'!$H$16",
         "U_PG": "'Gemisch &amp; Grenzen'!$H$14", "p_cricondenbar": "'Gemisch &amp; Grenzen'!$H$11"}
zus = "".join(f'<definedName name="{n}">{ref}</definedName>' for n, ref in namen.items())
wb = wb.replace("</definedNames>", zus + "</definedNames>")
# benannte Bereiche müssen alphabetisch sortiert sein
dn = re.search(r"<definedNames>(.*?)</definedNames>", wb, re.S)
eintr = re.findall(r'<definedName name="[^"]*"[^>]*>[^<]*</definedName>', dn.group(1))
eintr.sort(key=lambda e: re.search(r'name="([^"]*)"', e).group(1).lower())
wb = wb.replace(dn.group(0), "<definedNames>" + "".join(eintr) + "</definedNames>")
wb = re.sub(r"<calcPr[^>]*/>", '<calcPr calcId="191029" fullCalcOnLoad="1"/>', wb)
wb = re.sub(r'activeTab="\d+"', 'activeTab="0"', wb)
pk.setze("xl/workbook.xml", wb)
rels = pk.text("xl/_rels/workbook.xml.rels")
rels = rels.replace('<Relationship Id="rId12" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/calcChain" Target="calcChain.xml"/>', "")
rels = rels.replace("</Relationships>", '<Relationship Id="rId20" Type="http://schemas.openxmlformats.org/'
                    'officeDocument/2006/relationships/worksheet" Target="worksheets/sheet9.xml"/></Relationships>')
pk.setze("xl/_rels/workbook.xml.rels", rels)
ct = pk.text("[Content_Types].xml")
ct = ct.replace('<Override PartName="/xl/calcChain.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.calcChain+xml"/>', "")
ct = ct.replace("</Types>", '<Override PartName="/xl/worksheets/sheet9.xml" ContentType="application/vnd.'
                "openxmlformats-officedocument.spreadsheetml.worksheet+xml\"/></Types>")
pk.setze("[Content_Types].xml", ct)
pk.entferne("xl/calcChain.xml")
app = pk.text("docProps/app.xml")
app = re.sub(r"<vt:i4>8</vt:i4>", "<vt:i4>9</vt:i4>", app, count=1)
app = app.replace("<vt:lpstr>Übersicht</vt:lpstr>", "<vt:lpstr>Übersicht</vt:lpstr><vt:lpstr>Gemisch &amp; Grenzen</vt:lpstr>")
app = re.sub(r'(<vt:vector size=")(\d+)(" baseType="lpstr">)', lambda m_: f"{m_.group(1)}{int(m_.group(2)) + 1}{m_.group(3)}", app, count=1)
pk.setze("docProps/app.xml", app)

sst.schreibe()
formelwerte_loeschen(pk)
pk.speichere(ZIEL)
print(f"gespeichert: {ZIEL}")

# =============================================================================
# 12) Gespeicherte Formelergebnisse (optional, braucht LibreOffice)
# =============================================================================
# Excel rechnet beim Öffnen alles neu (fullCalcOnLoad). Damit auch Vorschauen
# ohne Rechenkern die richtigen Zahlen zeigen, wird eine Kopie mit
# LibreOffice neu berechnet und deren Ergebnisse als <v> eingetragen.
import os
import shutil
import subprocess
import tempfile

import openpyxl
from xlsx_xml import werte_einsetzen

soffice = shutil.which("soffice") or shutil.which("libreoffice")
if soffice:
    tmp = tempfile.mkdtemp()
    kopie = os.path.join(tmp, "kopie.xlsx")
    shutil.copy(ZIEL, kopie)
    subprocess.run([soffice, "--headless", "--convert-to", "xlsx", "--outdir", os.path.join(tmp, "out"), kopie],
                   check=True, capture_output=True, timeout=300)
    wbv = openpyxl.load_workbook(os.path.join(tmp, "out", "kopie.xlsx"), data_only=True)
    wbf = openpyxl.load_workbook(ZIEL)
    werte_ = {}
    for ws in wbf.worksheets:
        werte_[ws.title] = {c.coordinate: wbv[ws.title][c.coordinate].value
                            for row in ws.iter_rows() for c in row
                            if isinstance(c.value, str) and c.value.startswith("=")}
    pk2 = Paket(ZIEL)
    wbx = pk2.text("xl/workbook.xml")
    rels2 = pk2.text("xl/_rels/workbook.xml.rels")
    dateien = {}
    for name, rid in re.findall(r'<sheet name="([^"]*)" sheetId="\d+" r:id="(rId\d+)"/>', wbx):
        ziel_ = re.search(rf'Id="{rid}"[^>]*Target="([^"]*)"', rels2).group(1)
        dateien[name.replace("&amp;", "&")] = "xl/" + ziel_
    werte_einsetzen(pk2, dateien, werte_)
    pk2.speichere(ZIEL)
    print(f"Formelergebnisse eingetragen ({sum(len(v) for v in werte_.values())} Formeln)")
else:
    print("LibreOffice nicht gefunden - Excel rechnet die Formeln beim Öffnen neu.")
